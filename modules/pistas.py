"""La pista que decide: qué dice por dentro un documento que se llama 'bfddfa36'.

Cuando una propuesta llega con un nombre que no dice nada ('PROPUESTA COMERCIAL
bfddfa36 RFI Soporte y-o administración.pptx') la IA no tiene con qué
clasificarla: ese nombre no dice si es de QA, de Portales o de IA. Preguntando
solo por el nombre, o contesta cualquier cosa, o deja el archivo sin clasificar
(en la raíz del destino, que es lo que se quería evitar).

Aquí se saca lo mínimo para que sí pueda decidir: el título que el documento
guarda en sus propiedades y las primeras líneas (una propuesta lo dice en la
primera página: 'Propuesta de servicios de aseguramiento de calidad y pruebas').
Se leen de los BYTES del archivo (ya bajado de SharePoint), sin escribir nada en
disco y sin abrir Word.

Lo que no se pueda leer devuelve '' (un .ppt viejo, un documento protegido, o
python-docx/python-pptx sin instalar): quien llama decide qué hacer sin pista.
"""
from io import BytesIO

# Cuánto texto se le manda a la IA por cada archivo. Con el título y un par de
# líneas sobra para saber de qué va, y más que esto es pagar tokens de relleno en
# cada clasificación.
MAXIMO = 700
# Y cuánto se mira de cada documento: las primeras páginas son las que dicen de
# qué va una propuesta (portada, índice y el primer apartado).
DIAPOSITIVAS = 3
PARRAFOS = 25
# Las líneas se juntan en una sola: la pista ocupa un renglón del prompt y no se
# confunde con la lista de archivos (que son los nombres que hay que copiar).
SEPARADOR = " · "


def pista_de(datos, nombre="", maximo=MAXIMO):
    """Un trozo del texto de dentro del documento ('' si no se puede leer).

    ``datos`` son los bytes del archivo y ``nombre`` su nombre: por la extensión
    se sabe cómo leerlo. Nunca lanza: un archivo corrupto, uno de los formatos
    viejos (.ppt, .doc) o una librería que no esté instalada devuelven vacío, que
    es lo mismo que quedarse sin pista.
    """
    if not datos:
        return ""
    extension = (nombre or "").rsplit(".", 1)[-1].lower()
    if extension == "pptx":
        texto = _de_pptx(datos)
    elif extension == "docx":
        texto = _de_docx(datos)
    else:
        texto = ""
    return _recortar(texto, maximo)


def _de_pptx(datos):
    """Título y texto de las primeras diapositivas de una presentación."""
    try:
        from pptx import Presentation
    except Exception:
        return ""
    try:
        presentacion = Presentation(BytesIO(datos))
    except Exception:
        return ""
    lineas = [_titulo(presentacion)]
    for diapositiva in list(presentacion.slides)[:DIAPOSITIVAS]:
        for forma in diapositiva.shapes:
            if not getattr(forma, "has_text_frame", False):
                continue
            for parrafo in forma.text_frame.paragraphs:
                lineas.append("".join(trozo.text for trozo in parrafo.runs))
    return _unir(lineas)


def _de_docx(datos):
    """Título y primeros párrafos de un documento de Word."""
    try:
        import docx
    except Exception:
        return ""
    try:
        documento = docx.Document(BytesIO(datos))
    except Exception:
        return ""
    lineas = [_titulo(documento)]
    for parrafo in list(documento.paragraphs)[:PARRAFOS]:
        lineas.append(parrafo.text or "")
    return _unir(lineas)


def _titulo(documento):
    """El título que el documento guarda en sus propiedades ('' si no tiene).

    Muchas propuestas lo traen ('Propuesta Servicios de QA - Cliente 2024') y con
    eso ya se sabe de qué van, sin leer una sola página.
    """
    try:
        return (documento.core_properties.title or "").strip()
    except Exception:
        return ""


def _unir(lineas):
    """Las líneas con texto, en una sola (las vacías y los espacios de más fuera)."""
    limpias = []
    for linea in lineas:
        limpia = " ".join(str(linea or "").split())
        if limpia:
            limpias.append(limpia)
    return SEPARADOR.join(limpias)


def _recortar(texto, maximo):
    """Deja la pista en ``maximo`` letras como mucho, cortando por palabra."""
    limpio = " ".join(str(texto or "").split())
    if not limpio or not maximo or len(limpio) <= maximo:
        return limpio
    recorte = limpio[:maximo]
    if " " in recorte[maximo // 2:]:
        recorte = recorte[:recorte.rindex(" ")]
    return recorte + "…"
