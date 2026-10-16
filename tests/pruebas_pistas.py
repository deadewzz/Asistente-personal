"""Pruebas de la pista del documento (lo que hace clasificable un nombre de hash).

Se ejecuta tal cual, desde la raíz del proyecto:

    python tests/pruebas_pistas.py

Un archivo como 'PROPUESTA COMERCIAL bfddfa36 RFI Soporte y-o
administración.pptx' no dice de qué va, así que la IA no puede clasificarlo: lo
único que hay para decidir es lo que el documento dice por dentro. Aquí se prueba
`modules.pistas.pista_de`, que saca eso de los BYTES del archivo (sin escribirlo
en disco y sin abrir Word): el título de sus propiedades y las primeras líneas (o
diapositivas).

Los .docx y .pptx de las pruebas se fabrican en memoria con python-docx y
python-pptx (los mismos que usa el módulo), así que se prueba la lectura de
verdad. Y también se prueba lo que hace cuando NO se puede leer (un formato viejo
.ppt/.doc, un archivo corrupto, uno protegido): devuelve '' y no lanza nada,
porque de esos se clasifica por el nombre, como antes.
"""
import os
import sys
from io import BytesIO

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from modules import pistas

fallos = []
total = [0]


def comprobar(condicion, texto):
    total[0] += 1
    if condicion:
        print(f"  ok  {texto}")
    else:
        print(f"  FALLO  {texto}")
        fallos.append(texto)


def docx_de(titulo, parrafos=()):
    """Un .docx de verdad, en memoria (título en las propiedades y párrafos)."""
    import docx
    documento = docx.Document()
    if titulo:
        documento.core_properties.title = titulo
    for parrafo in parrafos:
        documento.add_paragraph(parrafo)
    guardado = BytesIO()
    documento.save(guardado)
    return guardado.getvalue()


def pptx_de(titulo, textos=("",)):
    """Un .pptx de verdad, en memoria (una diapositiva por texto)."""
    from pptx import Presentation
    presentacion = Presentation()
    for texto in textos:
        diapositiva = presentacion.slides.add_slide(
            presentacion.slide_layouts[1])
        if titulo:
            diapositiva.shapes.title.text = titulo
        diapositiva.placeholders[1].text = texto
    guardado = BytesIO()
    presentacion.save(guardado)
    return guardado.getvalue()


def pruebas_lo_que_dice_un_docx():
    print("\n== Lo que dice por dentro un .docx ==")
    datos = docx_de("Propuesta Servicios de QA - Banco X 2024",
                    ["Propuesta de servicios de aseguramiento de calidad.",
                     "   ", "Alcance: pruebas de portales y automatización."])
    pista = pistas.pista_de(datos, "PROPUESTA COMERCIAL bfddfa36 RFI.docx")
    comprobar("Propuesta Servicios de QA - Banco X 2024" in pista,
              "sale el título que el documento guarda en sus propiedades")
    comprobar("aseguramiento de calidad" in pista
              and "pruebas de portales" in pista,
              "y las primeras líneas del texto")
    comprobar(pista.startswith("Propuesta Servicios de QA"),
              "el título va primero, que es lo que más dice de qué va")
    comprobar("\n" not in pista and "  " not in pista,
              "todo en un renglón, sin saltos ni espacios de más (así no se "
              "confunde con la lista de archivos del prompt)")
    comprobar(pistas.SEPARADOR in pista,
              "con las partes separadas para poder leerlas")

    comprobar(pistas.pista_de(docx_de("", ["Sin título en las propiedades"]),
                             "notas.docx")
              == "Sin título en las propiedades",
              "si el documento no trae título, se queda con su texto")

    muchos = docx_de("Informe", [f"Párrafo {i}" for i in range(1, 41)])
    pista = pistas.pista_de(muchos, "informe.docx")
    comprobar(f"Párrafo {pistas.PARRAFOS}" in pista
              and f"Párrafo {pistas.PARRAFOS + 1}" not in pista,
              f"solo se leen los primeros {pistas.PARRAFOS} párrafos (la "
              "portada y el índice, no el documento entero)")


def pruebas_lo_que_dice_un_pptx():
    print("\n== Lo que dice por dentro un .pptx ==")
    datos = pptx_de("Propuesta Comercial Portales",
                    ("Cliente: Banco X", "RFI Soporte y administración"))
    pista = pistas.pista_de(datos, "PROPUESTA COMERCIAL bfddfa36.pptx")
    comprobar("Propuesta Comercial Portales" in pista,
              "sale el título de la diapositiva")
    comprobar("RFI Soporte" in pista, "y el texto que traiga")

    muchas = pptx_de("Propuesta", tuple(f"Diapositiva {i}" for i in range(1, 7)))
    pista = pistas.pista_de(muchas, "propuesta.pptx")
    comprobar(f"Diapositiva {pistas.DIAPOSITIVAS}" in pista
              and f"Diapositiva {pistas.DIAPOSITIVAS + 1}" not in pista,
              f"solo se miran las {pistas.DIAPOSITIVAS} primeras diapositivas")


def pruebas_lo_que_no_se_puede_leer():
    """Lo que no se puede leer devuelve '' y no lanza (de eso se decide igual).

    Es lo que hace que la clasificación nunca se caiga por un documento raro: los
    formatos viejos (.ppt, .doc), los protegidos, los que llegaron a medias y los
    que no son documentos de oficina se quedan sin pista, y de esos el plan se
    decide por el nombre, como se hacía antes de que existiera esto.
    """
    print("\n== Lo que no se puede leer ('' y sin lanzar) ==")
    comprobar(pistas.pista_de(b"", "acta.docx") == "", "sin datos no hay pista")
    comprobar(pistas.pista_de(None, "acta.docx") == "",
              "y con None tampoco (nunca revienta)")
    comprobar(pistas.pista_de(b"cualquier cosa", "acta.docx") == "",
              "un .docx corrupto o a medias devuelve '' (no lanza)")
    comprobar(pistas.pista_de(b"cualquier cosa", "propuesta.pptx") == "",
              "y lo mismo con una presentación rota")
    for nombre in ("transcripcion.doc", "presentacion.ppt", "acta.pdf",
                   "informe.txt", "sin_extension"):
        comprobar(pistas.pista_de(docx_de("Propuesta de QA"), nombre) == "",
                  f"'{nombre}' no se lee (formato viejo o no es un documento de "
                  "oficina nuevo)")
    comprobar(pistas.pista_de(docx_de("Propuesta de QA"), "") == "",
              "sin nombre no se sabe cómo leerlo")
    comprobar(pistas.pista_de(docx_de("Propuesta de QA"), "PROPUESTA.DOCX")
              != "",
              "y la extensión se mira sin importar mayúsculas")


def pruebas_el_recorte():
    """Se manda lo justo: el título y un poco de texto, no el documento entero.

    Cada letra de la pista se paga en cada clasificación (va en el prompt), y con
    el título y un par de líneas sobra para saber de qué va el documento.
    """
    print("\n== El recorte (lo justo para decidir, sin pagar de más) ==")
    largo = " ".join(f"palabra{i}" for i in range(500))
    pista = pistas.pista_de(docx_de("Título que dice de qué va", [largo]),
                            "propuesta.docx")
    comprobar(len(pista) <= pistas.MAXIMO + 1,
              f"la pista no pasa de {pistas.MAXIMO} letras (más es relleno "
              "pagado en cada archivo)")
    comprobar(pista.endswith("…"), "y se nota que está cortada (termina en '…')")
    comprobar(pista.startswith("Título que dice de qué va"),
              "el principio (el título) va siempre: es lo que más dice")
    comprobar("palabra499" not in pista, "y lo del final se queda fuera")

    comprobar(pistas.pista_de(docx_de("Corto", ["Una línea."]), "acta.docx")
              == "Corto" + pistas.SEPARADOR + "Una línea.",
              "si cabe entera, la pista va completa (sin '…' de más)")

    completa = "Corto" + pistas.SEPARADOR + "Una línea."
    pista = pistas.pista_de(docx_de("Corto", ["Una línea."]), "acta.docx",
                            maximo=12)
    comprobar(len(pista) <= 13 and pista.endswith("…"),
              "con un tope pequeño también corta")
    comprobar(pista.rstrip("…").split()[-1] in completa.split(),
              "y corta por palabra (no deja una a medias)")

    comprobar(pistas.pista_de(docx_de("Título", ["Texto"]), "acta.docx",
                              maximo=1000)
              == "Título" + pistas.SEPARADOR + "Texto",
              "con un tope grande cabe todo lo que hay")


def main_pruebas():
    pruebas_lo_que_dice_un_docx()
    pruebas_lo_que_dice_un_pptx()
    pruebas_lo_que_no_se_puede_leer()
    pruebas_el_recorte()
    print(f"\nTOTAL: {total[0]} comprobaciones, {len(fallos)} fallos")
    for fallo in fallos:
        print(f"  - {fallo}")
    return 1 if fallos else 0


if __name__ == "__main__":
    raise SystemExit(main_pruebas())
