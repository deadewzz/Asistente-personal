"""Pruebas de SharePoint sin Graph.

Se ejecuta tal cual, desde la raíz del proyecto:

    python tests/pruebas_sharepoint.py

Prueba la lógica que no necesita Brave ni red: parseo de URLs de
SharePoint/OneDrive, construcción de las llamadas `_api` (con una pestaña
falsa), descarga/subida/copia de archivos, resolución de la URL que pega el
usuario (la raíz de un sitio -> su biblioteca), clasificación y resumen de
desplazamientos, el filtro que deja pasar solo documentos de oficina (ppt,
pptx, doc y docx) y el catálogo de carpetas del DESTINO que se le enseña a la IA
(para que use las que ya existen en vez de inventarse otras parecidas). Las
llamadas reales se validan contra la sesión del usuario
con `tests/diagnostico_brave.py` y `tests/pruebas_copia_en_vivo.py`.
"""
import base64
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from modules import navegador_cdp
from modules.sharepoint import ErrorSharePoint, SharePoint

# De una biblioteca de SharePoint solo salen los documentos de oficina
# (``SharePoint.EXTENSIONES_UTILES``), pero casi todos los dobles de este
# archivo son de antes de ese filtro y usan nombres .pdf/.zip/.txt para probar
# el parseo de URLs, los listados y los movimientos. El filtro tiene sus
# propias pruebas (``pruebas_solo_documentos_de_oficina``), así que aquí se deja
# apagado (lista vacía = se pasa todo) para no renombrar medio centenar de
# dobles. Ese filtro no tiene nada que ver con lo que se prueba aquí (URLs,
# listados y movimientos), así que cada prueba sigue siendo sobre lo suyo.
EXTENSIONES_REALES = SharePoint.EXTENSIONES_UTILES
SharePoint.EXTENSIONES_UTILES = ()

fallos = []
total = [0]


def comprobar(condicion, texto):
    total[0] += 1
    if condicion:
        print(f"  ok  {texto}")
    else:
        print(f"  FALLO  {texto}")
        fallos.append(texto)


class PestanaFalsa:
    """Sustituye a navegador_cdp.Pestana: devuelve una respuesta fija.

    Si se le pasa una lista, va devolviendo una por llamada: así se prueban los
    reintentos (por ejemplo, `Files/add` -> `AddUsingPath`). Con `viva=False` se
    imita una pestaña que Brave ya cerró.
    """

    def __init__(self, respuesta, viva=True):
        self.respuesta = respuesta
        self.viva = viva
        self.visto = []

    def evaluar(self, expresion, navegar=None):
        self.visto.append(expresion)
        if isinstance(self.respuesta, list):
            indice = min(len(self.visto) - 1, len(self.respuesta) - 1)
            return self.respuesta[indice]
        return self.respuesta

    def sigue_viva(self):
        """Como la de verdad: dice si Brave todavía tiene esta pestaña."""
        return self.viva

    def cerrar(self):
        pass


SITIO = "https://contoso.sharepoint.com/sites/Tech"


def con_pestana(respuesta):
    """SharePoint con una pestaña falsa ya inyectada (sin tocar Brave)."""
    sp = SharePoint()
    pestana = PestanaFalsa(respuesta)
    sp._pestanas[SITIO] = pestana
    return sp, pestana


def pruebas_parseo_url():
    print("\n== Parseo de URLs de SharePoint ==")
    p = SharePoint.parsear_url(
        "https://contoso.sharepoint.com/sites/Tech/Documentos%20compartidos/Sub")
    comprobar(p["sitio"] == SITIO, "sitio de una URL /sites/")
    comprobar(p["carpeta_rel"] == "/sites/Tech/Documentos compartidos/Sub",
              "carpeta server-relative de una URL /sites/")
    comprobar(p["dominio"] == "contoso.sharepoint.com", "dominio extraído")
    comprobar(p["web_rel"] == "/sites/Tech", "ruta del sitio dentro del host")

    p = SharePoint.parsear_url("https://contoso.sharepoint.com/teams/Tech/Archivo")
    comprobar(p["sitio"] == "https://contoso.sharepoint.com/teams/Tech",
              "sitio de una URL /teams/")
    comprobar(p["carpeta_rel"] == "/teams/Tech/Archivo",
              "carpeta relativa de una URL /teams/")

    p = SharePoint.parsear_url(
        "https://contoso-my.sharepoint.com/personal/john_doe_contoso_com/Documents/Invoices")
    comprobar(p["sitio"] == "https://contoso-my.sharepoint.com/personal/john_doe_contoso_com",
              "sitio de OneDrive personal")
    comprobar(p["carpeta_rel"] == "/personal/john_doe_contoso_com/Documents/Invoices",
              "carpeta de OneDrive personal")

    p = SharePoint.parsear_url("https://contoso.sharepoint.com/")
    comprobar(p["sitio"] == "https://contoso.sharepoint.com/"
              and p["carpeta_rel"] == "/", "URL raíz del sitio")

    # Las URLs que se copian de SharePoint traen la vista web y sus parámetros.
    p = SharePoint.parsear_url(
        "https://contoso.sharepoint.com/sites/Tech/Oportunidad/Forms/"
        "AllItems.aspx?viewid=cc6180f7%2Dd03f%2D4ee0%2Da546%2D0c7a85fe3902"
        "&FolderCTID=0x0120002B8F763F31F15643983416479C046D67"
        "&CT=1785254039012")
    comprobar(p["sitio"] == SITIO, "sitio de una URL con vista Forms/AllItems.aspx")
    comprobar(p["carpeta_rel"] == "/sites/Tech/Oportunidad",
              "la vista AllItems.aspx se queda con su carpeta, no con el .aspx")

    p = SharePoint.parsear_url(
        "https://contoso.sharepoint.com/sites/Tech/Oportunidad/Forms/"
        "AllItems.aspx?RootFolder=%2Fsites%2FTech%2FOportunidad%2F2026"
        "&FolderCTID=0x0120")
    comprobar(p["carpeta_rel"] == "/sites/Tech/Oportunidad/2026",
              "RootFolder= manda sobre la ruta de la vista")

    p = SharePoint.parsear_url(
        "https://contoso.sharepoint.com/sites/Tech/Oportunidad/Forms/"
        "AllItems.aspx?id=%2Fsites%2FTech%2FOportunidad%2F2026%2Ffactura.pdf")
    comprobar(p["carpeta_rel"] == "/sites/Tech/Oportunidad/2026",
              "si id= apunta a un archivo, se usa su carpeta")

    p = SharePoint.parsear_url(
        "https://contoso.sharepoint.com/sites/Tech/Docs/Forms/OneUp.aspx?"
        "sourcedoc=%7BABC%7D")
    comprobar(p["carpeta_rel"] == "/sites/Tech/Docs",
              "cualquier vista .aspx (OneUp) se recorta")

    # La carpeta del parámetro puede estar en OTRO sitio del mismo host: la
    # vista 'Mi unidad' de OneDrive vive en el host -my, pero la carpeta es la
    # del sitio personal. Aquí se llamaba al web raíz del host y SharePoint
    # devolvía 403 (comprobado en vivo contra un tenant real).
    p = SharePoint.parsear_url(
        "https://contoso-my.sharepoint.com/my?"
        "id=%2Fpersonal%2Fjohn%5Fdoe%5Fcontoso%5Fcom%2FDocuments%2FPropuestas"
        "&viewid=7a54bd05%2D155b%2D4a81%2D8dfc%2D21990001e908")
    comprobar(p["sitio"] == "https://contoso-my.sharepoint.com/personal/"
                            "john_doe_contoso_com",
              "la vista 'Mi unidad' usa el sitio personal, no la raíz del host")
    comprobar(p["carpeta_rel"] == "/personal/john_doe_contoso_com/Documents/"
                                 "Propuestas",
              "y su carpeta sigue siendo la de id=")
    comprobar(p["web_rel"] == "/personal/john_doe_contoso_com",
              "la ruta del sitio sale de la carpeta de id=")

    p = SharePoint.parsear_url(
        "https://contoso.sharepoint.com/sites/Tech/Docs/Forms/AllItems.aspx?"
        "id=%2Fsites%2FOtro%2FDocumentos")
    comprobar(p["sitio"] == "https://contoso.sharepoint.com/sites/Otro",
              "un id= de otro sitio apunta a ese sitio")

    p = SharePoint.parsear_url(
        "https://contoso.sharepoint.com/Documentos%20compartidos/Forms/"
        "AllItems.aspx?id=%2FDocumentos%20compartidos%2FSub")
    comprobar(p["sitio"] == "https://contoso.sharepoint.com/",
              "un id= del web raíz no inventa un sitio que no existe")


def pruebas_odata():
    print("\n== Literales OData ==")
    comprobar(SharePoint._odata("/sites/Tech/Documentos compartidos")
              == "/sites/Tech/Documentos%20compartidos",
              "escapa espacios y deja las barras")
    comprobar(SharePoint._odata("/sites/Tech/O'Brien") == "/sites/Tech/O''Brien",
              "duplica la comilla simple")


def pruebas_api():
    print("\n== Llamadas a /_api (pestaña falsa) ==")
    sp, pestana = con_pestana({"estado": 200, "cuerpo": '{"d": {"results": [1]}}'})
    comprobar(sp._api(SITIO, "/web/x") == {"d": {"results": [1]}},
              "parsea el JSON de una respuesta 200")
    js = pestana.visto[-1]
    comprobar("credentials:'include'" in js, "el fetch manda las cookies de sesión")
    comprobar("_api/web/x" in js, "llama al endpoint pedido")

    sp, pestana = con_pestana({"estado": 200, "cuerpo": "{}"})
    sp._api(SITIO, "/web/folders", metodo="POST",
            datos={"__metadata": {"type": "SP.Folder"}})
    js = pestana.visto[-1]
    comprobar("contextinfo" in js and "X-RequestDigest" in js,
              "pide y reutiliza el token de escritura")
    comprobar("SP.Folder" in js, "envía el cuerpo JSON de la petición")
    comprobar("credentials:'include'" in js, "también las cookies en el POST")

    for estado, esperado, texto in (
            (401, "sesión", "el 401 avisa de que la sesión no vale"),
            (403, "permiso", "el 403 avisa de que falta permiso (no de sesión)"),
            (500, "500", "el 500 muestra el código de error")):
        sp, _ = con_pestana({"estado": estado, "cuerpo": "vaya"})
        try:
            sp._api(SITIO, "/web/x")
            comprobar(False, texto)
        except ErrorSharePoint as e:
            comprobar(esperado in str(e), texto)

    sp, _ = con_pestana({"estado": -1, "cuerpo": "NetworkError"})
    try:
        sp._api(SITIO, "/web/x")
        comprobar(False, "un fallo de red avisa de conexión/Brave cerrado")
    except ErrorSharePoint as e:
        comprobar("navegador" in str(e),
                  "un fallo de red avisa de conexión/Brave cerrado")

    sp, _ = con_pestana({"estado": 403, "cuerpo": "Acceso denegado",
                         "paso": "token"})
    try:
        sp._api(SITIO, "/web/x", metodo="POST", datos={})
        comprobar(False, "avisa si falla el token de escritura")
    except ErrorSharePoint as e:
        comprobar("token de escritura" in str(e),
                  "avisa si falla el token de escritura")

    sp, _ = con_pestana("texto suelto")
    try:
        sp._api(SITIO, "/web/x")
        comprobar(False, "una respuesta rara lanza ErrorSharePoint")
    except ErrorSharePoint:
        comprobar(True, "una respuesta rara lanza ErrorSharePoint")

    sp, _ = con_pestana({"estado": 200,
                         "cuerpo": "<html><head><title>Iniciar sesión"})
    try:
        sp._api(SITIO, "/web/x")
        comprobar(False, "detecta la página de login en vez de datos")
    except ErrorSharePoint as e:
        comprobar("sesión" in str(e), "detecta la página de login en vez de datos")


def pruebas_listados():
    print("\n== Listar archivos y carpetas ==")
    sp = SharePoint()
    visto = {}

    def api_falsa(sitio, ruta, metodo="GET", datos=None):
        visto.update(sitio=sitio, ruta=ruta, metodo=metodo)
        if ruta.endswith("Folders?$select=Name,ServerRelativeUrl"):
            return {"d": {"results": [{"Name": "Facturas",
                                       "ServerRelativeUrl": "/sites/Tech/Docs/Facturas"}]}}
        return {"d": {"results": [{"Name": "a.pdf",
                                   "ServerRelativeUrl": "/sites/Tech/Docs/a.pdf",
                                   "Length": 12}]}}

    sp._api = api_falsa
    archivos = sp.listar_archivos("https://contoso.sharepoint.com/sites/Tech/Docs")
    comprobar(archivos == [{"nombre": "a.pdf",
                            "ruta_rel": "/sites/Tech/Docs/a.pdf",
                            "subruta": "a.pdf", "tamano": 12}],
              "listar_archivos mapea nombre/ruta/subruta/tamaño")
    comprobar(visto["sitio"] == SITIO, "usa la URL raíz del sitio")
    comprobar("/sites/Tech/Docs" in visto["ruta"],
              "usa la ruta server-relative de la carpeta")

    carpetas = sp.listar_carpetas("https://contoso.sharepoint.com/sites/Tech/Docs")
    comprobar(carpetas == [{"nombre": "Facturas",
                            "ruta_rel": "/sites/Tech/Docs/Facturas"}],
              "listar_carpetas mapea nombre/ruta")

    # La API v2 manda los enteros grandes entre comillas (vienen de .NET).
    def api_texto(sitio, ruta, metodo="GET", datos=None):
        if ruta.endswith("Folders?$select=Name,ServerRelativeUrl"):
            return {"d": {"results": []}}
        return {"d": {"results": [
            {"Name": "grande.zip",
             "ServerRelativeUrl": "/sites/Tech/Docs/grande.zip",
             "Length": "25165824"},
            {"Name": "raro.bin",
             "ServerRelativeUrl": "/sites/Tech/Docs/raro.bin",
             "Length": None}]}}

    sp._api = api_texto
    archivos = sp.listar_archivos("https://contoso.sharepoint.com/sites/Tech/Docs")
    comprobar([a["tamano"] for a in archivos] == [25165824, 0],
              "el tamaño se convierte a número aunque venga como texto "
              "(o vacío)")


def pruebas_busqueda_en_subcarpetas():
    """La carpeta de arriba solo tiene carpetas: los archivos están dentro.

    Es el caso real de una biblioteca por clientes: al pasar 'la carpeta
    Colombia' hay que entrar en la carpeta de cada cliente (AXA, SURA...) para
    encontrar los documentos; si no, parecería vacía.
    """
    print("\n== Buscar archivos dentro de las subcarpetas ==")
    sp = SharePoint()
    pedidas = []

    def api_falsa(sitio, ruta, metodo="GET", datos=None):
        pedidas.append(ruta)
        if "'/sites/Tech/Colombia'" in ruta and "Folders?" in ruta:
            return {"d": {"results": [
                {"Name": "AXA", "ServerRelativeUrl": "/sites/Tech/Colombia/AXA",
                 "Files": {"results": [
                     {"Name": "acta-axa.pdf", "Length": "120",
                      "ServerRelativeUrl":
                          "/sites/Tech/Colombia/AXA/acta-axa.pdf"}]}},
                {"Name": "SURA", "ServerRelativeUrl": "/sites/Tech/Colombia/SURA",
                 "Files": {"results": []}}]}}
        if "'/sites/Tech/Colombia/AXA'" in ruta and "Folders?" in ruta:
            return {"d": {"results": [
                {"Name": "2024",
                 "ServerRelativeUrl": "/sites/Tech/Colombia/AXA/2024",
                 "Files": {"results": [
                     {"Name": "informe.pdf", "Length": 300,
                      "ServerRelativeUrl":
                          "/sites/Tech/Colombia/AXA/2024/informe.pdf"}]}}]}}
        if ruta.endswith("Files?$select=Name,ServerRelativeUrl,Length"):
            return {"d": {"results": []}}   # la carpeta pedida: solo carpetas
        return {}

    sp._api = api_falsa
    url = "https://contoso.sharepoint.com/sites/Tech/Colombia"
    comprobar(sp.listar_archivos(url) == [],
              "sin pedir profundidad no mira dentro (arriba no hay archivos)")

    archivos = sp.listar_archivos(url, profundidad=2)
    comprobar([a["subruta"] for a in archivos]
              == ["AXA/acta-axa.pdf", "AXA/2024/informe.pdf"],
              "con profundidad entra en cada cliente y saca sus archivos")
    comprobar(archivos[0]["ruta_rel"] == "/sites/Tech/Colombia/AXA/acta-axa.pdf"
              and archivos[1]["tamano"] == 300,
              "guarda la ruta real y el tamaño de cada archivo")
    comprobar(sp.nota_listado == "",
              "y no avisa de nada, porque los pudo ver todos")

    antes = len(pedidas)
    un_nivel = sp.listar_archivos(url, profundidad=1)
    comprobar([a["subruta"] for a in un_nivel] == ["AXA/acta-axa.pdf"],
              "con profundidad=1 se queda en las carpetas de primer nivel")
    comprobar(sum(1 for r in pedidas[antes:] if "Folders?" in r) == 1,
              "el primer nivel se pide en UNA llamada (subcarpetas con sus "
              "archivos dentro)")

    # Topes: con una biblioteca enorme se para y lo deja dicho.
    sp2 = SharePoint()
    sp2.MAXIMO_CARPETAS = 1

    def api_ancha(sitio, ruta, metodo="GET", datos=None):
        if ruta.endswith("Files?$select=Name,ServerRelativeUrl,Length"):
            return {"d": {"results": []}}
        if "'/sites/Tech/Colombia'" in ruta:
            return {"d": {"results": [
                {"Name": f"Cliente{i}", "ServerRelativeUrl": f"/x/Cliente{i}",
                 "Files": {"results": [
                     {"Name": f"f{i}.pdf", "Length": 1,
                      "ServerRelativeUrl": f"/x/Cliente{i}/f{i}.pdf"}]}}
                for i in range(3)]}}
        return {"d": {"results": []}}

    sp2._api = api_ancha
    archivos = sp2.listar_archivos(url, profundidad=2)
    comprobar(len(archivos) == 3, "recoge lo que sí pudo ver")
    comprobar("Me quedé en 1 carpetas" in sp2.nota_listado,
              "y avisa de que dejó carpetas sin mirar")


def pruebas_paginado():
    """SharePoint corta los listados: sin seguir ``d.__next`` faltan archivos."""
    print("\n== Listados con más de una página ($skiptoken) ==")
    sp = SharePoint()
    llamadas = []
    enlace = ("https://contoso.sharepoint.com/sites/Tech/_api/web/"
              "GetFolderByServerRelativeUrl('/sites/Tech/Docs')/Files"
              "?$skiptoken=Paged=TRUE")

    def api_falsa(sitio, ruta, metodo="GET", datos=None):
        llamadas.append(ruta)
        if "skiptoken" in ruta:
            return {"d": {"results": [{"Name": "b.pdf", "Length": 20,
                                       "ServerRelativeUrl":
                                           "/sites/Tech/Docs/b.pdf"}]}}
        if ruta.endswith("Files?$select=Name,ServerRelativeUrl,Length"):
            return {"d": {"results": [{"Name": "a.pdf", "Length": 10,
                                       "ServerRelativeUrl":
                                           "/sites/Tech/Docs/a.pdf"}],
                          "__next": enlace}}
        return {}

    sp._api = api_falsa
    archivos = sp.listar_archivos("https://contoso.sharepoint.com/sites/Tech/Docs")
    comprobar([a["nombre"] for a in archivos] == ["a.pdf", "b.pdf"],
              "sigue la página siguiente del listado")
    comprobar(llamadas[-1].startswith("/web/") and "skiptoken" in llamadas[-1],
              "la pide como ruta /_api (sin el host ni '/_api')")

    comprobar(sp.listar_carpetas("https://contoso.sharepoint.com/sites/Tech/Docs")
              == [], "un listado de carpetas vacío no rompe nada")
    comprobar(SharePoint._siguiente_pagina("", "/web/x?$top=1") == "",
              "sin enlace no hay página siguiente")
    comprobar(SharePoint._siguiente_pagina("$skiptoken=ABC", "/web/x?$top=1")
              == "/web/x?$skiptoken=ABC",
              "también acepta el enlace en forma de consulta suelta")


def pruebas_avisos_de_progreso():
    """El módulo va contando lo que hace, para que el chat no se quede mudo.

    Una biblioteca por clientes tarda minutos: qué nivel está mirando, cuántos
    archivos lleva, si el listado viene por partes, qué carpeta va a usar... Con
    esos avisos se ve que sigue trabajando (y de qué sale lo que tarda).
    """
    print("\n== Avisos de progreso (lo que el chat va diciendo) ==")
    sp = SharePoint()
    comprobar(sp.on_aviso is None, "sin ventana que escuche, no hay oyente")
    sp._aviso("hola")            # no debe fallar ni avisar a nadie
    comprobar(True, "un aviso sin oyente no rompe nada")

    avisos = []
    sp.on_aviso = avisos.append
    sp._aviso("hola")
    comprobar(avisos == ["hola"], "con oyente, el texto llega tal cual")

    def aviso_roto(texto):
        raise RuntimeError("el chat se ha caído")

    sp.on_aviso = aviso_roto
    sp._aviso("hola")
    comprobar(True, "si el canal de avisos falla, la tarea continúa")

    # Buscar dentro de las subcarpetas: se avisa de cada nivel y del resultado.
    sp = SharePoint()
    avisos = []
    sp.on_aviso = avisos.append

    def api_falsa(sitio, ruta, metodo="GET", datos=None):
        if ruta.endswith("Files?$select=Name,ServerRelativeUrl,Length"):
            return {"d": {"results": [
                {"Name": "notas.txt", "Length": 5,
                 "ServerRelativeUrl": "/sites/Tech/Colombia/notas.txt"}]}}
        if "'/sites/Tech/Colombia'" in ruta and "Folders?" in ruta:
            return {"d": {"results": [
                {"Name": "AXA", "ServerRelativeUrl": "/sites/Tech/Colombia/AXA",
                 "Files": {"results": [
                     {"Name": "acta.pdf", "Length": 10,
                      "ServerRelativeUrl":
                          "/sites/Tech/Colombia/AXA/acta.pdf"}]}}]}}
        return {"d": {"results": []}}

    sp._api = api_falsa
    archivos = sp.listar_archivos(
        "https://contoso.sharepoint.com/sites/Tech/Colombia", profundidad=2)
    comprobar([a["subruta"] for a in archivos] == ["notas.txt", "AXA/acta.pdf"],
              "lista lo de la carpeta y lo que hay dentro de la subcarpeta")
    comprobar(any("mirando también dentro de las subcarpetas" in a
                  for a in avisos),
              "avisa de que va a mirar dentro de las subcarpetas")
    comprobar(any("nivel 1 de 2" in a for a in avisos), "dice por qué nivel va")
    comprobar(any(a.startswith("Encontré 1 archivo(s) dentro de las subcarpetas")
                  for a in avisos),
              "y termina contando cuántos archivos encontró")

    # Listado paginado: se avisa de que sigue pidiendo páginas.
    sp = SharePoint()
    avisos = []
    sp.on_aviso = avisos.append

    def api_paginas(sitio, ruta, metodo="GET", datos=None):
        if "skiptoken" in ruta:
            return {"d": {"results": [
                {"Name": "b.pdf", "Length": 20,
                 "ServerRelativeUrl": "/sites/Tech/Docs/b.pdf"}]}}
        if ruta.endswith("Files?$select=Name,ServerRelativeUrl,Length"):
            return {"d": {
                "results": [{"Name": "a.pdf", "Length": 10,
                             "ServerRelativeUrl": "/sites/Tech/Docs/a.pdf"}],
                "__next": "https://contoso.sharepoint.com/sites/Tech/_api"
                          "/web/x?$skiptoken=Paged=TRUE"}}
        return {"d": {"results": []}}

    sp._api = api_paginas
    sp.listar_archivos("https://contoso.sharepoint.com/sites/Tech/Docs")
    comprobar(any("pidiendo la página 2" in a for a in avisos),
              "si el listado viene por partes, lo cuenta")

    # URL del sitio a secas: se avisa de qué carpeta de documentos se va a usar.
    sp = SharePoint()
    avisos = []
    sp.on_aviso = avisos.append
    sp._api = lambda sitio, ruta, metodo="GET", datos=None: {
        "d": {"results": [{"RootFolder":
                           {"ServerRelativeUrl": "/sites/Tech/Docs"}}]}}
    resultado = sp.resolver("https://contoso.sharepoint.com/sites/Tech")
    comprobar(resultado.endswith("/sites/Tech/Docs"),
              "la URL del sitio se resuelve a su biblioteca de documentos")
    comprobar(any("usaré su carpeta de documentos" in a for a in avisos),
              "y lo avisa (así se sabe de dónde salen los archivos)")

    # Trasiego de archivos: cada uno que pasa se cuenta.
    sp = SharePoint()
    avisos = []
    sp.on_aviso = avisos.append
    sp.carpeta_existe = lambda sitio, ruta: True
    sp.archivo_existe = lambda sitio, ruta: False
    sp.listar_carpetas = lambda url: []
    sp.crear_carpeta = lambda sitio, ruta: None
    sp.mover = lambda sitio, origen, destino, nombre, tamano=None: None
    resumen = sp.clasificar_y_mover(
        "https://contoso.sharepoint.com/sites/Tech/Colombia",
        "https://contoso.sharepoint.com/sites/Tech/Archivo",
        [("AXA/acta.pdf", "Actas"), ("SURA/informe.pdf", "Informes")])
    comprobar(any("Moviendo 2 archivo(s)" in a for a in avisos),
              "al empezar dice cuántos archivos va a pasar")
    comprobar([a for a in avisos if a.startswith("Moviendo 1 de 2")
               and a.endswith("→ Actas")],
              "y va contando archivo a archivo, con la carpeta de destino")
    comprobar("Movidos 2 de 2" in resumen, "el resumen final sigue igual")


def pruebas_plan_con_rutas():
    """Con archivos en subcarpetas, el plan lleva la carpeta de cada uno.

    'AXA/2024/acta.pdf' es la carpeta del ORIGEN (con ella se encuentra el
    archivo); al destino va solo 'acta.pdf', dentro de la subcarpeta de la regla.
    """
    print("\n== Plan de clasificación con archivos en subcarpetas ==")
    archivos = [
        {"nombre": "acta.pdf", "subruta": "AXA/acta.pdf", "tamano": 1},
        {"nombre": "acta.pdf", "subruta": "SURA/acta.pdf", "tamano": 1},
        {"nombre": "informe.pdf", "subruta": "AXA/2024/informe.pdf", "tamano": 1},
    ]
    plan = SharePoint.ajustar_plan(
        [("axa/acta.pdf", "Clientes"), ("informe.pdf", "Legal"),
         ("no-existe.pdf", "Otros"), ("acta.pdf", "X")], archivos)
    comprobar(plan[0] == ("AXA/acta.pdf", "Clientes"),
              "respeta la ruta completa (sin importar mayúsculas)")
    comprobar(plan[1] == ("AXA/2024/informe.pdf", "Legal"),
              "un nombre suelto se completa si identifica a un solo archivo")
    comprobar(plan[2] == ("no-existe.pdf", "Otros"),
              "lo que no existe se deja tal cual (el fallo se verá en el resumen)")
    comprobar(plan[3] == ("acta.pdf", "X"),
              "con dos 'acta.pdf' no adivina cuál es")

    sp = SharePoint()
    movidos = []
    sp.listar_carpetas = lambda url: []
    sp.crear_carpeta = lambda sitio, rel: None
    sp.carpeta_existe = lambda sitio, rel: True
    sp.archivo_existe = lambda sitio, rel: False
    sp.mover = lambda sitio, origen, destino, nombre, tamano=None: movidos.append(
        (origen, destino, nombre))
    resumen = sp.clasificar_y_mover(
        "https://contoso.sharepoint.com/sites/Tech/Colombia",
        "https://contoso.sharepoint.com/sites/Tech/Archivo",
        [("AXA/2024/informe.pdf", "Legal")])
    comprobar(movidos == [("/sites/Tech/Colombia/AXA/2024/informe.pdf",
                           "/sites/Tech/Archivo/Legal", "informe.pdf")],
              "lo saca de su subcarpeta y al destino llega con su nombre")
    comprobar("AXA/2024/informe.pdf → Legal" in resumen,
              "el resumen enseña de dónde sale y a dónde va")


def con_api_estado(estado, cuerpo=""):
    """SharePoint cuyo `_api` siempre responde con ese código HTTP."""
    sp = SharePoint()

    def api_falsa(sitio, ruta, metodo="GET", datos=None):
        if estado == 200:
            return {}
        error = ErrorSharePoint(f"SharePoint respondió {estado}: {cuerpo}")
        error.estado = estado
        raise error

    sp._api = api_falsa
    return sp


def pruebas_existencia():
    """SharePoint contesta 200 con lista vacía a rutas que no existen.

    Por eso el código pregunta explícitamente por la carpeta/archivo antes de
    fiarse de un listado: si no, una carpeta mal escrita parecería vacía y un
    archivo del destino sin detectar se sobrescribiría.
    """
    print("\n== Existencia de carpetas y archivos ==")
    sp = con_api_estado(404, "Archivo no encontrado")
    comprobar(sp.carpeta_existe(SITIO, "/sites/Tech/Docs") is False,
              "404 en la carpeta = no existe")
    comprobar(sp.archivo_existe(SITIO, "/sites/Tech/Docs/a.pdf") is False,
              "404 en el archivo = no existe")

    sp = con_api_estado(200, "{}")
    comprobar(sp.carpeta_existe(SITIO, "/sites/Tech/Docs") is True,
              "200 en la carpeta = existe")
    comprobar(sp.archivo_existe(SITIO, "/sites/Tech/Docs/a.pdf") is True,
              "200 en el archivo = existe")

    sp = con_api_estado(403, "Acceso denegado")
    try:
        sp.carpeta_existe(SITIO, "/sites/Tech/Docs")
        comprobar(False, "un 403 no se toma por 'no existe' (se propaga)")
    except ErrorSharePoint as e:
        comprobar("403" in str(e), "un 403 no se toma por 'no existe' (se propaga)")

    sp = con_api_estado(404)
    try:
        sp.listar_archivos("https://contoso.sharepoint.com/sites/Tech/NoExiste")
        comprobar(False, "listar_archivos avisa cuando la carpeta no existe")
    except ErrorSharePoint as e:
        comprobar("no existe" in str(e),
                  "listar_archivos avisa cuando la carpeta no existe")
    try:
        sp.listar_carpetas("https://contoso.sharepoint.com/sites/Tech/NoExiste")
        comprobar(False, "listar_carpetas avisa cuando la carpeta no existe")
    except ErrorSharePoint as e:
        comprobar("no existe" in str(e),
                  "listar_carpetas avisa cuando la carpeta no existe")


def pruebas_clasificar_y_mover():
    print("\n== Clasificar y mover ==")
    sp = SharePoint()
    creadas, movidos = [], []
    sp.listar_carpetas = lambda url: [{"nombre": "Otros",
                                       "ruta_rel": "/sites/Tech/Destino/Otros"}]
    sp.crear_carpeta = lambda sitio, rel: creadas.append(rel)
    sp.carpeta_existe = lambda sitio, rel: True
    sp.archivo_existe = lambda sitio, rel: False
    sp.mover = lambda sitio, origen, destino, nombre, tamano=None: movidos.append(
        (origen, destino, nombre))

    resumen = sp.clasificar_y_mover(
        "https://contoso.sharepoint.com/sites/Tech/Origen",
        "https://contoso.sharepoint.com/sites/Tech/Destino",
        [("a.pdf", "Estimaciones"), ("b.docx", "Otros"), ("c.txt", "")])

    comprobar(creadas == ["/sites/Tech/Destino/Estimaciones"],
              "crea solo la subcarpeta que falta")
    comprobar(len(movidos) == 3, "mueve los tres archivos")
    comprobar(movidos[0] == ("/sites/Tech/Origen/a.pdf",
                             "/sites/Tech/Destino/Estimaciones", "a.pdf"),
              "origen y destino correctos para el primer archivo")
    comprobar(movidos[2][1] == "/sites/Tech/Destino",
              "sin subcarpeta deja el archivo en la raíz del destino")
    comprobar("Movidos 3 de 3" in resumen, "resume los movimientos")
    comprobar("(raíz)" in resumen, "marca los que van a la raíz")

    sp2 = SharePoint()
    sp2.crear_carpeta = lambda sitio, rel: None
    sp2.listar_carpetas = lambda url: []
    sp2.carpeta_existe = lambda sitio, rel: True
    sp2.archivo_existe = lambda sitio, rel: False

    def mover_falla(sitio, origen, destino, nombre, tamano=None):
        raise ErrorSharePoint("no me dejan")

    sp2.mover = mover_falla
    resumen = sp2.clasificar_y_mover(
        "https://contoso.sharepoint.com/sites/Tech/Origen",
        "https://contoso.sharepoint.com/sites/Tech/Destino", [("a.pdf", "X")])
    comprobar("Movidos 0 de 1" in resumen and "no me dejan" in resumen,
              "los fallos no abortan el resto y se listan")

    sp3 = SharePoint()
    movidos3 = []
    sp3.listar_carpetas = lambda url: []
    sp3.crear_carpeta = lambda sitio, rel: None
    sp3.carpeta_existe = lambda sitio, rel: True
    # SharePoint no distingue mayúsculas: la comprobación la hace él, así que
    # 'a.pdf' existe aunque el destino guarde 'A.PDF'.
    sp3.archivo_existe = lambda sitio, rel: rel.lower().endswith("a.pdf")
    sp3.mover = lambda sitio, origen, destino, nombre, tamano=None: \
        movidos3.append(nombre)
    resumen = sp3.clasificar_y_mover(
        "https://contoso.sharepoint.com/sites/Tech/Origen",
        "https://contoso.sharepoint.com/sites/Tech/Destino",
        [("A.PDF", "X"), ("b.pdf", "X")])
    comprobar(movidos3 == ["b.pdf"], "no pisa el archivo que ya existía en destino")
    comprobar("Movidos 1 de 2" in resumen and "Ya existían" in resumen,
              "el resumen avisa de los que ya existían")

    sp4 = SharePoint()
    movidos4 = []
    sp4.listar_carpetas = lambda url: []
    sp4.crear_carpeta = lambda sitio, rel: None
    sp4.archivo_existe = lambda sitio, rel: False
    sp4.mover = lambda sitio, origen, destino, nombre, tamano=None: \
        movidos4.append(nombre)

    def carpeta_sin_permiso(sitio, rel):
        if rel.endswith("/X"):
            error = ErrorSharePoint("SharePoint respondió 403 (Acceso denegado)")
            error.estado = 403
            raise error
        return True

    sp4.carpeta_existe = carpeta_sin_permiso
    resumen = sp4.clasificar_y_mover(
        "https://contoso.sharepoint.com/sites/Tech/Origen",
        "https://contoso.sharepoint.com/sites/Tech/Destino", [("a.pdf", "X")])
    comprobar(movidos4 == [],
              "si no puede comprobar el destino no mueve (así no pisa nada)")
    comprobar("no pude comprobar el destino" in resumen.lower() and "403" in resumen,
              "el resumen explica por qué no lo movió")

    sp5 = SharePoint()
    creados5 = []
    sp5.listar_carpetas = lambda url: []
    sp5.crear_carpeta = lambda sitio, rel: creados5.append(rel)
    sp5.carpeta_existe = lambda sitio, rel: rel != "/sites/Tech/Nuevo"
    sp5.archivo_existe = lambda sitio, rel: False
    sp5.mover = lambda sitio, origen, destino, nombre, tamano=None: None
    sp5.clasificar_y_mover(
        "https://contoso.sharepoint.com/sites/Tech/Origen",
        "https://contoso.sharepoint.com/sites/Tech/Nuevo", [("a.pdf", "")])
    comprobar(creados5 == ["/sites/Tech/Nuevo"],
              "crea la carpeta destino si todavía no existe")

    sp6 = SharePoint()
    sp6.listar_carpetas = lambda url: []
    sp6.crear_carpeta = lambda sitio, rel: None
    sp6.archivo_existe = lambda sitio, rel: False

    def carpeta_destino_fallida(sitio, rel):
        error = ErrorSharePoint("SharePoint respondió 403 (Acceso denegado)")
        error.estado = 403
        raise error

    sp6.carpeta_existe = carpeta_destino_fallida
    resumen = sp6.clasificar_y_mover(
        "https://contoso.sharepoint.com/sites/Tech/Origen",
        "https://contoso.sharepoint.com/sites/Tech/Otro", [("a.pdf", "")])
    comprobar("No moví nada" in resumen and "403" in resumen,
              "si el destino entero no se puede ni mirar, no intenta nada")

    # Entre sitios distintos no se mueve ni se falla: se copia. Eso se prueba
    # con el resto de la copia en 'pruebas_binario_y_copiar'.


def pruebas_binario_y_copiar():
    print("\n== Bajar, subir y copiar archivos ==")
    sp, pestana = con_pestana({"estado": 200, "cuerpo": "",
                               "base64": base64.b64encode(b"hola").decode()})
    comprobar(sp.descargar(SITIO, "/sites/Tech/Docs/a.pdf") == b"hola",
              "descargar devuelve los bytes del archivo")
    js = pestana.visto[-1]
    comprobar("$value" in js and "btoa" in js,
              "la descarga pide el contenido ($value) y lo pasa en base64")
    comprobar("GetFileByServerRelativeUrl('/sites/Tech/Docs/a.pdf')" in js,
              "la descarga apunta al archivo pedido")

    sp, _ = con_pestana({"estado": 404, "cuerpo": "no existe"})
    try:
        sp.descargar(SITIO, "/sites/Tech/Docs/nada.pdf")
        comprobar(False, "si la descarga falla, lo dice")
    except ErrorSharePoint as e:
        comprobar("404" in str(e), "si la descarga falla, lo dice")

    sp, _ = con_pestana({"estado": 200, "cuerpo": "",
                         "base64": base64.b64encode(b"x" * 20).decode()})
    try:
        sp.descargar(SITIO, "/sites/Tech/Docs/gordo.zip", maximo=10)
        comprobar(False, "avisa si el archivo no cabe por la pestaña")
    except ErrorSharePoint as e:
        comprobar("MB" in str(e), "avisa si el archivo no cabe por la pestaña")

    sp, pestana = con_pestana({"estado": 200, "cuerpo": "{}"})
    sp.subir(SITIO, "/sites/Tech/Docs", "acta.docx", b"contenido")
    js = pestana.visto[-1]
    comprobar("Files/add" in js and "overwrite=true" in js,
              "subir usa Files/add y sobrescribe")
    comprobar("contextinfo" in js and "X-RequestDigest" in js,
              "la subida pide y usa el token de escritura")
    comprobar(base64.b64encode(b"contenido").decode() in js,
              "la subida manda el contenido del archivo en el cuerpo")

    sp, pestana = con_pestana([{"estado": 400, "cuerpo": "nombre raro"},
                               {"estado": 200, "cuerpo": "{}"}])
    sp.subir(SITIO, "/sites/Tech/Docs", "informe #3.pdf", b"pdf")
    comprobar(len(pestana.visto) == 2 and "AddUsingPath" in pestana.visto[-1],
              "si Files/add falla con el nombre, reintenta con AddUsingPath")
    comprobar("informe #3.pdf" in pestana.visto[-1],
              "AddUsingPath usa el nombre sin codificar")

    sp, _ = con_pestana({"estado": 403, "cuerpo": "sin permiso"})
    try:
        sp.subir(SITIO, "/sites/Tech/Docs", "a.pdf", b"x")
        comprobar(False, "si no puede subir de ninguna forma, lo dice")
    except ErrorSharePoint as e:
        comprobar("403" in str(e), "si no puede subir de ninguna forma, lo dice")

    # Copiar dentro del MISMO sitio: lo hace SharePoint, sin bajarlo.
    sp, pestana = con_pestana({"estado": 200, "cuerpo": "{}"})
    sp.copiar(SITIO, "/sites/Tech/Origen/a.pdf", SITIO,
              "/sites/Tech/Destino", "a.pdf")
    js = pestana.visto[-1]
    comprobar("CopyTo" in js and "bOverwrite=true" in js,
              "copiar en el mismo sitio lo hace SharePoint (CopyTo)")
    comprobar("Origen/a.pdf" in js and "Destino/a.pdf" in js,
              "CopyTo apunta al archivo de origen y al destino")

    # Copiar a OTRO sitio (tu OneDrive): hay que bajarlo y subirlo.
    sp = SharePoint()
    bajados, subidos = [], []

    def descargar_falso(sitio, rel, maximo=None):
        bajados.append((sitio, rel))
        return b"bytes del archivo"

    sp.descargar = descargar_falso
    sp._tamano_o_nada = lambda sitio, rel: None  # sin preguntar a SharePoint
    sp.subir = lambda sitio, carpeta, nombre, datos, sobrescribir=True: \
        subidos.append((sitio, carpeta, nombre, datos))
    sp.copiar("https://a.sharepoint.com/sites/X", "/sites/X/Docs/a.pdf",
              "https://b-my.sharepoint.com/personal/u_b_com",
              "/personal/u_b_com/Documents", "a.pdf")
    comprobar(bajados == [("https://a.sharepoint.com/sites/X",
                           "/sites/X/Docs/a.pdf")],
              "copiar entre sitios baja el archivo del origen")
    comprobar(subidos == [("https://b-my.sharepoint.com/personal/u_b_com",
                           "/personal/u_b_com/Documents", "a.pdf",
                           b"bytes del archivo")],
              "copiar entre sitios sube el archivo al destino")

    # Y el clasificador, entre sitios distintos, COPIA (no mueve ni falla).
    sp = SharePoint()
    copiados, movidos = [], []
    sp.listar_carpetas = lambda url: []
    sp.crear_carpeta = lambda sitio, rel: None
    sp.carpeta_existe = lambda sitio, rel: True
    sp.archivo_existe = lambda sitio, rel: False
    sp.mover = lambda *a, **k: movidos.append(a)
    sp.copiar = lambda *a, **k: copiados.append(a)
    resumen = sp.clasificar_y_mover(
        "https://a.sharepoint.com/sites/X/Docs",
        "https://b-my.sharepoint.com/personal/u_b_com/Documents",
        [("a.pdf", "Facturas")])
    comprobar(movidos == [] and len(copiados) == 1,
              "entre sitios clasifica copiando, no moviendo")
    comprobar(copiados[0] == ("https://a.sharepoint.com/sites/X",
                              "/sites/X/Docs/a.pdf",
                              "https://b-my.sharepoint.com/personal/u_b_com",
                              "/personal/u_b_com/Documents/Facturas", "a.pdf"),
              "copia del sitio de origen a la subcarpeta del destino")
    comprobar("Copiados 1 de 1" in resumen and "no los toco" in resumen,
              "el resumen dice que copió y que los originales siguen ahí")

    # En el mismo sitio sigue moviendo, y con copiar=True copia (CopyTo).
    sp2 = SharePoint()
    copiados2 = []
    sp2.listar_carpetas = lambda url: []
    sp2.crear_carpeta = lambda sitio, rel: None
    sp2.carpeta_existe = lambda sitio, rel: True
    sp2.archivo_existe = lambda sitio, rel: False
    sp2.mover = lambda *a, **k: None
    sp2.copiar = lambda *a, **k: copiados2.append(a)
    resumen = sp2.clasificar_y_mover(
        "https://a.sharepoint.com/sites/X/Docs",
        "https://a.sharepoint.com/sites/X/Destino", [("a.pdf", "")],
        copiar=True)
    comprobar(len(copiados2) == 1 and "Copiados 1 de 1" in resumen,
              "copiar=True copia incluso dentro del mismo sitio")
    comprobar("no los toco" not in resumen,
              "en el mismo sitio no habla de originales que no se tocan")

    sp, pestana = con_pestana({"estado": 200, "cuerpo": "{}"})
    sp.borrar_archivo(SITIO, "/sites/Tech/Docs/copia.pdf")
    js = pestana.visto[-1]
    comprobar("X-HTTP-Method" in js and "DELETE" in js and "IF-MATCH" in js,
              "borrar usa POST con X-HTTP-Method: DELETE de SharePoint")


def pruebas_archivos_grandes():
    """Los archivos grandes (más de 10 MB) se bajan a tu PC y se suben por trozos.

    SharePoint deja a medias los movimientos de unos pocos MB y por la pestaña
    de Brave no cabe un archivo entero: a partir de ``TAMANO_POR_DISCO`` se baja
    a disco por rangos, se sube con la subida por trozos y, si era un
    movimiento, el original se borra SOLO si lo que llegó pesa lo mismo.
    """
    print("\n== Archivos grandes: bajar a tu PC y subir por trozos ==")
    umbral = SharePoint.TAMANO_POR_DISCO
    gordo = umbral + 1

    # Cuánto pesa: se le pregunta a SharePoint pidiendo solo el tamaño.
    sp, pestana = con_pestana({"estado": 200,
                               "cuerpo": json.dumps({"d": {"Length": "15728640"}})})
    comprobar(sp.tamano_de(SITIO, "/sites/Tech/Docs/gordo.zip") == 15728640,
              "pregunta a SharePoint cuánto pesa un archivo")
    comprobar("$select=Length" in pestana.visto[-1],
              "y lo pregunta pidiendo solo el tamaño")
    sp, _ = con_pestana({"estado": 200, "cuerpo": "{}"})
    comprobar(sp.tamano_de(SITIO, "/x/a.pdf") is None,
              "si SharePoint no lo dice, el tamaño queda en nada (no se inventa)")
    sp, _ = con_pestana({"estado": 403, "cuerpo": "sin permiso"})
    comprobar(sp._tamano_o_nada(SITIO, "/x/a.pdf") is None,
              "si no se puede preguntar, se sigue sin romperse")

    # Mover un archivo grande: nada de MoveTo, se hace por disco.
    sp, pestana = con_pestana({"estado": 200, "cuerpo": "{}"})
    pasados = []
    sp._pasar_por_disco = lambda *a, **k: pasados.append((a, k))
    sp.mover(SITIO, "/sites/Tech/Origen/gordo.zip", "/sites/Tech/Destino",
             "gordo.zip", tamano=gordo)
    comprobar(len(pasados) == 1 and pasados[0][1].get("borrar_origen") is True,
              "un archivo grande se mueve por disco y borra el original al final")
    comprobar(pestana.visto == [],
              "y no se intenta MoveTo, que es el que se queda a medias")

    # Mover uno pequeño: como siempre (MoveTo y sin pasar por disco).
    sp, pestana = con_pestana({"estado": 200, "cuerpo": "{}"})

    def no_deberia(*a, **k):
        comprobar(False, "un archivo pequeño no debería pasar por disco")

    sp._pasar_por_disco = no_deberia
    sp.mover(SITIO, "/sites/Tech/Origen/acta.pdf", "/sites/Tech/Destino",
             "acta.pdf", tamano=1024)
    comprobar("MoveTo" in pestana.visto[-1],
              "uno pequeño se mueve con MoveTo, como siempre")

    # Si SharePoint se atraganta con uno que no es tan grande, se reintenta
    # por disco: es justo lo que pasa con los de unos pocos MB.
    sp, _ = con_pestana({"estado": 500, "cuerpo": "no puedo mover"})
    pasados = []
    sp._pasar_por_disco = lambda *a, **k: pasados.append((a, k))
    sp.mover(SITIO, "/sites/Tech/Origen/medio.zip", "/sites/Tech/Destino",
             "medio.zip", tamano=5 * 1024 * 1024)
    comprobar(len(pasados) == 1 and pasados[0][1].get("borrar_origen") is True,
              "si el MoveTo falla, se reintenta por disco (y se borra el original)")

    # Un 403 (sin permiso) no se reintenta por disco: bajar tampoco va a dejar.
    sp, _ = con_pestana({"estado": 403, "cuerpo": "sin permiso"})
    pasados = []
    sp._pasar_por_disco = lambda *a, **k: pasados.append((a, k))
    try:
        sp.mover(SITIO, "/sites/Tech/Origen/a.pdf", "/sites/Tech/Destino",
                 "a.pdf", tamano=1024)
        comprobar(False, "un 403 al mover se devuelve tal cual")
    except ErrorSharePoint as e:
        comprobar("403" in str(e) and pasados == [],
                  "un 403 al mover se devuelve tal cual (y no se insiste por disco)")

    # Bajar por trozos: se piden rangos y se escribe el archivo en tu PC.
    sp = SharePoint()
    sp.TAMANO_TROZO = 4
    partes = [b"abcd", b"efgh", b"ij"]
    pedidos = []

    def rango_falso(sitio, ruta, desde, hasta):
        pedidos.append((desde, hasta))
        return (partes[len(pedidos) - 1] if len(pedidos) <= len(partes)
                else b""), 10

    sp._api_rango = rango_falso
    destino = os.path.join(tempfile.gettempdir(), "asistente_gordo.zip")
    try:
        bajado = sp.descargar_a_disco(SITIO, "/sites/Tech/Docs/gordo.zip",
                                      destino, tamano=10)
        with open(destino, "rb") as archivo:
            escrito = archivo.read()
        comprobar(bajado == 10 and escrito == b"abcdefghij",
                  "el archivo grande se baja a tu PC completo")
        comprobar(pedidos == [(0, 3), (4, 7), (8, 11)],
                  "pidiéndolo por trozos (bytes=0-3, 4-7, 8-11)")
    finally:
        if os.path.exists(destino):
            os.remove(destino)

    # Si lo que llega no pesa lo que dice SharePoint, no se pasa nada.
    sp = SharePoint()
    sp.TAMANO_TROZO = 4
    respuestas = [(b"ab", 10), (b"", None)]
    sp._api_rango = lambda *a: respuestas.pop(0)
    destino = os.path.join(tempfile.gettempdir(), "asistente_cortado.zip")
    try:
        sp.descargar_a_disco(SITIO, "/sites/Tech/Docs/gordo.zip", destino,
                             tamano=10)
        comprobar(False, "una bajada a medias no se da por buena")
    except ErrorSharePoint as e:
        comprobar("no llegó completo" in str(e),
                  "si la bajada llega a medias, lo dice (y no se sube nada)")
    finally:
        if os.path.exists(destino):
            os.remove(destino)

    # Subir desde disco: los grandes van por trozos de SharePoint.
    sp, pestana = con_pestana({"estado": 204, "cuerpo": ""})
    sp.TAMANO_TROZO = 4
    origen = os.path.join(tempfile.gettempdir(), "asistente_trozos.bin")
    contenido = b"0123456789"  # 10 bytes: 4 + 4 + 2
    with open(origen, "wb") as archivo:
        archivo.write(contenido)
    try:
        subido = sp.subir_desde_disco(SITIO, "/sites/Tech/Docs", "gordo.bin",
                                      origen)
        js = " ".join(pestana.visto)
        comprobar(subido == 10, "sube el archivo entero desde tu PC")
        comprobar("StartUpload(uploadId=guid'" in js,
                  "empieza preparando el archivo (StartUpload)")
        comprobar("ContinueUpload(uploadId=guid'" in js and "fileOffset=4" in js,
                  "manda los trozos con ContinueUpload y su desplazamiento")
        comprobar("FinishUpload(uploadId=guid'" in js and "fileOffset=8" in js,
                  "y cierra el archivo con FinishUpload")
        mandado = b"".join(
            base64.b64decode(t) for t in
            re.findall(r'atob\("([A-Za-z0-9+/=]*)"\)', js))
        comprobar(mandado == contenido,
                  "los bytes que viajan son los del archivo, sin perder nada")
    finally:
        if os.path.exists(origen):
            os.remove(origen)

    # Un archivo pequeño se sube de una vez (el camino de siempre).
    sp, pestana = con_pestana({"estado": 200, "cuerpo": "{}"})
    origen = os.path.join(tempfile.gettempdir(), "asistente_pequeno.bin")
    with open(origen, "wb") as archivo:
        archivo.write(b"pequeno")
    try:
        sp.subir_desde_disco(SITIO, "/sites/Tech/Docs", "pequeno.bin", origen)
        comprobar("Files/add" in pestana.visto[-1]
                  and "StartUpload" not in pestana.visto[-1],
                  "uno pequeño se sube de una vez, sin trozos")
    finally:
        if os.path.exists(origen):
            os.remove(origen)


    # El camino completo: baja a tu PC, sube al destino y, ya comprobado que
    # pesa lo mismo, borra el original.
    sp = SharePoint()
    local = os.path.join(tempfile.gettempdir(), "asistente_pasado.zip")
    sp._ruta_temporal = lambda nombre: local
    hechos = {"bajados": [], "subidos": [], "borrados": []}

    def bajar_falso(sitio, rel, ruta, tamano=None):
        hechos["bajados"].append((sitio, rel))
        with open(ruta, "wb") as archivo:
            archivo.write(b"x" * 10)
        return 10

    def subir_falso(sitio, carpeta, nombre, ruta, sobrescribir=True):
        hechos["subidos"].append((sitio, carpeta, nombre))
        return 10

    sp.descargar_a_disco = bajar_falso
    sp.subir_desde_disco = subir_falso
    sp.tamano_de = lambda sitio, rel: 10  # lo que quedó en el destino
    sp.borrar_archivo = lambda sitio, rel: hechos["borrados"].append((sitio, rel))
    sp._pasar_por_disco(SITIO, "/sites/Tech/Origen/gordo.zip", SITIO,
                        "/sites/Tech/Destino", "gordo.zip", tamano=10,
                        borrar_origen=True)
    comprobar(hechos["bajados"] == [(SITIO, "/sites/Tech/Origen/gordo.zip")],
              "baja el archivo del origen a tu PC")
    comprobar(hechos["subidos"] == [(SITIO, "/sites/Tech/Destino", "gordo.zip")],
              "lo sube a la carpeta de destino")
    comprobar(hechos["borrados"] == [(SITIO, "/sites/Tech/Origen/gordo.zip")],
              "y borra el original (queda en la papelera del sitio)")
    comprobar(not os.path.exists(local),
              "el archivo de tu PC se limpia cuando todo va bien")

    # Si en el destino no pesa lo mismo, no se borra NADA y se dice dónde quedó.
    sp = SharePoint()
    local = os.path.join(tempfile.gettempdir(), "asistente_dudoso.zip")
    sp._ruta_temporal = lambda nombre: local
    sp.descargar_a_disco = bajar_falso
    sp.subir_desde_disco = subir_falso
    sp.tamano_de = lambda sitio, rel: 5
    borrados = []
    sp.borrar_archivo = lambda sitio, rel: borrados.append(rel)
    try:
        sp._pasar_por_disco(SITIO, "/sites/Tech/Origen/gordo.zip", SITIO,
                            "/sites/Tech/Destino", "gordo.zip", tamano=10,
                            borrar_origen=True)
        comprobar(False, "si lo que llegó no pesa lo mismo, no se sigue")
    except ErrorSharePoint as e:
        comprobar("no me fío" in str(e),
                  "si lo que llegó no pesa lo mismo, lo dice y no toca nada")
        comprobar("está en tu PC" in str(e),
                  "y avisa de dónde quedó el archivo bajado")
    comprobar(borrados == [] and os.path.exists(local),
              "no borra el original y deja el archivo en tu PC (antes duplicado "
              "que perdido)")
    if os.path.exists(local):
        os.remove(local)

    # Si no se puede comprobar el peso del destino, tampoco se borra el original.
    sp = SharePoint()
    local = os.path.join(tempfile.gettempdir(), "asistente_sinpeso.zip")
    sp._ruta_temporal = lambda nombre: local
    sp.descargar_a_disco = bajar_falso
    sp.subir_desde_disco = subir_falso
    sp.tamano_de = lambda sitio, rel: None
    avisos = []
    sp.on_aviso = avisos.append
    borrados = []
    sp.borrar_archivo = lambda sitio, rel: borrados.append(rel)
    sp._pasar_por_disco(SITIO, "/sites/Tech/Origen/gordo.zip", SITIO,
                        "/sites/Tech/Destino", "gordo.zip", tamano=10,
                        borrar_origen=True)
    comprobar(borrados == [] and any("dejo el original" in a for a in avisos),
              "si no puede comprobar cuánto pesa en el destino, deja el original")
    if os.path.exists(local):
        os.remove(local)

    # Copiar un archivo grande: por disco también, pero sin borrar el original.
    sp = SharePoint()
    pasados = []
    sp._pasar_por_disco = lambda *a, **k: pasados.append((a, k))
    sp.copiar("https://a.sharepoint.com/sites/X", "/sites/X/Docs/gordo.zip",
              "https://b-my.sharepoint.com/personal/u_b_com",
              "/personal/u_b_com/Documents", "gordo.zip", tamano=gordo)
    comprobar(len(pasados) == 1 and not pasados[0][1].get("borrar_origen"),
              "copiar un archivo grande va por disco y NO borra el original")

    # Y si el archivo no cabe de una vez (más de TAMANO_MAXIMO), se pasa por
    # disco en vez de rendirse: es lo que hacía antes con los de 30 MB.
    sp = SharePoint()
    pasados = []

    def no_cabe(sitio, rel, maximo=None):
        error = ErrorSharePoint("no cabe en un mensaje de la pestaña")
        error.pesado = True
        raise error

    sp.descargar = no_cabe
    sp._tamano_o_nada = lambda sitio, rel: None
    sp._pasar_por_disco = lambda *a, **k: pasados.append((a, k))
    sp.copiar("https://a.sharepoint.com/sites/X", "/sites/X/Docs/g.zip",
              "https://b-my.sharepoint.com/personal/u_b_com",
              "/personal/u_b_com/Documents", "g.zip")
    comprobar(len(pasados) == 1,
              "si no cabe de una vez, se copia igual pasándolo por disco")

    # El clasificador ya sabe cuánto pesa cada archivo (viene en el listado) y
    # se lo pasa a mover/copiar para no preguntarlo otra vez, avisando antes de
    # cuántos van por disco (que es más lento).
    sp = SharePoint()
    avisos = []
    sp.on_aviso = avisos.append
    sp.listar_carpetas = lambda url: []
    sp.crear_carpeta = lambda sitio, rel: None
    sp.carpeta_existe = lambda sitio, rel: True
    sp.archivo_existe = lambda sitio, rel: False
    tamanos_vistos = []
    sp.mover = lambda sitio, origen, destino, nombre, tamano=None: \
        tamanos_vistos.append((nombre, tamano))
    resumen = sp.clasificar_y_mover(
        "https://contoso.sharepoint.com/sites/Tech/Origen",
        "https://contoso.sharepoint.com/sites/Tech/Destino",
        [("AXA/acta.pdf", "Actas"), ("AXA/gordo.zip", "Actas")],
        tamanos={"AXA/acta.pdf": 100, "AXA/gordo.zip": gordo})
    comprobar(tamanos_vistos == [("acta.pdf", 100), ("gordo.zip", gordo)],
              "el tamaño del listado llega a mover (sin preguntarlo otra vez)")
    comprobar("Movidos 2 de 2" in resumen, "y el trasiego sigue igual")
    comprobar(any("1 de esos archivo(s) pasan de 10 MB" in a for a in avisos),
              "avisa de cuántos van por disco antes de empezar")


def pruebas_resolver_url():
    print("\n== Resolver la URL que pega el usuario ==")
    sp = SharePoint()
    llamadas = []

    def api_falsa(sitio, ruta, metodo="GET", datos=None, cabeceras=None):
        llamadas.append(sitio)
        if "PeopleManager" in ruta:
            return {"d": {"PersonalUrl":
                          "https://contoso-my.sharepoint.com/personal/"
                          "john_doe_contoso_com/"}}
        if "/personal/john_doe_contoso_com" in sitio:
            return {"d": {"results": [{"RootFolder": {"ServerRelativeUrl":
                          "/personal/john_doe_contoso_com/Documents"}}]}}
        if "/sites/Tech" in sitio:
            return {"d": {"results": [{"RootFolder": {"ServerRelativeUrl":
                          "/sites/Tech/Documentos compartidos"}}]}}
        return {"d": {"results": []}}

    sp._api = api_falsa
    comprobar(sp.resolver("https://contoso.sharepoint.com/sites/Tech")
              == "https://contoso.sharepoint.com/sites/Tech/"
                 "Documentos compartidos",
              "la raíz de un sitio pasa a su biblioteca (el nombre cambia "
              "con el idioma)")
    comprobar(sp.resolver("https://contoso-my.sharepoint.com/")
              == "https://contoso-my.sharepoint.com/personal/"
                 "john_doe_contoso_com/Documents",
              "la URL de OneDrive pasa a tu unidad (/personal/<usuario>)")

    antes = len(llamadas)
    url_carpeta = ("https://contoso-my.sharepoint.com/personal/"
                   "john_doe_contoso_com/Documents/Invoices")
    comprobar(sp.resolver(url_carpeta) == url_carpeta,
              "una URL que ya es carpeta se deja tal cual")
    comprobar(len(llamadas) == antes,
              "y resolverla no gasta ninguna llamada de más")

    sp2 = SharePoint()

    def api_sin_permiso(*a, **k):
        raise ErrorSharePoint("SharePoint respondió 403 (Acceso denegado)")

    sp2._api = api_sin_permiso
    comprobar(sp2.resolver("https://contoso.sharepoint.com/sites/Tech")
              == "https://contoso.sharepoint.com/sites/Tech/Documents",
              "si no puede preguntar por la biblioteca, usa 'Documents'")

    sp3 = SharePoint()

    def api_sin_brave(*a, **k):
        raise navegador_cdp.ErrorNavegador("no encuentro Brave")

    sp3._api = api_sin_brave
    try:
        sp3.resolver("https://contoso-my.sharepoint.com/")
        comprobar(False, "sin saber de quién es la unidad, lo dice en vez de "
                         "inventarse una carpeta")
    except ErrorSharePoint as e:
        comprobar("Mi unidad" in str(e) and "Documents" in str(e),
                  "sin saber de quién es la unidad, lo dice en vez de "
                  "inventarse una carpeta")


def pruebas_pestana_muerta():
    """Una pestaña que Brave ya cerró no puede dejar el módulo inservible.

    Brave contesta 500 en el *handshake* cuando el id de la pestaña ya no existe
    (la cerró el usuario, se reinició Brave, el ahorro de memoria la descartó):
    eso salía como ``Invalid response status`` y la operación se daba por
    perdida. Aquí se prueba que el módulo lo reconoce y sigue adelante.
    """
    print("\n== Pestañas muertas y reintentos ==")
    comprobar(issubclass(navegador_cdp.ErrorPestana,
                         navegador_cdp.ErrorNavegador),
              "ErrorPestana sigue siendo un ErrorNavegador (lo que ya lo "
              "capturaba sigue funcionando)")

    # 1. La caché reutiliza una pestaña viva...
    sp = SharePoint()
    viva = PestanaFalsa({"estado": 200, "cuerpo": "{}"})
    sp._pestanas[SITIO] = viva
    comprobar(sp._pestana(SITIO) is viva,
              "se reutiliza la pestaña que sigue viva")

    # 2. ...y cambia por otra la que Brave ya cerró.
    abiertas = []

    class PestanaFalsaNueva:
        """Como navegador_cdp.Pestana, pero sin abrir Brave de verdad."""

        def __init__(self, url=None, timeout=40):
            self.url = url
            abiertas.append(url)

        def abrir(self):
            return self

        def sigue_viva(self):
            return True

        def evaluar(self, expresion, navegar=None, intentos=3):
            return {"estado": 200, "cuerpo": "{}"}

        def cerrar(self):
            pass

    original_pestana = navegador_cdp.Pestana
    navegador_cdp.Pestana = PestanaFalsaNueva
    try:
        sp._pestanas[SITIO] = PestanaFalsa({"estado": 200}, viva=False)
        nueva = sp._pestana(SITIO)
        comprobar(abiertas == [SITIO] and isinstance(nueva, PestanaFalsaNueva),
                  "si Brave ya la cerró, se abre otra pestaña del mismo sitio")
    finally:
        navegador_cdp.Pestana = original_pestana

    # 3. Reiniciar Brave invalida todas las pestañas apuntadas.
    original_asegurar = navegador_cdp.asegurar
    try:
        sp = SharePoint()
        sp._pestanas[SITIO] = PestanaFalsa({"estado": 200})
        navegador_cdp.asegurar = lambda url="", reiniciar=True: "Reinicié Brave"
        sp.asegurar_navegador(SITIO)
        comprobar(sp._pestanas == {},
                  "al reiniciar Brave se olvidan las pestañas viejas")

        sp._pestanas[SITIO] = PestanaFalsa({"estado": 200})
        navegador_cdp.asegurar = lambda url="", reiniciar=True: (
            navegador_cdp.NOTA_YA_ABIERTO)
        sp.asegurar_navegador(SITIO)
        comprobar(SITIO in sp._pestanas,
                  "si Brave ya estaba abierto, se conservan las pestañas")
    finally:
        navegador_cdp.asegurar = original_asegurar

    # 4. El mensaje del rechazo cuenta qué pasó de verdad.
    comprobar("ya no existe" in navegador_cdp._rechazo(
        500, "ws://127.0.0.1:9222/devtools/page/PESTANA-QUE-YA-NO-EXISTE"),
              "un 500 se explica como 'esa pestaña ya no existe'")
    comprobar("--remote-allow-origins" in navegador_cdp._rechazo(
        403, "ws://127.0.0.1:9222/devtools/page/x"),
              "un 403 habla del origen que Brave no acepta")
    comprobar("cerrado" in navegador_cdp._rechazo(0, "ws://127.0.0.1:9222/x"),
              "sin respuesta se sospecha que Brave está cerrado")

    # 5. Con la pestaña muerta, `evaluar` insiste con una nueva.
    conexiones = []

    async def evaluar_falso(ws_url, expresion, navegar, timeout):
        conexiones.append(ws_url)
        if len(conexiones) == 1:
            raise navegador_cdp.ErrorPestana("esa pestaña ya no existe")
        return 7

    reabiertas = []

    class PestanaQueRenace(navegador_cdp.Pestana):
        """Pestaña de verdad cuyo `abrir` no toca Brave (solo lo apunta)."""

        def abrir(self):
            reabiertas.append(self.url)
            self.id = "renacida"
            self.ws_url = "ws://127.0.0.1:9222/devtools/page/renacida"
            return self

        def sigue_viva(self):
            return False

    original_evaluar = navegador_cdp._evaluar_async
    navegador_cdp._evaluar_async = evaluar_falso
    try:
        pestana = PestanaQueRenace(SITIO)
        pestana.id = "muerta"
        pestana.ws_url = "ws://127.0.0.1:9222/devtools/page/muerta"
        comprobar(pestana.evaluar("2 + 5") == 7,
                  "tras el rechazo se reintenta y la evaluación acaba bien")
        comprobar(reabiertas == [SITIO],
                  "el reintento abre una pestaña nueva en la misma URL")
        comprobar(conexiones[-1].endswith("/renacida"),
                  "y se conecta a la pestaña nueva, no a la muerta")
    finally:
        navegador_cdp._evaluar_async = original_evaluar

    # 6. También cuando lo que falla es *cuándo* nos conectamos (Brave ocupado).
    conexiones.clear()
    reabiertas.clear()

    async def evaluar_falla_dos_veces(ws_url, expresion, navegar, timeout):
        conexiones.append(ws_url)
        if len(conexiones) <= 2:
            raise navegador_cdp.ErrorPestana("Brave respondió 500")
        return "vale"

    class PestanaSiempreViva(navegador_cdp.Pestana):
        def abrir(self):
            reabiertas.append(self.url)
            self.id = "otra"
            self.ws_url = "ws://127.0.0.1:9222/devtools/page/otra"
            return self

        def sigue_viva(self):
            return True

    navegador_cdp._evaluar_async = evaluar_falla_dos_veces
    try:
        pestana = PestanaSiempreViva(SITIO)
        pestana.id = "viva"
        pestana.ws_url = "ws://127.0.0.1:9222/devtools/page/viva"
        comprobar(pestana.evaluar("x") == "vale",
                  "si la conexión falla varias veces, también se reintenta")
        comprobar(len(reabiertas) == 1 and conexiones[-1].endswith("/otra"),
                  "a la segunda, se prueba con una pestaña recién abierta")
    finally:
        navegador_cdp._evaluar_async = original_evaluar

    # 7. Y si de verdad no hay forma, se queja con el motivo.
    async def evaluar_siempre_falla(ws_url, expresion, navegar, timeout):
        raise navegador_cdp.ErrorPestana("esa pestaña ya no existe")

    navegador_cdp._evaluar_async = evaluar_siempre_falla
    try:
        pestana = PestanaQueRenace(SITIO)
        try:
            pestana.evaluar("x")
            comprobar(False, "si nunca conecta, lo dice con el motivo")
        except navegador_cdp.ErrorPestana as e:
            comprobar("ya no existe" in str(e),
                      "si nunca conecta, lo dice con el motivo")
    finally:
        navegador_cdp._evaluar_async = original_evaluar


def pruebas_javascript_valido():
    """Si hay Node instalado, comprueba que el JS construido compila.

    El `_api` arma las llamadas concatenando cadenas; un fallo ahí solo se
    vería al hablar con Brave. Con Node lo detectamos sin abrir el navegador.
    """
    print("\n== El JavaScript generado compila ==")
    node = shutil.which("node")
    if not node:
        print("  --  node no está instalado: se omite la comprobación")
        return
    sp, pestana = con_pestana({"estado": 200, "cuerpo": "{}"})
    sp._api(SITIO, "/web/x")
    sp._api(SITIO, "/web/x", metodo="POST", datos={"a": "b'c", "n": 1})
    sp._api(SITIO, "/web/x", metodo="POST", datos={},
            cabeceras={"X-HTTP-Method": "DELETE", "IF-MATCH": "*"})
    sp._api_binario(SITIO, "/web/GetFileByServerRelativeUrl('/x/a.pdf')/$value")
    sp._api_binario(SITIO, "/web/x/Files/add(url='a.docx',overwrite=true)",
                    contenido=b"\x00\x01binario")
    sp._api_rango(SITIO, "/web/GetFileByServerRelativeUrl('/x/a.pdf')/$value",
                  0, 1023)
    # Y las subidas grandes: preparar el archivo, los trozos y el cierre.
    sp.TAMANO_TROZO = 4
    ruta_trozos = os.path.join(tempfile.gettempdir(), "asistente_js_trozos.bin")
    with open(ruta_trozos, "wb") as archivo:
        archivo.write(b"0123456789")
    try:
        sp.subir_desde_disco(SITIO, "/sites/Tech/Docs", "gordo.bin", ruta_trozos)
    finally:
        os.remove(ruta_trozos)
    ruta = os.path.join(tempfile.gettempdir(), "asistente_sharepoint_js.js")
    try:
        with open(ruta, "w", encoding="utf-8") as f:
            for i, js in enumerate(pestana.visto):
                f.write(f"const j{i} = {js};\n")
        proceso = subprocess.run([node, "--check", ruta],
                                 capture_output=True, text=True)
        comprobar(proceso.returncode == 0,
                  "las expresiones JS (GET, POST, binario, rangos y trozos) "
                  "compilan sin errores"
                  + ("" if proceso.returncode == 0
                     else f": {proceso.stderr[:200]}"))
    finally:
        if os.path.exists(ruta):
            os.remove(ruta)


def pruebas_clasificacion_naive():
    print("\n== Clasificación de respaldo ==")
    plan = SharePoint.clasificar_naive(["a.txt", "b.pdf"], "cualquier regla")
    comprobar(plan == [("a.txt", ""), ("b.pdf", "")], "sin IA todo va a la raíz")


def pruebas_solo_documentos_de_oficina():
    """El filtro: de esas carpetas solo salen documentos de oficina.

    En una biblioteca de verdad (una carpeta por cliente, con sus años dentro)
    hay vídeos y demás mezclados con las actas y los informes, y son justo los
    que peor sientan: bajar un vídeo de 1,5 GB para clasificarlo es tirar disco
    y tiempo. Aquí se comprueba lo pedido: que un vídeo (o un PDF, o un zip) ni
    se liste, ni se baje, ni se mueva, y que se diga en qué carpeta está.
    """
    print("\n== Solo documentos de oficina (ppt, pptx, doc y docx) ==")

    def api_falsa(sitio, ruta, metodo="GET", datos=None):
        """Clientes: en la raíz y dentro de AXA, un documento y un vídeo."""
        if ruta.endswith("Files?$select=Name,ServerRelativeUrl,Length"):
            return {"d": {"results": [
                {"Name": "presentacion-2026.pptx", "Length": 5000,
                 "ServerRelativeUrl":
                     "/sites/Tech/Colombia/presentacion-2026.pptx"},
                {"Name": "charla.mp4", "Length": 1500000000,
                 "ServerRelativeUrl": "/sites/Tech/Colombia/charla.mp4"}]}}
        if "'/sites/Tech/Colombia'" in ruta and "Folders?" in ruta:
            return {"d": {"results": [
                {"Name": "AXA",
                 "ServerRelativeUrl": "/sites/Tech/Colombia/AXA",
                 "Files": {"results": [
                     {"Name": "acta.docx", "Length": 120,
                      "ServerRelativeUrl":
                          "/sites/Tech/Colombia/AXA/acta.docx"},
                     {"Name": "reunion-2024.mov", "Length": 900,
                      "ServerRelativeUrl":
                          "/sites/Tech/Colombia/AXA/reunion-2024.mov"}]}}]}}
        return {}

    SharePoint.EXTENSIONES_UTILES = EXTENSIONES_REALES
    try:
        comprobar(SharePoint.EXTENSIONES_UTILES
                  == (".ppt", ".pptx", ".doc", ".docx"),
                  "por defecto solo se pasan ppt, pptx, doc y docx")

        sp = SharePoint()
        comprobar(sp.extension_de("Informe 2024.12.DOCX") == ".docx",
                  "la extensión se saca del último punto y en minúsculas")
        comprobar(sp.extension_de("carpeta/acta") == "",
                  "un archivo sin extensión no tiene ninguna")
        comprobar(sp.extensiones_texto() == "ppt, pptx, doc y docx",
                  "las extensiones se dicen como en el chat")
        comprobar(sp.es_util("presentacion.PPTX") and sp.es_util("acta.doc"),
                  "valen las de oficina con cualquier mayúscula")
        comprobar(not sp.es_util("charla.mp4") and not sp.es_util("foto.jpg")
                  and not sp.es_util("acta.pdf")
                  and not sp.es_util("informe.docx.exe"),
                  "vídeos, fotos, PDF y nombres con truco se quedan fuera")

        # --- Listado: los vídeos no salen, y se dice en qué carpeta están.
        avisos = []
        sp.on_aviso = avisos.append
        sp._api = api_falsa
        archivos = sp.listar_archivos(
            "https://contoso.sharepoint.com/sites/Tech/Colombia", profundidad=2)
        comprobar([a["subruta"] for a in archivos]
                  == ["presentacion-2026.pptx", "AXA/acta.docx"],
                  "solo se listan los documentos de oficina (también dentro de "
                  "las subcarpetas)")
        comprobar(sp.descartados
                  == {"total": 2,
                      "carpetas": ["/sites/Tech/Colombia", "AXA"]},
                  "los que me salté quedan apuntados con la carpeta donde están")
        comprobar("Me salté 2 archivo(s) que no son ppt, pptx, doc y docx"
                  in sp.nota_listado
                  and "/sites/Tech/Colombia" in sp.nota_listado
                  and "AXA" in sp.nota_listado,
                  "y el aviso dice cuántos me salté y en qué carpetas están")
        comprobar(any("Me salté 2 archivo(s)" in a for a in avisos),
                  "el aviso también sale por el chat")

        # --- Bajar, mover y copiar: el filtro es la red de seguridad.
        sp2 = SharePoint()
        llamadas = []

        def binario_falso(sitio, ruta_api, contenido=None, correctos=(200,)):
            llamadas.append(ruta_api)
            return b"x"

        def api_falsa2(sitio, ruta_api, metodo="GET", datos=None,
                       cabeceras=None):
            llamadas.append(ruta_api)
            return {}

        sp2._api_binario = binario_falso
        sp2._api = api_falsa2

        def comprobar_que_protesta(texto, hacer):
            """El filtro para el paso, y dice la carpeta donde está el vídeo."""
            try:
                hacer()
            except ErrorSharePoint as e:
                comprobar("solo ppt, pptx, doc y docx" in str(e)
                          and "AXA" in str(e), texto)
            else:
                comprobar(False, texto)

        comprobar_que_protesta(
            "un vídeo no se baja aunque lo pidan por su nombre (y se dice "
            "dónde está)",
            lambda: sp2.descargar(SITIO,
                                  "/sites/Tech/Colombia/AXA/charla.mp4"))
        en_disco = os.path.join(tempfile.gettempdir(), "asistente_video.mp4")
        comprobar_que_protesta(
            "un vídeo no se baja a tu PC",
            lambda: sp2.descargar_a_disco(
                SITIO, "/sites/Tech/Colombia/AXA/charla.mp4", en_disco))
        comprobar(not os.path.exists(en_disco),
                  "ni se llegó a crear el archivo en tu PC")
        comprobar_que_protesta(
            "un vídeo no se mueve",
            lambda: sp2.mover(SITIO, "/sites/Tech/Colombia/AXA/charla.mp4",
                              "/sites/Tech/Archivo", "charla.mp4", tamano=10))
        comprobar_que_protesta(
            "un vídeo no se copia (ni al mismo sitio ni a otro)",
            lambda: sp2.copiar(SITIO, "/sites/Tech/Colombia/AXA/charla.mp4",
                               SITIO, "/sites/Tech/Archivo", "charla.mp4",
                               tamano=10))
        comprobar(llamadas == [],
                  "de un vídeo no se pidió nada a SharePoint (ni MoveTo ni "
                  "$value)")

        # --- El plan: lo que venga de fuera tampoco se toca, y se dice.
        sp3 = SharePoint()
        utiles, omitidos, otros = sp3._solo_del_plan(
            [("AXA/acta.docx", "Actas"), ("AXA/charla.mp4", "Vídeos"),
             ("notas.txt", "Otros"), ("AXA/transcripcion.docx", "omitir")])
        comprobar(utiles == [("AXA/acta.docx", "Actas")],
                  "del plan solo pasan los documentos de oficina")
        comprobar(otros == [("AXA/charla.mp4", "Vídeos"),
                            ("notas.txt", "Otros")],
                  "y los demás se apartan para decirlo en el resumen")
        comprobar(omitidos == [("AXA/transcripcion.docx", "omitir")],
                  "lo que la regla deja fuera se aparta aparte (no es un fallo)")

        movidos, avisos3 = [], []
        sp3.on_aviso = avisos3.append
        sp3.carpeta_existe = lambda sitio, ruta: True
        sp3.archivo_existe = lambda sitio, ruta: False
        sp3.listar_carpetas = lambda url: []
        sp3.crear_carpeta = lambda sitio, ruta: None
        sp3.mover = lambda sitio, origen, destino, nombre, tamano=None: \
            movidos.append(nombre)
        resumen = sp3.clasificar_y_mover(
            "https://contoso.sharepoint.com/sites/Tech/Colombia",
            "https://contoso.sharepoint.com/sites/Tech/Archivo",
            [("AXA/acta.docx", "Actas"), ("AXA/charla.mp4", "Vídeos")])
        comprobar(movidos == ["acta.docx"], "el documento sí se mueve")
        comprobar("Movidos 1 de 2" in resumen
                  and "No los toqué (solo ppt, pptx, doc y docx):" in resumen
                  and "AXA/charla.mp4 → Vídeos" in resumen,
                  "y el resumen dice lo que sí pasó y lo que se quedó fuera")
        comprobar(any("Me salté 1 archivo(s) que no son" in a for a in avisos3),
                  "avisando por el chat antes de empezar")

        sp4 = SharePoint()
        resumen = sp4.clasificar_y_mover(
            "https://contoso.sharepoint.com/sites/Tech/Colombia",
            "https://contoso.sharepoint.com/sites/Tech/Archivo",
            [("AXA/charla.mp4", "Vídeos")])
        comprobar(resumen.startswith("No moví nada")
                  and "AXA/charla.mp4 → Vídeos" in resumen,
                  "si el plan era todo vídeos, no se mueve nada y se explica")

        # --- Y con la lista vacía se pasa todo (así se quita el filtro).
        SharePoint.EXTENSIONES_UTILES = ()
        comprobar(sp.es_util("charla.mp4")
                  and sp.extensiones_texto() == "cualquier archivo",
                  "con la lista vacía se pasa todo (filtro apagado)")
    finally:
        SharePoint.EXTENSIONES_UTILES = EXTENSIONES_REALES


def pruebas_omitir_lo_que_la_regla_deja_fuera():
    """Un archivo que la regla deja fuera no se pasa: ni se copia ni se baja.

    La regla de propuestas manda omitir transcripciones, anexos técnicos,
    Consultoría y lo que hay que revisar. Antes eso acababa copiado (y, si era
    grande, bajado a tu PC) porque el plan tenía que decir algo.
    """
    print("\n== Lo que la regla deja fuera no se toca ==")
    sp = SharePoint()
    for palabra in ("omitir", "Omitir", "SE OMITE", "se omite: no es propuesta",
                    "[REVISAR]", "revisar: falta el alcance", "no aplica",
                    "pendiente", "n/a", "Ninguna", "descartar"):
        comprobar(sp.es_omitir(palabra), f"'{palabra}' se lee como 'no tocar'")
    for palabra in ("Actas", "Otros", "Assessment", "Assessment/Data",
                    "Consultoría", "Obligaciones", "Revisadas", ""):
        comprobar(not sp.es_omitir(palabra),
                  f"'{palabra}' es una carpeta de verdad (no se omite)")

    utiles, omitidos, otros = sp._solo_del_plan(
        [("a.docx", "Actas"), ("b.docx", "omitir"), ("c.docx", "revisar"),
         ("charla.mp4", "Vídeos")])
    comprobar(utiles == [("a.docx", "Actas")],
              "el plan de verdad solo lleva lo que se pasa")
    comprobar(omitidos == [("b.docx", "omitir"), ("c.docx", "revisar")],
              "y aparte los que la regla deja fuera")
    comprobar(otros == [("charla.mp4", "Vídeos")],
              "un vídeo sigue siendo 'no es documento', no 'omitido'")

    van, fuera = sp.separar_omitidos(
        [("a.docx", "Actas"), ("b.docx", "omitir"), ("c.docx", "revisar")])
    comprobar(van == [("a.docx", "Actas")],
              "'separar_omitidos' deja el plan tal como viene, sin adornos")
    comprobar(fuera == [("b.docx", "omitir"), ("c.docx", "revisar")],
              "y aparte los que no se pasan")

    # De punta a punta: aunque el plan lo traiga, no se mueve ni crea carpeta.
    sp2 = SharePoint()
    creadas, movidos, avisos = [], [], []
    sp2.listar_carpetas = lambda url: []
    sp2.crear_carpeta = lambda sitio, rel: creadas.append(rel)
    sp2.carpeta_existe = lambda sitio, rel: True
    sp2.archivo_existe = lambda sitio, rel: False
    sp2.on_aviso = avisos.append
    sp2.mover = lambda sitio, origen, destino, nombre, tamano=None: \
        movidos.append(nombre)
    resumen = sp2.clasificar_y_mover(
        "https://contoso.sharepoint.com/sites/Tech/Colombia",
        "https://contoso.sharepoint.com/sites/Tech/Archivo",
        [("acta.docx", "Actas"), ("acta-vieja.docx", "omitir"),
         ("transcripcion.docx", "revisar: no es propuesta")])
    comprobar(movidos == ["acta.docx"], "solo se pasa el que la regla manda pasar")
    comprobar(creadas == ["/sites/Tech/Archivo/Actas"],
              "y no se crea una carpeta 'omitir' con esa excusa")
    comprobar("No los toqué (tu regla los deja fuera):" in resumen
              and "acta-vieja.docx → omitir" in resumen,
              "el resumen dice cuáles no se tocaron")
    comprobar(any("dejo 2 archivo(s) sin tocar" in a for a in avisos),
              "y lo avisa por el chat antes de empezar")

    # Si la regla deja fuera a todos, no se hace nada y se dice (sin drama).
    sp3 = SharePoint()
    resumen = sp3.clasificar_y_mover(
        "https://contoso.sharepoint.com/sites/Tech/Colombia",
        "https://contoso.sharepoint.com/sites/Tech/Archivo",
        [("a.docx", "omitir"), ("b.docx", "omitir")])
    comprobar(resumen.startswith("No moví nada") and "a.docx → omitir" in resumen,
              "si todos están omitidos, no se mueve nada y se explica")


def pruebas_subcarpetas_anidadas():
    """La regla puede pedir rama y subrama ('Assessment/Data'): dos niveles."""
    print("\n== Subcarpetas de dos niveles (rama y subrama) ==")
    sp = SharePoint()
    creadas = []
    sp.crear_carpeta = lambda sitio, rel: creadas.append(rel)
    comprobar(sp._asegurar_carpetas(SITIO, "/sites/Tech/Archivo",
                                    "Assessment/Data")
              == "/sites/Tech/Archivo/Assessment/Data",
              "devuelve la ruta del destino con sus dos niveles")
    comprobar(creadas == ["/sites/Tech/Archivo/Assessment",
                          "/sites/Tech/Archivo/Assessment/Data"],
              "crea antes la carpeta madre y después la hija")

    creadas.clear()
    sp._asegurar_carpetas(SITIO, "/sites/Tech/Archivo", "Assessment/Data",
                          {"/sites/Tech/Archivo/Assessment/Data"})
    comprobar(creadas == ["/sites/Tech/Archivo/Assessment"],
              "lo que ya venía en el listado del destino no se vuelve a crear")

    creadas.clear()
    comprobar(sp._asegurar_carpetas(SITIO, "/sites/Tech/Archivo", "")
              == "/sites/Tech/Archivo",
              "sin subcarpeta, el destino es la carpeta tal cual")
    comprobar(creadas == [], "y no se crea nada por el camino")

    # De punta a punta: el archivo acaba dentro de su subrama, no en la rama.
    sp2 = SharePoint()
    creadas2, movidos2 = [], []
    sp2.listar_carpetas = lambda url: []
    sp2.crear_carpeta = lambda sitio, rel: creadas2.append(rel)
    sp2.carpeta_existe = lambda sitio, rel: True
    sp2.archivo_existe = lambda sitio, rel: False
    sp2.mover = lambda sitio, origen, destino, nombre, tamano=None: \
        movidos2.append((destino, nombre))
    resumen = sp2.clasificar_y_mover(
        "https://contoso.sharepoint.com/sites/Tech/Colombia",
        "https://contoso.sharepoint.com/sites/Tech/Archivo",
        [("acta.docx", "Assessment/Data"), ("otra.docx", "Assessment/Data")])
    comprobar(movidos2 == [("/sites/Tech/Archivo/Assessment/Data", "acta.docx"),
                           ("/sites/Tech/Archivo/Assessment/Data", "otra.docx")],
              "los dos archivos quedan en la subrama (rama y subrama)")
    comprobar(creadas2 == ["/sites/Tech/Archivo/Assessment",
                           "/sites/Tech/Archivo/Assessment/Data"],
              "las carpetas se crean una sola vez, no una por archivo")
    comprobar("acta.docx → Assessment/Data" in resumen,
              "el resumen enseña la subrama donde quedó cada uno")

    # Y los adornos de la IA no llegan al nombre de la carpeta.
    comprobar(SharePoint.limpiar_subcarpeta(" **Assessment** / 'Data' ")
              == "Assessment/Data",
              "los adornos de la IA ('**', comillas) se quitan del nombre")
    comprobar(SharePoint.limpiar_subcarpeta("Assessment|Data")
              == "AssessmentData",
              "y los caracteres que SharePoint no acepta también")


def pruebas_catalogo_del_destino():
    """Al destino se le mira lo que YA tiene, para no inventar carpetas parecidas.

    El caso real: la regla dice 'Assessment/Data Analysis' y en el destino ya
    está la 'Assessment/Data' de siempre. Como solo se listaba el ORIGEN, la IA
    clasificaba a ciegas y acababa creando una carpeta nueva al lado de la que
    ya existía (y al mes siguiente otra vez). Aquí se prueba que el destino se
    lee con sus dos niveles (``catalogo_carpetas``), que esa lista hace que el
    archivo acabe en la carpeta que está hecha (``alinear_carpeta``), que se
    corta y se dice cuando hay demasiadas carpetas, y que al mover se usan las
    que ya existen aunque nadie le pase el catálogo (``clasificar_y_mover``).
    """
    print("\n== Las carpetas que ya existen en el destino ==")
    url_archivo = "https://contoso.sharepoint.com/sites/Tech/Archivo"
    sp = SharePoint()
    avisos = []
    sp.on_aviso = avisos.append
    sp.carpeta_existe = lambda sitio, rel: True

    def api_falsa(sitio, ruta, metodo="GET", datos=None):
        # Lo más específico primero: la subrama se pide por su ruta completa.
        if "'/sites/Tech/Archivo/Assessment'" in ruta:
            return {"d": {"results": [
                {"Name": "Data", "ServerRelativeUrl":
                    "/sites/Tech/Archivo/Assessment/Data"},
                {"Name": "   ", "ServerRelativeUrl":
                    "/sites/Tech/Archivo/Assessment/blanco"},
                {"ServerRelativeUrl":
                    "/sites/Tech/Archivo/Assessment/sin_nombre"}]}}
        if "'/sites/Tech/Archivo'" in ruta:
            return {"d": {"results": [
                {"Name": "Assessment",
                 "ServerRelativeUrl": "/sites/Tech/Archivo/Assessment"},
                {"Name": "Informes",
                 "ServerRelativeUrl": "/sites/Tech/Archivo/Informes"}]}}
        return {"d": {"results": []}}

    sp._api = api_falsa
    catalogo = sp.catalogo_carpetas(url_archivo)
    comprobar(catalogo == ["Assessment", "Informes", "Assessment/Data"],
              "lee el destino con sus dos niveles: la rama y, dentro, su subrama")
    comprobar(sp.nota_catalogo == "", "y no avisa de nada: los pudo ver todos")
    comprobar(any("nivel 1 de 2" in a for a in avisos)
              and any("nivel 2 de 2" in a for a in avisos),
              "el chat va diciendo qué nivel del destino está mirando")
    comprobar(all("blanco" not in c and "sin_nombre" not in c for c in catalogo),
              "lo que viene sin nombre (o en blanco) no entra en el catálogo")

    # Un destino que todavía no existe no es un error: no hay nada que reusar.
    sp2 = SharePoint()

    def api_prohibida(*argumentos, **opciones):
        raise AssertionError("no tenía que preguntar por las carpetas")

    sp2.carpeta_existe = lambda sitio, rel: False
    sp2._api = api_prohibida
    comprobar(sp2.catalogo_carpetas(url_archivo) == [],
              "un destino que aún no existe da una lista vacía (se creará luego)")

    # Tope de carpetas devueltas: se corta y se dice, que callarse sería peor.
    sp3 = SharePoint()
    avisos3 = []
    sp3.on_aviso = avisos3.append
    sp3.carpeta_existe = lambda sitio, rel: True
    sp3.MAXIMO_CARPETAS_CATALOGO = 2
    sp3._api = api_falsa
    comprobar(sp3.catalogo_carpetas(url_archivo) == ["Assessment", "Informes"],
              "si hay más carpetas de las que se enseñan, el catálogo se corta")
    comprobar("no miré más a fondo" in sp3.nota_catalogo,
              "y lo deja escrito ('nota_catalogo'): no estando completo, que se "
              "sepa")
    comprobar(any("no miré más a fondo" in a for a in avisos3),
              "y también lo dice por el chat")
    comprobar(sp.nota_catalogo == "",
              "mientras que cuando sí se ven todas, no hay nada que decir")

    sp3b = SharePoint()
    sp3b.carpeta_existe = lambda sitio, rel: True
    sp3b.MAXIMO_CARPETAS_CATALOGO = 1
    sp3b._api = api_falsa
    comprobar(sp3b.catalogo_carpetas(url_archivo) == ["Assessment"],
              "y si el tope se llena a media lista, se corta por donde iba")
    comprobar("te enseño las primeras 1" in sp3b.nota_catalogo,
              "diciendo cuántas enseña y que hay más")

    # Tope de carpetas miradas: si se queda corto, lo avisa en vez de callarse.
    sp4 = SharePoint()
    sp4.carpeta_existe = lambda sitio, rel: True
    sp4.MAXIMO_CARPETAS = 1
    sp4._api = api_falsa
    sp4.catalogo_carpetas(url_archivo)
    comprobar("Me quedé en 1 carpetas del destino" in sp4.nota_catalogo,
              "si deja carpetas del destino sin mirar, lo dice (puede haber más)")

    # Y lo que se enseña acaba eligiendo la carpeta que YA existe.
    catalogo_destino = ["Assessment", "Assessment/Data",
                        "Assessment/Data Analysis", "Informes"]
    comprobar(SharePoint.alinear_carpeta("assessment/data", catalogo_destino)
              == "Assessment/Data",
              "'assessment/data' se guarda en la 'Assessment/Data' que ya está")
    comprobar(SharePoint.alinear_carpeta("Assessment / DATA", catalogo_destino)
              == "Assessment/Data",
              "igual con 'Assessment / DATA': mismo nombre, otra grafía")
    comprobar(SharePoint.alinear_carpeta("Assessment/Data Analysis",
                                         catalogo_destino)
              == "Assessment/Data Analysis",
              "y 'Data Analysis' no se mete dentro de 'Data', que es otra carpeta")
    comprobar(SharePoint.alinear_carpeta("Assessment/data analysis",
                                         catalogo_destino)
              == "Assessment/Data Analysis",
              "aunque la IA la escriba en minúsculas")
    comprobar(SharePoint.alinear_carpeta("Assessment/Data/2024",
                                         catalogo_destino)
              == "Assessment/Data/2024",
              "un nivel que el destino no tiene se deja como viene (será nuevo)")
    comprobar(SharePoint.alinear_carpeta("Clientes/Nueva", catalogo_destino)
              == "Clientes/Nueva", "una carpeta nueva de verdad no se toca")
    comprobar(SharePoint.alinear_carpeta("  Assessment/Data  ", [])
              == "Assessment/Data",
              "sin catálogo solo se quitan los adornos de siempre")
    comprobar(SharePoint.alinear_carpeta("", catalogo_destino) == "",
              "sin carpeta no hay nada que alinear (el archivo va a la raíz)")

    # Si hay dos carpetas distintas que se escriben casi igual, no se adivina.
    ambiguo = ["Assessment/Data-X", "Assessment/Data X"]
    comprobar(SharePoint.alinear_carpeta("Assessment/data x", ambiguo)
              == "Assessment/data x",
              "con dos carpetas parecidas se deja tal cual: mejor no meterlo mal")

    # De lo que ya existe (rutas absolutas) salen las rutas relativas del destino.
    existentes = {"/sites/Tech/Archivo",
                  "/sites/Tech/Archivo/Assessment",
                  "/sites/Tech/Archivo/Assessment/Data",
                  "/sites/Tech/Archivo/Operacion/2024/Informes",
                  "/sites/Otro/Archivo/Assessment"}
    del_destino = SharePoint._catalogo_de_existentes(existentes,
                                                     "/sites/Tech/Archivo")
    comprobar(sorted(del_destino) == ["Assessment", "Assessment/Data"],
              "saca la ruta relativa de lo que ya está (y solo lo del destino)")
    comprobar("Operacion/2024/Informes" not in del_destino
              and "Otro/Archivo/Assessment" not in del_destino,
              "lo que está a más de dos niveles (o en otro sitio) no entra")
    comprobar(SharePoint._catalogo_de_existentes(existentes,
                                                 "/sites/Tech/Archivo",
                                                 profundidad=1) == ["Assessment"],
              "con profundidad=1 se queda en el primer nivel")

    # De punta a punta: sin que nadie le pase el catálogo, las carpetas que el
    # destino ya lista (el primer nivel) sirven de red y no se crea otra.
    sp5 = SharePoint()
    creadas, movidos, avisos5 = [], [], []
    sp5.listar_carpetas = lambda url: [
        {"nombre": "Assessment", "ruta_rel": "/sites/Tech/Archivo/Assessment"}]
    sp5.crear_carpeta = lambda sitio, rel: creadas.append(rel)
    sp5.carpeta_existe = lambda sitio, rel: True
    sp5.archivo_existe = lambda sitio, rel: False
    sp5.on_aviso = avisos5.append
    sp5.mover = lambda sitio, origen, destino, nombre, tamano=None: \
        movidos.append((destino, nombre))
    resumen = sp5.clasificar_y_mover(
        "https://contoso.sharepoint.com/sites/Tech/Colombia", url_archivo,
        [("acta.docx", "assessment")])
    comprobar(movidos == [("/sites/Tech/Archivo/Assessment", "acta.docx")],
              "sin catálogo se usan las carpetas que el destino ya tiene")
    comprobar(creadas == [], "y no se crea ninguna carpeta por el camino")
    comprobar(any("uso las que están hechas" in a for a in avisos5),
              "el cambio se cuenta por el chat, no se hace a escondidas")
    comprobar("acta.docx → Assessment" in resumen,
              "y el resumen enseña la carpeta de verdad, no la que pidió la IA")

    # Con el catálogo del destino (el que se le dio a la IA) va igual de derecho.
    creadas.clear()
    movidos.clear()
    avisos5.clear()
    sp5.clasificar_y_mover(
        "https://contoso.sharepoint.com/sites/Tech/Colombia", url_archivo,
        [("acta.docx", "Assessment/Data")],
        carpetas=["Assessment", "Assessment/Data"])
    comprobar(movidos == [("/sites/Tech/Archivo/Assessment/Data", "acta.docx")],
              "con el catálogo del destino el archivo va a la subrama que existe")
    comprobar(creadas == ["/sites/Tech/Archivo/Assessment/Data"],
              "y solo se crea el nivel que de verdad faltaba")
    comprobar(not any("uso las que están hechas" in a for a in avisos5),
              "si el plan ya venía bien, no hay nada que avisar")

def _docx_de_prueba(titulo, parrafos=()):
    """Un .docx de verdad, en memoria, para probar lo que se lee por dentro.

    Se hace con python-docx (el mismo que usa ``modules.pistas``) y se guarda en
    un ``BytesIO``: así se prueba la lectura de verdad, sin archivos por medio.
    """
    from io import BytesIO
    import docx
    documento = docx.Document()
    documento.core_properties.title = titulo
    for parrafo in parrafos:
        documento.add_paragraph(parrafo)
    guardado = BytesIO()
    documento.save(guardado)
    return guardado.getvalue()


def _pptx_de_prueba(titulo, texto=""):
    """Un .pptx de verdad, en memoria (una diapositiva con su título)."""
    from io import BytesIO
    from pptx import Presentation
    presentacion = Presentation()
    diapositiva = presentacion.slides.add_slide(presentacion.slide_layouts[1])
    diapositiva.shapes.title.text = titulo
    diapositiva.placeholders[1].text = texto
    guardado = BytesIO()
    presentacion.save(guardado)
    return guardado.getvalue()


def pruebas_lo_que_no_se_supo_clasificar():
    """Un archivo 'sin clasificar' no se pasa: se aparta y se le mira por dentro.

    El caso real: 'PROPUESTA COMERCIAL bfddfa36 RFI Soporte y-o
    administración.pptx'. Ese nombre no dice de qué disciplina es, así que la IA
    no puede saber a qué carpeta va: antes eso salía como 'omitir' (no se movía
    nada y nadie entendía por qué) o como raíz del destino (todo amontonado).
    Ahora el plan lo dice con ``DUDOSO`` ('sin clasificar'), que NO es lo mismo
    que lo que la regla deja fuera: es una duda del programa y tiene arreglo.

    Aquí se prueba que se reconoce, que se aparta de los otros dos grupos (para
    no contarlo como si el usuario lo hubiera pedido), que nunca se pasa aunque
    nadie haga el segundo intento, que el resumen lo cuenta como duda y que
    ``pistas_de`` saca de los bytes lo que dice el documento por dentro (título y
    primeras líneas) solo de los archivos que se le piden.
    """
    print("\n== Lo que la IA no supo clasificar ('sin clasificar') ==")
    sp = SharePoint()
    for palabra in ("sin clasificar", "SIN CLASIFICAR", "  sin   clasificar  ",
                    "[sin clasificar]", "Sin Clasificar", "**sin clasificar**"):
        comprobar(sp.es_dudoso(palabra),
                  f"'{palabra}' se lee como 'no sé a qué carpeta va'")
    for palabra in ("Actas", "", "omitir", "revisar", "(raíz)", "Otros"):
        comprobar(not sp.es_dudoso(palabra), f"'{palabra}' no es una duda")

    comprobar(sp.es_omitir(SharePoint.DUDOSO),
              "'sin clasificar' tampoco se pasa: es_omitir lo reconoce")
    utiles, omitidos, otros = sp._solo_del_plan(
        [("a.docx", "Actas"), ("b.docx", SharePoint.DUDOSO)])
    comprobar(utiles == [("a.docx", "Actas")]
              and omitidos == [("b.docx", SharePoint.DUDOSO)] and not otros,
              "si nadie hace el segundo intento, el dudoso no se pasa igual")

    van, dudosos = sp.separar_dudosos(
        [("a.docx", "Actas"), ("b.docx", SharePoint.DUDOSO),
         ("c.docx", "omitir")])
    comprobar(van == [("a.docx", "Actas"), ("c.docx", "omitir")],
              "'separar_dudosos' deja el plan del primer intento")
    comprobar(dudosos == [("b.docx", SharePoint.DUDOSO)],
              "y aparte solo los que no se supieron clasificar")

    fuera, dudas = sp._partir_omitidos(
        [("a.docx", "omitir"), ("b.docx", "revisar: falta el alcance"),
         ("c.docx", SharePoint.DUDOSO)])
    comprobar(fuera == [("a.docx", "omitir"),
                        ("b.docx", "revisar: falta el alcance")],
              "'_partir_omitidos' junta lo que la regla deja fuera")
    comprobar(dudas == [("c.docx", SharePoint.DUDOSO)],
              "y aparte las dudas del programa (no es lo mismo)")

    comprobar(SharePoint.alinear_carpeta(SharePoint.DUDOSO,
                                        ["Sin clasificar"]) == SharePoint.DUDOSO,
              "una duda no se confunde con una carpeta que se llame parecido")

    # De punta a punta: el plan lo trae, pero no se mueve ni crea carpeta.
    sp2 = SharePoint()
    creadas, movidos, avisos = [], [], []
    sp2.listar_carpetas = lambda url: []
    sp2.crear_carpeta = lambda sitio, rel: creadas.append(rel)
    sp2.carpeta_existe = lambda sitio, rel: True
    sp2.archivo_existe = lambda sitio, rel: False
    sp2.on_aviso = avisos.append
    sp2.mover = lambda sitio, origen, destino, nombre, tamano=None: \
        movidos.append(nombre)
    resumen = sp2.clasificar_y_mover(
        "https://contoso.sharepoint.com/sites/Tech/Colombia",
        "https://contoso.sharepoint.com/sites/Tech/Archivo",
        [("acta.docx", "Actas"), ("hash.pptx", SharePoint.DUDOSO)])
    comprobar(movidos == ["acta.docx"],
              "el 'sin clasificar' no se mueve (mejor quieto que mal puesto)")
    comprobar("hash.pptx" not in creadas and "/sites/Tech/Archivo/Actas" in creadas,
              "y no se crea una carpeta con esa duda como nombre")
    comprobar("No los toqué: no supe en qué carpeta ponerlos:" in resumen
              and f"hash.pptx → {SharePoint.DUDOSO}" in resumen,
              "el resumen lo cuenta como duda del programa")
    comprobar(any("no supe en qué carpeta ponerlos" in a for a in avisos)
              and not any("tu regla los deja fuera" in a for a in avisos),
              "y se avisa por el chat sin echarle la culpa a la regla")

    # Y lo que hace posible el segundo intento: leer el documento por dentro.
    sp3 = SharePoint()
    bajados = []
    docx = _docx_de_prueba("Propuesta de servicios de QA",
                           ["Cliente: Banco X.", "Alcance: pruebas de portales."])
    pptx = _pptx_de_prueba("Propuesta Comercial Portales", "RFI Soporte")

    def descargar_falso(sitio, rel, maximo=None):
        nombre = rel.rsplit("/", 1)[-1]
        bajados.append(rel)
        if nombre == "Protegida.docx":
            raise ErrorSharePoint("no pude bajar el archivo (está protegido)")
        return {"Propuesta.docx": docx, "Propuesta.pptx": pptx}.get(
            nombre, b"es un PDF: no se puede leer por dentro")

    sp3.descargar = descargar_falso
    sp3.on_aviso = lambda texto: None
    pistas = sp3.pistas_de(
        "https://contoso.sharepoint.com/sites/Tech/Colombia",
        ["AXA/Propuesta.docx", "AXA/2024/Propuesta.pptx", "AXA/anexo.pdf",
         "Protegida.docx"])
    comprobar(bajados == ["/sites/Tech/Colombia/AXA/Propuesta.docx",
                          "/sites/Tech/Colombia/AXA/2024/Propuesta.pptx",
                          "/sites/Tech/Colombia/AXA/anexo.pdf",
                          "/sites/Tech/Colombia/Protegida.docx"],
              "se baja solo lo que se le pide, con la ruta de su carpeta")
    comprobar("Propuesta de servicios de QA" in pistas["AXA/Propuesta.docx"]
              and "Banco X" in pistas["AXA/Propuesta.docx"],
              "de un .docx saca el título y las primeras líneas")
    comprobar("Propuesta Comercial Portales" in pistas["AXA/2024/Propuesta.pptx"],
              "y de un .pptx el título de la diapositiva")
    comprobar("AXA/anexo.pdf" not in pistas,
              "del que no se puede leer nada (un PDF) no hay pista")
    comprobar("Protegida.docx" not in pistas,
              "y de un documento que no baja tampoco (y no se cae)")
    comprobar(sp3.pistas_de("https://x.sharepoint.com/sites/T", []) == {},
              "sin archivos que mirar no se baja nada")


def main_pruebas():
    pruebas_parseo_url()
    pruebas_odata()
    pruebas_api()
    pruebas_listados()
    pruebas_busqueda_en_subcarpetas()
    pruebas_paginado()
    pruebas_avisos_de_progreso()
    pruebas_existencia()
    pruebas_clasificar_y_mover()
    pruebas_plan_con_rutas()
    pruebas_binario_y_copiar()
    pruebas_archivos_grandes()
    pruebas_resolver_url()
    pruebas_pestana_muerta()
    pruebas_javascript_valido()
    pruebas_clasificacion_naive()
    pruebas_solo_documentos_de_oficina()
    pruebas_omitir_lo_que_la_regla_deja_fuera()
    pruebas_lo_que_no_se_supo_clasificar()
    pruebas_subcarpetas_anidadas()
    pruebas_catalogo_del_destino()
    print(f"\nTOTAL: {total[0]} comprobaciones, {len(fallos)} fallos")
    for fallo in fallos:
        print(f"  - {fallo}")
    return 1 if fallos else 0


if __name__ == "__main__":
    raise SystemExit(main_pruebas())
