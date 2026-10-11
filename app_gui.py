import customtkinter as ctk
import glob
import json
import os
import re
import webbrowser
import threading
from datetime import datetime
from tkinter import filedialog

from PIL import Image

from modules.api_brain import APIBrain
from modules.system_actions import SystemActions
from modules.voice_engine import VoiceEngine
from modules.voice_output import VoiceOutput
from modules.ocr_engine import ocr_imagen, motor as ocr_motor
from modules.cloud_matrix import CloudMatrix, COLUMNAS_DEFECTO
from modules import rutas

ctk.set_appearance_mode("Dark")
ctk.set_default_color_theme("blue")

# ---------------- Paleta oscura estilo Claude ----------------
COLOR_FONDO = "#1b1a19"
COLOR_SIDEBAR = "#141312"
COLOR_PANEL = "#1f1e1d"
COLOR_TARJETA = "#2a2926"
COLOR_HOVER = "#34332f"
COLOR_BORDE = "#3b3a37"
COLOR_TEXTO = "#ece9e2"
COLOR_TEXTO_SUAVE = "#a5a19a"
COLOR_ACENTO = "#d97757"
COLOR_ACENTO_HOVER = "#c2653f"
COLOR_USUARIO = "#33322e"

# ---------------- Adjuntos ----------------
CARPETA_ADJUNTOS = rutas.ADJUNTOS
EXT_IMAGEN = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tif", ".tiff", ".ico"}
EXT_TEXTO = {".txt", ".md", ".markdown", ".py", ".json", ".csv", ".log", ".html", ".htm",
             ".xml", ".js", ".ts", ".css", ".c", ".cpp", ".h", ".hpp", ".java", ".cs",
             ".sql", ".yaml", ".yml", ".ini", ".cfg", ".bat", ".ps1", ".sh", ".env",
             ".toml", ".rst", ".pyw", ".tsx", ".jsx", ".php", ".rb", ".go", ".rs"}
MAX_TEXTO_ADJUNTO = 12000  # caracteres máximos que se envían por archivo a la IA

# El Entry de Tk se vuelve extremadamente lento (O(n²)) al insertar cadenas
# enormes: 20.000 car. ~0,5 s y 100.000 car. ~13 s, lo que congela la interfaz.
# Por encima de este umbral, el texto pegado se guarda como documento adjunto.
UMBRAL_PEGADO_TEXTO = 8000

# ---------------- Documentos de texto ----------------
CARPETA_DOCUMENTOS = rutas.DOCUMENTOS

# ---------------- Ubicaciones (rutas/URLs) citadas en los mensajes ----------------
# Se detectan para ofrecer un botón que las abra con un solo clic.
RE_UBICACION = re.compile(
    r'"([A-Za-z]:[\\/][^"\r\n]+)"'        # "C:\ruta con espacios\archivo.txt"
    r'|`([A-Za-z]:[\\/][^`\r\n]+)`'       # `C:\ruta\archivo.txt`
    r'|(\\\\[^\s<>"\'`)\]]+)'             # \\servidor\recurso\archivo.txt (UNC)
    r'|([A-Za-z]:[\\/][^\s<>"\'`)\]]*)'   # C:\ruta\archivo.txt (sin espacios)
    r'|(https?://[^\s<>"\'`)\]]+)'        # https://ejemplo.com
)


class AppGUI(ctk.CTk):
    def __init__(self, config_path=rutas.CONFIG_USUARIO, voice_output=None):
        super().__init__()

        self.config_path = config_path
        self.escuchando = False
        self.chat_activo = None      # None = sin definir, False = bienvenida, True = conversación
        self._procesando = False
        self._ind_pensando = None
        self._pegado_abajo = True    # True si el usuario está viendo el final del chat
        self.adjuntos = []          # lista de dicts con los archivos adjuntos pendientes
        self._burbujas = []         # registros de las burbujas visibles (para selección)
        self._seleccion = []        # registros de las burbujas seleccionadas
        self.conversacion_actual = None  # ruta del archivo de la conversación activa
        os.makedirs(CARPETA_ADJUNTOS, exist_ok=True)

        # Instanciar módulos
        self.brain = APIBrain(config_path)
        self.system = SystemActions(config_path)
        self.voice = VoiceEngine(config_path)

        # Estado del OCR (lectura de texto dentro de imágenes adjuntas)
        self.ocr_motor = ocr_motor()
        print(f"OCR activo: {self.ocr_motor}" if self.ocr_motor
              else "OCR no disponible: no se leerá el texto de las imágenes.")

        # Salida de voz (si se recibe desde main.py o se crea por defecto)
        self.voice_output = voice_output if voice_output else VoiceOutput()

        # Alertas por voz (pomodoro) apuntan al reproductor
        self.system.productivity.on_alerta = self.reproducir_voz

        # Asistente de reuniones: alertas por voz + resumen con la IA + monitoreo
        self.system.meeting.on_alerta = self.reproducir_voz
        self.system.meeting.on_resumen = self._resumir_reunion
        self.system.meeting.on_metadatos = self._metadatos_reunion
        self.system.meeting.iniciar_monitoreo()

        # Herramientas de preventa: generación de documentos con la IA
        self.system.preventa.on_generar = self._generar_preventa

        # Configuración de la ventana principal
        self.title("Jarvis")
        self.geometry("1120x760")
        self.resizable(True, True)
        self.minsize(900, 600)
        self.configure(fg_color=COLOR_FONDO)

        self._crear_interfaz()
        self._cargar_config_en_inputs()
        self._cargar_conversaciones_sidebar()

    # ---------------- Construcción de la interfaz ----------------
    def _crear_interfaz(self):
        contenedor = ctk.CTkFrame(self, fg_color="transparent")
        contenedor.pack(fill="both", expand=True)

        self._crear_sidebar(contenedor)

        # Área principal (derecha)
        self.area_principal = ctk.CTkFrame(contenedor, fg_color=COLOR_PANEL, corner_radius=0)
        self.area_principal.pack(side="left", fill="both", expand=True)

        # Dos vistas dentro del área principal
        self.vista_chat = ctk.CTkFrame(self.area_principal, fg_color="transparent")
        self.vista_ajustes = ctk.CTkFrame(self.area_principal, fg_color="transparent")

        self._crear_vista_chat(self.vista_chat)
        self._crear_vista_ajustes(self.vista_ajustes)

        self._mostrar_vista("chat")
        self._mostrar_estado_chat(False)

    def _crear_sidebar(self, parent):
        self.sidebar = ctk.CTkFrame(parent, width=272, fg_color=COLOR_SIDEBAR, corner_radius=0)
        self.sidebar.pack(side="left", fill="y")
        self.sidebar.pack_propagate(False)

        # Cabecera con logo + título
        cabecera = ctk.CTkFrame(self.sidebar, fg_color="transparent")
        cabecera.pack(fill="x", padx=16, pady=(18, 10))
        ctk.CTkLabel(cabecera, text="\u2726", font=("Segoe UI", 20),
                     text_color=COLOR_ACENTO).pack(side="left")
        ctk.CTkLabel(cabecera, text="Jarvis", font=("Segoe UI", 17, "bold"),
                     text_color=COLOR_TEXTO).pack(side="left", padx=(8, 0))

        # Botón "Nuevo chat"
        self.btn_nuevo = ctk.CTkButton(
            self.sidebar, text="\uff0b   Nuevo chat", anchor="w",
            font=("Segoe UI", 13), height=38, corner_radius=8,
            fg_color=COLOR_TARJETA, hover_color=COLOR_HOVER, text_color=COLOR_TEXTO,
            command=self._nueva_conversacion,
        )
        self.btn_nuevo.pack(fill="x", padx=14, pady=(4, 16))

        # Encabezado de sección
        ctk.CTkLabel(self.sidebar, text="CONVERSACIONES", font=("Segoe UI", 10, "bold"),
                     text_color=COLOR_TEXTO_SUAVE).pack(anchor="w", padx=20, pady=(0, 4))

        # Lista scrollable de conversaciones
        self.lista_conv = ctk.CTkScrollableFrame(self.sidebar, fg_color="transparent")
        self.lista_conv.pack(fill="both", expand=True, padx=8, pady=(0, 8))

        # Pie: perfil + ajustes
        pie = ctk.CTkFrame(self.sidebar, fg_color="transparent")
        pie.pack(side="bottom", fill="x", padx=14, pady=12)
        self.avatar = ctk.CTkLabel(pie, text="J", width=30, height=30, corner_radius=15,
                                   fg_color=COLOR_ACENTO, text_color="#ffffff",
                                   font=("Segoe UI", 13, "bold"))
        self.avatar.pack(side="left")
        self.lbl_usuario = ctk.CTkLabel(pie, text="Usuario \u00b7 Free",
                                        font=("Segoe UI", 13), text_color=COLOR_TEXTO)
        self.lbl_usuario.pack(side="left", padx=(10, 0))
        ctk.CTkButton(pie, text="\u2699", width=32, height=32, corner_radius=8,
                      fg_color="transparent", hover_color=COLOR_HOVER,
                      text_color=COLOR_TEXTO_SUAVE, font=("Segoe UI", 16),
                      command=self._toggle_ajustes).pack(side="right")

    def _mostrar_vista(self, vista):
        self.vista_chat.pack_forget()
        self.vista_ajustes.pack_forget()
        if vista == "ajustes":
            self.vista_ajustes.pack(fill="both", expand=True)
            self._cargar_reglas_ui()
        else:
            self.vista_chat.pack(fill="both", expand=True)

    def _toggle_ajustes(self):
        if self.vista_ajustes.winfo_ismapped():
            self._mostrar_vista("chat")
        else:
            self._mostrar_vista("ajustes")

    def _crear_vista_chat(self, parent):
        # Espaciadores para centrar la bienvenida verticalmente. Se les da un alto
        # mínimo (por defecto un CTkFrame pide 200 px): así el reparto del espacio
        # sobrante es igual para los dos y el hero queda realmente centrado.
        self.spacer_top = ctk.CTkFrame(parent, fg_color="transparent", height=1)
        self.spacer_bottom = ctk.CTkFrame(parent, fg_color="transparent", height=1)

        # Hero (pantalla de bienvenida)
        self.hero = ctk.CTkFrame(parent, fg_color="transparent")
        ctk.CTkLabel(self.hero, text="\u2726", font=("Segoe UI", 44),
                     text_color=COLOR_ACENTO).pack()
        ctk.CTkLabel(self.hero, text="\u00bfEn qué puedo ayudarte hoy?",
                     font=("Segoe UI", 26), text_color=COLOR_TEXTO).pack(pady=(10, 0))
        ctk.CTkLabel(self.hero, text="Escribe tu mensaje o usa el control por voz.",
                     font=("Segoe UI", 13), text_color=COLOR_TEXTO_SUAVE).pack(pady=(6, 0))

        # Historial de conversación (burbujas)
        self.chat_scroll = ctk.CTkScrollableFrame(parent, fg_color="transparent")
        # Al cambiar el tamaño de la vista (ventana, barra de selección, maximizar...)
        # se reancla el final: el <Configure> del lienzo es la señal fiable del cambio
        # de alto disponible. add="+" conserva el manejador interno de
        # CTkScrollableFrame (que ajusta el ancho del contenido).
        self.chat_scroll._parent_canvas.bind("<Configure>", self._chat_redimensionado, add="+")
        # El lienzo avisa (yscrollcommand) cada vez que se desplaza: así sabemos si
        # el usuario está viendo el final y podemos reanclarlo al redimensionar.
        self.chat_scroll._parent_canvas.configure(yscrollcommand=self._scroll_actualizado)

        # Barra de entrada (compartida por ambos estados)
        self._crear_barra_input(parent)

    def _crear_barra_input(self, parent):
        self.barra_input = ctk.CTkFrame(parent, fg_color="transparent")

        # Fila de chips con los adjuntos pendientes (se rellena dinámicamente).
        # No se empaqueta aquí a propósito: un CTkFrame sin hijos conserva su alto
        # por defecto (200 px) y dejaría una banda muerta entre el chat y la caja
        # de escritura, cortando los mensajes por encima del cuadro de texto.
        # _actualizar_adjuntos_ui() la muestra solo cuando hay adjuntos.
        self.adjuntos_frame = ctk.CTkFrame(self.barra_input, fg_color="transparent")

        caja = ctk.CTkFrame(self.barra_input, fg_color=COLOR_TARJETA, corner_radius=18,
                            border_width=1, border_color=COLOR_BORDE)
        caja.pack(fill="x")
        self.caja_input = caja  # referencia para recolocar la fila de chips encima

        # Botón para adjuntar archivos (📎)
        self.btn_adjuntar = ctk.CTkButton(
            caja, text="\U0001F4CE", width=34, height=34, corner_radius=17,
            fg_color="transparent", hover_color=COLOR_HOVER, text_color=COLOR_TEXTO_SUAVE,
            font=("Segoe UI", 16), command=self._adjuntar_archivos)
        self.btn_adjuntar.pack(side="left", padx=(8, 0), pady=9)

        self.entry_chat = ctk.CTkEntry(
            caja, placeholder_text="Escribe un mensaje a Jarvis...",
            fg_color="transparent", border_width=0, text_color=COLOR_TEXTO,
            placeholder_text_color=COLOR_TEXTO_SUAVE, font=("Segoe UI", 14))
        self.entry_chat.pack(side="left", fill="x", expand=True, padx=(4, 8), pady=12)

        self.btn_enviar = ctk.CTkButton(
            caja, text="\u2191", width=36, height=36, corner_radius=18,
            fg_color=COLOR_ACENTO, hover_color=COLOR_ACENTO_HOVER, text_color="#ffffff",
            font=("Segoe UI", 18, "bold"), command=self._enviar_chat)
        self.btn_enviar.pack(side="right", padx=(0, 8), pady=8)

        self.entry_chat.bind("<Return>", lambda e: self._enviar_chat())
        self.entry_chat.bind("<Control-v>", self._pegar_portapapeles)
        self.entry_chat.bind("<<Paste>>", self._pegar_portapapeles)

        ctk.CTkLabel(self.barra_input,
                     text="Jarvis puede cometer errores. Verifica la información importante.",
                     font=("Segoe UI", 10), text_color=COLOR_TEXTO_SUAVE).pack(pady=(6, 0))

        # Barra de acciones de la selección de mensajes (oculta por defecto)
        self._crear_barra_seleccion(parent)

    def _mostrar_estado_chat(self, activo):
        """Alterna entre la pantalla de bienvenida (False) y la conversación (True)."""
        if activo == self.chat_activo:
            return
        self.chat_activo = activo
        for w in (self.spacer_top, self.spacer_bottom, self.hero,
                  self.chat_scroll, self.barra_input, self.barra_seleccion):
            w.pack_forget()
        if activo:
            self.chat_scroll.pack(fill="both", expand=True, padx=6, pady=(12, 0))
            self.barra_input.pack(fill="x", padx=110, pady=(4, 16))
            # Fuerza el cálculo de geometría del área scrollable ANTES de añadir
            # burbujas. Sin esto, al re-empaquetar el CTkScrollableFrame el canvas
            # mantiene un ancho inicial erróneo (p. ej. 1 px), las etiquetas se
            # envuelven a una altura enorme y la vista acaba mostrando un área
            # vacía (chat "en blanco").
            try:
                self.chat_scroll.update_idletasks()
            except Exception:
                pass
        else:
            self.spacer_top.pack(fill="both", expand=True)
            self.hero.pack(fill="x", padx=30, pady=(0, 20))
            self.barra_input.pack(fill="x", padx=110, pady=(0, 8))
            self.spacer_bottom.pack(fill="both", expand=True)
        # Refresca la barra de selección (visible solo si hay mensajes marcados)
        self._actualizar_barra_seleccion()

    # ---------------- Render de mensajes ----------------
    def _ruta_nueva_conversacion(self):
        """Ruta de un archivo nuevo: una conversación independiente por sesión."""
        base = os.path.join(rutas.CONVERSACIONES,
                            "historial_chat_" + datetime.now().strftime("%Y-%m-%d_%H%M%S"))
        ruta = base + ".md"
        i = 2
        while os.path.exists(ruta):
            ruta = f"{base}_{i}.md"
            i += 1
        return ruta

    def _agregar_chat(self, autor, texto, adjuntos=None):
        """Persiste el mensaje en la conversación activa y lo muestra como burbuja.

        Cada conversación se guarda en su propio archivo
        (historial_chat_AAAA-MM-DD_HHMMSS.md). Si no hay conversación activa se
        crea una nueva, de modo que "Nuevo chat" siempre abre un hilo aparte.
        """
        adjuntos = adjuntos or []
        linea_adj = "".join(f" 📎 {a['nombre']}" for a in adjuntos)
        nueva = False
        try:
            os.makedirs(rutas.CONVERSACIONES, exist_ok=True)
            if not self.conversacion_actual:
                self.conversacion_actual = self._ruta_nueva_conversacion()
                nueva = True
            hora = datetime.now().strftime("%H:%M")
            with open(self.conversacion_actual, "a", encoding="utf-8") as f:
                f.write(f"**[{hora}] {autor}:** {texto}{linea_adj}\n\n")
        except Exception as e:
            print(f"Error guardando historial: {e}")

        self.after(0, lambda: self._insertar_burbuja(autor, texto, adjuntos))
        if nueva:
            # Refresca la lista lateral para mostrar el hilo recién creado
            self.after(0, self._cargar_conversaciones_sidebar)

    def _insertar_burbuja(self, autor, texto, adjuntos=None):
        adjuntos = adjuntos or []
        self._mostrar_estado_chat(True)

        es_usuario = str(autor).strip().lower() in ("tú", "tu", "usuario", "yo")
        fila = ctk.CTkFrame(self.chat_scroll, fg_color="transparent")
        fila.pack(fill="x", padx=18, pady=5)

        color = COLOR_USUARIO if es_usuario else COLOR_TARJETA
        burbuja = ctk.CTkFrame(fila, fg_color=color, corner_radius=14)
        burbuja.pack(side="right" if es_usuario else "left")

        ancho_panel = self.area_principal.winfo_width()
        if ancho_panel <= 1:
            ancho_panel = 820
        wraplength = max(320, int(ancho_panel * 0.6))

        # Imágenes adjuntas: vista previa dentro de la burbuja
        for adj in adjuntos:
            if adj.get("tipo") == "imagen":
                img = self._crear_ctkimage(adj["ruta"], max_lado=240)
                if img is not None:
                    lbl_img = ctk.CTkLabel(burbuja, text="", image=img)
                    lbl_img._ctk_image = img  # evita recolección de basura
                    lbl_img.pack(padx=12, pady=(12, 2), anchor="w")

        # Chips de los archivos que no son imágenes
        for adj in adjuntos:
            if adj.get("tipo") != "imagen":
                ctk.CTkLabel(burbuja, text=f"\U0001F4C4 {adj['nombre']}",
                             font=("Segoe UI", 11), text_color=COLOR_TEXTO_SUAVE,
                             fg_color=COLOR_HOVER, corner_radius=8).pack(
                                 padx=12, pady=(6, 0), anchor="w")

        if texto:
            pad_y_sup = 8 if adjuntos else 11
            ctk.CTkLabel(burbuja, text=texto, wraplength=wraplength, justify="left",
                         anchor="w", font=("Segoe UI", 13),
                         text_color=COLOR_TEXTO).pack(padx=16, pady=(pad_y_sup, 11))

        # Registra la burbuja y permite seleccionarla con un clic
        registro = {"autor": autor, "texto": texto, "burbuja": burbuja, "seleccionada": False}
        self._burbujas.append(registro)
        self._vincular_clic(burbuja, lambda e, r=registro: self._toggle_seleccion(r))

        # Enlaces clicables a rutas/URLs del mensaje (se añaden DESPUÉS de
        # vincular el clic para no pisar el manejador de selección de burbuja).
        if texto:
            self._agregar_ubicaciones_burbuja(burbuja, texto)

        self.after(30, self._scroll_final)

    # ---------------- Selección de mensajes ----------------
    def _crear_barra_seleccion(self, parent):
        """Construye la barra de acciones que aparece al seleccionar mensajes."""
        self.barra_seleccion = ctk.CTkFrame(parent, fg_color=COLOR_TARJETA,
                                            corner_radius=10, border_width=1,
                                            border_color=COLOR_BORDE)
        ctk.CTkButton(self.barra_seleccion, text="\u2715", width=28, height=28,
                      corner_radius=8, fg_color="transparent",
                      hover_color=COLOR_ACENTO_HOVER, text_color=COLOR_TEXTO_SUAVE,
                      font=("Segoe UI", 12),
                      command=self._limpiar_seleccion).pack(side="left", padx=(10, 6), pady=6)
        self.lbl_sel = ctk.CTkLabel(self.barra_seleccion, text="0 mensajes seleccionados",
                                    font=("Segoe UI", 12), text_color=COLOR_TEXTO)
        self.lbl_sel.pack(side="left", pady=8)
        ctk.CTkButton(self.barra_seleccion, text="Guardar .txt", width=110, height=28,
                      corner_radius=8, fg_color=COLOR_ACENTO,
                      hover_color=COLOR_ACENTO_HOVER, text_color="#ffffff",
                      font=("Segoe UI", 12, "bold"),
                      command=self._guardar_seleccion_txt).pack(side="right", padx=(4, 10), pady=6)
        ctk.CTkButton(self.barra_seleccion, text="Copiar", width=80, height=28,
                      corner_radius=8, fg_color="transparent", hover_color=COLOR_HOVER,
                      text_color=COLOR_TEXTO, font=("Segoe UI", 12),
                      command=self._copiar_seleccion).pack(side="right", padx=4, pady=6)
        ctk.CTkButton(self.barra_seleccion, text="Todo", width=60, height=28,
                      corner_radius=8, fg_color="transparent", hover_color=COLOR_HOVER,
                      text_color=COLOR_TEXTO_SUAVE, font=("Segoe UI", 12),
                      command=self._seleccionar_todos).pack(side="right", padx=4, pady=6)

    def _vincular_clic(self, widget, callback):
        """Vincula el clic izquierdo a los widgets que reciben clics en la burbuja.

        customtkinter dibuja sobre un Canvas interno y expone los textos en un
        Label interno; se recorren esos componentes reales (sin duplicar) para
        que el clic funcione tanto sobre el texto como sobre el fondo.
        """
        for w in self._widgets_clickeables(widget):
            try:
                w.bind("<Button-1>", callback)
            except Exception:
                pass
            try:
                w.configure(cursor="hand2")
            except Exception:
                pass

    def _widgets_clickeables(self, raiz):
        """Devuelve los widgets tk (hojas) de un árbol customtkinter sin duplicados."""
        hojas = []
        vistos = set()
        pila = [raiz]
        while pila:
            w = pila.pop()
            if id(w) in vistos:
                continue
            vistos.add(id(w))
            es_customtk = any(hasattr(w, a) for a in ("_canvas", "_label", "_text_label"))
            if es_customtk:
                # Los widgets customtkinter no reciben el clic: lo reciben sus
                # componentes internos (_canvas para el fondo, _label para el texto).
                for attr in ("_canvas", "_label", "_text_label"):
                    hijo = getattr(w, attr, None)
                    if hijo is not None and id(hijo) not in vistos:
                        pila.append(hijo)
            else:
                hojas.append(w)
            try:
                hijos = w.winfo_children()
            except Exception:
                hijos = []
            for hijo in hijos:
                if id(hijo) not in vistos:
                    pila.append(hijo)
        return hojas

    # ---------------- Ubicaciones clicables (rutas/URLs) ----------------
    def _detectar_ubicaciones(self, texto):
        """Devuelve las rutas de archivo y URLs únicas presentes en el texto."""
        if not texto:
            return []
        encontradas = []
        vistos = set()
        for m in RE_UBICACION.finditer(texto):
            valor = next((g for g in m.groups() if g), None)
            if not valor:
                continue
            valor = valor.rstrip(".,;:")
            clave = valor.lower()
            if valor and clave not in vistos:
                vistos.add(clave)
                encontradas.append(valor)
        return encontradas

    def _es_url(self, destino):
        """True si el texto es un enlace con esquema (http://, https://, etc.)."""
        return bool(re.match(r"^[a-zA-Z][a-zA-Z0-9+.\-]*://", destino or ""))

    def _agregar_ubicaciones_burbuja(self, parent, texto):
        """Añade botones para abrir con un clic las rutas/URLs citadas."""
        ubicaciones = self._detectar_ubicaciones(texto)
        for i, destino in enumerate(ubicaciones):
            icono = "\U0001F310" if self._es_url(destino) else "\U0001F4C1"
            etiqueta = destino if len(destino) <= 80 else "..." + destino[-77:]
            ultimo = i == len(ubicaciones) - 1
            btn = ctk.CTkButton(
                parent, text=f"{icono}  {etiqueta}", anchor="w",
                height=26, corner_radius=8, fg_color=COLOR_HOVER,
                hover_color=COLOR_ACENTO_HOVER, text_color=COLOR_ACENTO,
                font=("Segoe UI", 11),
                command=lambda d=destino: self._abrir_ubicacion(d))
            btn.pack(padx=14, pady=(4, 10 if ultimo else 0), anchor="w")

    def _abrir_ubicacion(self, destino):
        """Abre en el Explorador la ubicación indicada (o la URL en el navegador).

        Si es un archivo, abre su carpeta con el archivo seleccionado; si es una
        carpeta, la abre directamente. Si no existe, intenta abrir la carpeta
        contenedora más cercana y avisa al usuario.
        """
        destino = (destino or "").strip().strip('"').strip("'").strip("`")
        if not destino:
            return
        if self._es_url(destino):
            import webbrowser
            try:
                webbrowser.open(destino)
            except Exception as e:
                self._aviso_input(f"No pude abrir el enlace: {e}")
            return
        import subprocess
        try:
            ruta = os.path.normpath(
                os.path.expandvars(os.path.expanduser(destino)))
            if not os.path.isabs(ruta):
                ruta = os.path.abspath(ruta)
        except Exception:
            ruta = destino
        try:
            if os.path.isfile(ruta):
                subprocess.Popen(f'explorer /select,"{ruta}"')
                return
            if os.path.isdir(ruta):
                os.startfile(ruta)
                return
            # No existe: abro la carpeta contenedora más cercana que sí exista
            carpeta = os.path.dirname(ruta)
            while carpeta and not os.path.isdir(carpeta):
                padre = os.path.dirname(carpeta)
                if padre == carpeta:
                    break
                carpeta = padre
            if carpeta and os.path.isdir(carpeta):
                os.startfile(carpeta)
                self._aviso_input(
                    f"No encontré «{os.path.basename(ruta)}»; abrí {carpeta}")
            else:
                self._aviso_input(f"No encontré la ubicación: {destino}")
        except Exception as e:
            self._aviso_input(f"No pude abrir la ubicación: {e}")

    def _toggle_seleccion(self, registro):
        """Marca o desmarca una burbuja como seleccionada."""
        if registro in self._seleccion:
            self._seleccion.remove(registro)
            registro["seleccionada"] = False
            try:
                registro["burbuja"].configure(border_width=0)
            except Exception:
                pass
        else:
            self._seleccion.append(registro)
            registro["seleccionada"] = True
            try:
                registro["burbuja"].configure(border_width=2,
                                              border_color=COLOR_ACENTO)
            except Exception:
                pass
        self._actualizar_barra_seleccion()

    def _actualizar_barra_seleccion(self):
        """Muestra u oculta la barra de selección según los mensajes marcados."""
        if not hasattr(self, "barra_seleccion"):
            return
        n = len(self._seleccion)
        if n and self.chat_activo:
            plural = "s" if n != 1 else ""
            self.lbl_sel.configure(
                text=f"{n} mensaje{plural} seleccionado{plural}")
            try:
                self.barra_seleccion.pack(fill="x", padx=110, pady=(0, 6),
                                          before=self.barra_input)
            except Exception:
                self.barra_seleccion.pack(fill="x", padx=110, pady=(0, 6))
        else:
            self.barra_seleccion.pack_forget()

    def _seleccionar_todos(self):
        self._seleccion = list(self._burbujas)
        for registro in self._burbujas:
            registro["seleccionada"] = True
            try:
                registro["burbuja"].configure(border_width=2,
                                              border_color=COLOR_ACENTO)
            except Exception:
                pass
        self._actualizar_barra_seleccion()

    def _limpiar_seleccion(self):
        for registro in self._seleccion:
            registro["seleccionada"] = False
            try:
                registro["burbuja"].configure(border_width=0)
            except Exception:
                pass
        self._seleccion = []
        self._actualizar_barra_seleccion()

    def _formatear_seleccion(self):
        """Devuelve el texto de los mensajes seleccionados listo para exportar."""
        lineas = []
        for registro in self._seleccion:
            texto = (registro.get("texto") or "").strip()
            if texto:
                lineas.append(f"{registro.get('autor', '')}: {texto}")
        return "\n\n".join(lineas)

    def _feedback_seleccion(self, mensaje):
        """Muestra un mensaje temporal en la barra de selección."""
        if not hasattr(self, "lbl_sel"):
            return
        self.lbl_sel.configure(text=mensaje)
        self.after(1600, self._actualizar_barra_seleccion)

    def _copiar_seleccion(self):
        texto = self._formatear_seleccion()
        if not texto:
            self._feedback_seleccion("No hay texto que copiar")
            return
        self.clipboard_clear()
        self.clipboard_append(texto)
        self._feedback_seleccion("Copiado al portapapeles")

    def _guardar_seleccion_txt(self):
        contenido = self._formatear_seleccion()
        if not contenido:
            self._feedback_seleccion("No hay texto que guardar")
            return
        encabezado = ("Conversación de Jarvis\n"
                      f"Fecha: {datetime.now().strftime('%Y-%m-%d %H:%M')}\n"
                      + "-" * 40 + "\n\n")
        ruta = self.crear_documento_txt(encabezado + contenido + "\n")
        if ruta:
            self._feedback_seleccion("Documento guardado")
            self._agregar_chat("Jarvis", f"Documento .txt creado en:\n{ruta}")
        else:
            self._feedback_seleccion("Guardado cancelado")

    # ---------------- Documentos de texto ----------------
    def crear_documento_txt(self, contenido, ruta=None):
        """Crea un documento .txt con el contenido indicado.

        Si no se indica 'ruta', pide al usuario dónde guardarlo (por defecto en
        la carpeta de documentos de Jarvis con un nombre basado en la fecha y
        hora). Devuelve la ruta del archivo creado, o None si cancela o falla.
        """
        contenido = (contenido or "").strip()
        if not contenido:
            return None
        if not ruta:
            try:
                os.makedirs(CARPETA_DOCUMENTOS, exist_ok=True)
            except Exception:
                pass
            nombre_def = f"documento_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt"
            ruta = filedialog.asksaveasfilename(
                title="Guardar documento de texto",
                initialdir=os.path.abspath(CARPETA_DOCUMENTOS),
                initialfile=nombre_def,
                defaultextension=".txt",
                filetypes=[("Documento de texto", "*.txt"),
                           ("Todos los archivos", "*.*")],
            )
            if not ruta:
                return None
        if not ruta.lower().endswith(".txt"):
            ruta += ".txt"
        try:
            os.makedirs(os.path.dirname(os.path.abspath(ruta)), exist_ok=True)
            with open(ruta, "w", encoding="utf-8") as f:
                f.write(contenido + "\n")
            return ruta
        except Exception as e:
            print(f"Error creando documento .txt: {e}")
            return None

    # ---------------- Gestión de adjuntos ----------------
    def _crear_ctkimage(self, ruta, max_lado=220):
        """Devuelve un CTkImage escalado conservando la proporción, o None si falla."""
        try:
            with Image.open(ruta) as im:
                im = im.convert("RGBA")
                im.thumbnail((max_lado, max_lado))
                return ctk.CTkImage(light_image=im, dark_image=im, size=im.size)
        except Exception as e:
            print(f"No se pudo cargar la imagen {ruta}: {e}")
            return None

    def _adjuntar_archivos(self):
        rutas = filedialog.askopenfilenames(
            title="Adjuntar archivos",
            filetypes=[
                ("Todos los archivos", "*.*"),
                ("Imágenes", "*.png *.jpg *.jpeg *.gif *.bmp *.webp *.tiff"),
                ("Documentos", "*.pdf *.docx *.txt *.md *.csv"),
                ("Hojas de cálculo", "*.xlsx *.xlsm"),
            ])
        for ruta in rutas:
            self._agregar_adjunto(ruta)

    def _agregar_adjunto(self, ruta):
        ruta = os.path.abspath(ruta)
        if not os.path.isfile(ruta):
            return
        if any(a["ruta"] == ruta for a in self.adjuntos):
            return
        ext = os.path.splitext(ruta)[1].lower()
        tipo = "imagen" if ext in EXT_IMAGEN else "archivo"
        self.adjuntos.append({"ruta": ruta, "nombre": os.path.basename(ruta), "tipo": tipo})
        self._actualizar_adjuntos_ui()

    def _quitar_adjunto(self, ruta):
        ruta = os.path.abspath(ruta)
        self.adjuntos = [a for a in self.adjuntos if a["ruta"] != ruta]
        self._actualizar_adjuntos_ui()

    def _limpiar_adjuntos(self):
        self.adjuntos = []
        self._actualizar_adjuntos_ui()

    def _actualizar_adjuntos_ui(self):
        for w in self.adjuntos_frame.winfo_children():
            w.destroy()
        # La fila de chips solo reserva espacio si hay adjuntos: un CTkFrame vacío
        # mantiene 200 px de alto y cortaría el área visible del chat.
        if not self.adjuntos:
            self.adjuntos_frame.pack_forget()
            return
        try:
            self.adjuntos_frame.pack(fill="x", before=self.caja_input)
        except Exception:
            self.adjuntos_frame.pack(fill="x")
        for adj in self.adjuntos:
            chip = ctk.CTkFrame(self.adjuntos_frame, fg_color=COLOR_TARJETA,
                                corner_radius=12, border_width=1, border_color=COLOR_BORDE)
            chip.pack(side="left", padx=(0, 8), pady=(0, 8))

            if adj["tipo"] == "imagen":
                img = self._crear_ctkimage(adj["ruta"], max_lado=26)
                if img is not None:
                    lbl_mini = ctk.CTkLabel(chip, text="", image=img)
                    lbl_mini._ctk_image = img
                    lbl_mini.pack(side="left", padx=(6, 4), pady=5)

            icono = "" if adj["tipo"] == "imagen" else "\U0001F4C4 "
            nombre = adj["nombre"]
            if len(nombre) > 22:
                nombre = nombre[:19] + "..."
            ctk.CTkLabel(chip, text=f"{icono}{nombre}", font=("Segoe UI", 11),
                         text_color=COLOR_TEXTO).pack(side="left", padx=(6, 4), pady=5)

            ctk.CTkButton(chip, text="\u2715", width=20, height=20, corner_radius=10,
                          fg_color="transparent", hover_color=COLOR_ACENTO_HOVER,
                          text_color=COLOR_TEXTO_SUAVE, font=("Segoe UI", 11),
                          command=lambda r=adj["ruta"]: self._quitar_adjunto(r)).pack(
                              side="left", padx=(0, 6), pady=5)

    def _pegar_portapapeles(self, event=None):
        """Procesa el pegado (Ctrl+V).

        - Imágenes y archivos copiados se adjuntan a la conversación.
        - El texto corto se pega con normalidad en el campo de entrada.
        - El texto muy largo se guarda como documento adjunto (.txt): insertar
          miles de caracteres en el Entry de Tk es O(n²) y congelaba la interfaz.
        """
        # 1) Archivos o imágenes del portapapeles (vía Pillow)
        datos = None
        try:
            from PIL import ImageGrab
            datos = ImageGrab.grabclipboard()
        except Exception:
            datos = None

        # Lista de rutas de archivos copiados desde el Explorador
        if isinstance(datos, list):
            agregado = False
            for ruta in datos:
                if os.path.isfile(ruta):
                    self._agregar_adjunto(ruta)
                    agregado = True
            if agregado:
                return "break"
            datos = None

        # Imagen (captura de pantalla o copiada al portapapeles)
        if datos is not None and not isinstance(datos, str):
            try:
                os.makedirs(CARPETA_ADJUNTOS, exist_ok=True)
                nombre = f"pegado_{datetime.now().strftime('%Y%m%d_%H%M%S')}.png"
                ruta = os.path.join(CARPETA_ADJUNTOS, nombre)
                datos.save(ruta, "PNG")
                self._agregar_adjunto(ruta)
                return "break"
            except Exception as e:
                print(f"No se pudo guardar la imagen del portapapeles: {e}")
                datos = None

        # 2) Texto del portapapeles
        texto = datos if isinstance(datos, str) else None
        if texto is None:
            try:
                texto = self.clipboard_get()
            except Exception:
                texto = None  # no hay texto: el pegado normal decide

        if texto and len(texto) > UMBRAL_PEGADO_TEXTO:
            ruta = self._guardar_texto_pegado(texto)
            if ruta:
                self._agregar_adjunto(ruta)
                self._aviso_input(
                    f"Texto largo ({len(texto):,} caracteres) adjuntado como "
                    f"{os.path.basename(ruta)}")
                print(f"Texto pegado adjuntado como documento: {ruta}")
                return "break"
        return None  # texto corto: continúa el pegado normal del campo

    def _guardar_texto_pegado(self, texto):
        """Guarda un texto pegado como archivo .txt dentro de la carpeta de adjuntos."""
        try:
            os.makedirs(CARPETA_ADJUNTOS, exist_ok=True)
            nombre = f"texto_pegado_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt"
            ruta = os.path.join(CARPETA_ADJUNTOS, nombre)
            with open(ruta, "w", encoding="utf-8") as f:
                f.write(texto)
            return ruta
        except Exception as e:
            print(f"No se pudo guardar el texto pegado: {e}")
            return None

    def _aviso_input(self, mensaje):
        """Muestra un aviso temporal bajo el campo de entrada (se oculta solo)."""
        self._ocultar_aviso()
        try:
            self._lbl_aviso = ctk.CTkLabel(self.barra_input, text=mensaje,
                                           font=("Segoe UI", 11),
                                           text_color=COLOR_ACENTO)
            self._lbl_aviso.pack(pady=(6, 0))
            lbl = self._lbl_aviso
            self.after(3500, lambda: self._ocultar_aviso(lbl))
        except Exception:
            self._lbl_aviso = None

    def _ocultar_aviso(self, lbl=None):
        """Elimina el aviso temporal; si se indica 'lbl' solo si sigue siendo el actual."""
        actual = getattr(self, "_lbl_aviso", None)
        if actual is None:
            return
        if lbl is not None and lbl is not actual:
            return
        try:
            actual.destroy()
        except Exception:
            pass
        self._lbl_aviso = None

    # ---------------- Extracción de contenido de los adjuntos ----------------
    def _extraer_contenido(self, ruta):
        """Devuelve el texto/contenido extraído de un archivo para enviarlo a la IA."""
        ext = os.path.splitext(ruta)[1].lower()
        nombre = os.path.basename(ruta)
        try:
            if ext in EXT_IMAGEN:
                return self._extraer_imagen(ruta)
            if ext in EXT_TEXTO:
                with open(ruta, "r", encoding="utf-8", errors="ignore") as f:
                    return f.read()[:MAX_TEXTO_ADJUNTO]
            if ext == ".pdf":
                return self._extraer_pdf(ruta)
            if ext == ".docx":
                return self._extraer_docx(ruta)
            if ext in (".xlsx", ".xlsm"):
                return self._extraer_xlsx(ruta)
            if ext in (".pptx", ".pptm", ".potx"):
                return self._extraer_pptx(ruta)
            if ext in (".ppt", ".doc", ".xls"):
                return (f"[Archivo de Office en formato antiguo ({ext}); no se puede "
                        f"extraer su texto. Indica al usuario que lo guarde como "
                        f".pptx/.docx/.xlsx o en PDF.]")
            # Intento final: leerlo como texto plano
            with open(ruta, "r", encoding="utf-8", errors="strict") as f:
                return f.read()[:MAX_TEXTO_ADJUNTO]
        except Exception:
            pass
        try:
            tam = os.path.getsize(ruta)
        except Exception:
            tam = 0
        return f"[Archivo binario no legible como texto: {nombre} ({tam} bytes)]"

    def _extraer_imagen(self, ruta):
        nombre = os.path.basename(ruta)
        info = f"[Imagen: {nombre}]"
        try:
            with Image.open(ruta) as im:
                info = f"[Imagen: {nombre}, formato {im.format}, {im.width}x{im.height}px]"
        except Exception:
            pass
        try:
            texto_ocr = ocr_imagen(ruta) if ocr_motor() else ""
        except Exception as e:
            print(f"No se pudo aplicar OCR a {nombre}: {e}")
            texto_ocr = ""
        if texto_ocr:
            return f"{info}\nTexto detectado en la imagen (OCR):\n{texto_ocr[:MAX_TEXTO_ADJUNTO]}"
        return (f"{info}\n(No se pudo extraer texto de esta imagen. "
                f"El archivo está en: {ruta})")

    def _extraer_pdf(self, ruta):
        try:
            from pypdf import PdfReader
        except Exception:
            try:
                from PyPDF2 import PdfReader
            except Exception:
                return (f"[PDF adjunto: {os.path.basename(ruta)}. "
                        "Instala 'pypdf' para extraer su texto.]")
        try:
            lector = PdfReader(ruta)
            partes = [(pag.extract_text() or "") for pag in lector.pages[:30]]
            return "\n".join(partes)[:MAX_TEXTO_ADJUNTO] or "[PDF sin texto extraíble]"
        except Exception as e:
            return f"[No se pudo leer el PDF: {e}]"

    def _extraer_docx(self, ruta):
        try:
            import docx
        except Exception:
            return (f"[Documento Word: {os.path.basename(ruta)}. "
                    "Instala 'python-docx' para extraer su texto.]")
        try:
            doc = docx.Document(ruta)
            return "\n".join(p.text for p in doc.paragraphs)[:MAX_TEXTO_ADJUNTO]
        except Exception as e:
            return f"[No se pudo leer el documento Word: {e}]"

    def _extraer_xlsx(self, ruta):
        try:
            import openpyxl
        except Exception:
            return (f"[Hoja de cálculo: {os.path.basename(ruta)}. "
                    "Instala 'openpyxl' para extraer su contenido.]")
        try:
            wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
            partes = []
            for hoja in wb.worksheets:
                partes.append(f"# Hoja: {hoja.title}")
                for i, fila in enumerate(hoja.iter_rows(values_only=True)):
                    if i > 300:
                        partes.append("... (truncado)")
                        break
                    partes.append(" | ".join("" if c is None else str(c) for c in fila))
            return "\n".join(partes)[:MAX_TEXTO_ADJUNTO]
        except Exception as e:
            return f"[No se pudo leer la hoja de cálculo: {e}]"

    def _extraer_pptx(self, ruta):
        try:
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE
        except Exception:
            return (f"[Presentación: {os.path.basename(ruta)}. "
                    "Instala 'python-pptx' para extraer su contenido.]")
        try:
            prs = Presentation(ruta)
        except Exception as e:
            return f"[No se pudo leer la presentación: {e}]"

        def _texto_formas(formas):
            """Recorre las formas (incluidas las agrupadas) y devuelve sus textos."""
            lineas = []
            for forma in formas:
                try:
                    if forma.shape_type == MSO_SHAPE_TYPE.GROUP:
                        lineas.extend(_texto_formas(forma.shapes))
                        continue
                except Exception:
                    pass
                if getattr(forma, "has_table", False) and forma.has_table:
                    for fila in forma.table.rows:
                        celdas = [c.text.strip() for c in fila.cells]
                        if any(celdas):
                            lineas.append(" | ".join(celdas))
                    continue
                try:
                    if forma.has_text_frame:
                        txt = forma.text_frame.text.strip()
                        if txt:
                            lineas.append(txt)
                except Exception:
                    pass
            return lineas

        partes = []
        for i, slide in enumerate(prs.slides, start=1):
            if i > 60:
                partes.append("... (se omitieron las diapositivas restantes)")
                break
            lineas = _texto_formas(slide.shapes)
            try:
                notas = slide.notes_slide.notes_text_frame.text.strip()
            except Exception:
                notas = ""
            bloque = f"# Diapositiva {i}"
            if lineas:
                bloque += "\n" + "\n".join(lineas)
            if notas:
                bloque += f"\n[Notas: {notas}]"
            partes.append(bloque)
        contenido = "\n\n".join(partes).strip()
        return contenido[:MAX_TEXTO_ADJUNTO] or "[Presentación sin texto extraíble]"

    def _construir_prompt(self, texto, adjuntos):
        """Combina el texto del usuario con el contenido extraído de los adjuntos."""
        if not adjuntos:
            return texto
        bloques = [f"--- Archivo adjunto: {a['nombre']} ---\n{self._extraer_contenido(a['ruta'])}"
                   for a in adjuntos]
        encabezado = texto.strip() or "Analiza el/los archivo(s) adjunto(s)."
        return (f"{encabezado}\n\nEl usuario adjuntó los siguientes archivos; "
                "ten en cuenta su contenido al responder:\n\n" + "\n\n".join(bloques))

    def _guardar_regla_si_aplica(self, texto):
        """Guarda una regla permanente si el mensaje la declara de forma explícita.

        Se reconocen los prefijos 'regla:' y '/regla' (p. ej. 'regla: las tareas
        van al padre 174080'). Devuelve True si se guardó una regla.
        """
        t = (texto or "").strip()
        bajo = t.lower()
        if bajo.startswith("/regla"):
            regla = t[len("/regla"):].lstrip(": ").strip()
        elif bajo.startswith("regla:"):
            regla = t[len("regla:"):].strip()
        else:
            return False
        if not regla:
            return False
        self.system.productivity.agregar_regla(regla)
        print(f"Regla permanente guardada: {regla}")
        return True

    def _scroll_final(self):
        """Desplaza la conversación al final recalculando antes la región de scroll.

        Recalcular la región evita que, tras insertar muchas burbujas de forma
        síncrona (p. ej. al abrir una conversación), el scroll use una geometría
        obsoleta y deje los mensajes fuera de la vista (chat en blanco)."""
        try:
            canvas = self.chat_scroll._parent_canvas
            canvas.update_idletasks()
            canvas.configure(scrollregion=canvas.bbox("all"))
            canvas.yview_moveto(1.0)
            self._pegado_abajo = True
        except Exception:
            pass

    def _scroll_actualizado(self, first, last):
        """Registra si el final del chat está a la vista (lo llama el lienzo).

        Sustituye al yscrollcommand de CTkScrollableFrame conservando su trabajo
        (mover la barra lateral) y recuerda la posición para reanclar el final
        cuando cambia el tamaño de la vista."""
        try:
            self._pegado_abajo = float(last) >= 0.999
        except Exception:
            self._pegado_abajo = True
        try:
            self.chat_scroll._scrollbar.set(first, last)
        except Exception:
            pass

    def _chat_redimensionado(self, event=None):
        """Reancla el final de la conversación cuando cambia el tamaño de la vista.

        Sin esto, al encoger o maximizar la ventana el lienzo cambia de alto y los
        últimos mensajes quedan por debajo del borde visible (el chat parece
        cortado). Solo se reancla si el usuario estaba viendo el final, para no
        arrastrar la vista mientras lee mensajes anteriores."""
        if not self.chat_activo or not self._pegado_abajo:
            return
        # En diferido: no conviene reconfigurar el lienzo dentro de su <Configure>
        # (llegan varios eventos seguidos hasta que la ventana tiene su tamaño final).
        self.after_idle(self._scroll_final)

    def _mostrar_pensando(self, activo):
        self._procesando = activo
        if activo:
            if self._ind_pensando is None:
                self._ind_pensando = ctk.CTkFrame(self.chat_scroll,
                                                  fg_color=COLOR_TARJETA, corner_radius=14)
                self._lbl_pensando = ctk.CTkLabel(self._ind_pensando,
                                                  text="Jarvis está escribiendo",
                                                  font=("Segoe UI", 12, "italic"),
                                                  text_color=COLOR_TEXTO_SUAVE)
                self._lbl_pensando.pack(padx=16, pady=10)
            self._ind_pensando.pack(anchor="w", padx=18, pady=5)
            self.after(30, self._scroll_final)
            self._animar_pensando(0)
        else:
            if self._ind_pensando is not None:
                self._ind_pensando.pack_forget()

    def _animar_pensando(self, i):
        if not self._procesando or self._ind_pensando is None:
            return
        self._lbl_pensando.configure(text="Jarvis está escribiendo" + "." * (i % 4))
        self.after(400, lambda: self._animar_pensando(i + 1))

    def _nueva_conversacion(self):
        """Inicia un hilo nuevo: el próximo mensaje creará su propio archivo."""
        for w in self.chat_scroll.winfo_children():
            w.destroy()
        self._ind_pensando = None
        self._pegado_abajo = True    # un hilo nuevo se ve desde el final
        self._procesando = False
        self.chat_activo = None
        self._burbujas = []
        self._seleccion = []
        self.conversacion_actual = None
        self._limpiar_adjuntos()
        self._mostrar_vista("chat")
        self._mostrar_estado_chat(False)

    # ---------------- Envío y procesamiento ----------------
    def _enviar_chat(self):
        texto = self.entry_chat.get().strip()
        adjuntos = list(self.adjuntos)
        if not texto and not adjuntos:
            return
        self.entry_chat.delete(0, "end")
        self._limpiar_adjuntos()
        self._agregar_chat("Tú", texto, adjuntos)
        self.after(0, lambda: self._mostrar_pensando(True))
        threading.Thread(target=self._procesar_chat, args=(texto, adjuntos),
                         daemon=True).start()

    def _procesar_chat(self, texto, adjuntos=None):
        adjuntos = adjuntos or []
        try:
            self._guardar_regla_si_aplica(texto)
            prompt = self._construir_prompt(texto, adjuntos)
            respuesta = self._procesar_mensaje(prompt)
        except Exception as e:
            respuesta = f"Ocurrió un error al procesar tu mensaje: {e}"
        finally:
            self.after(0, lambda: self._mostrar_pensando(False))
        if respuesta:
            self._agregar_chat("Jarvis", respuesta)

    def _procesar_mensaje(self, texto):
        """Procesa un mensaje de usuario (voz o texto) y devuelve la respuesta final."""
        respuesta_texto, comando_accion = self.brain.procesar_mensaje(texto)
        resultado = None
        if comando_accion:
            resultado = self.system.ejecutar_comando(comando_accion)
            if comando_accion.strip().upper() == "SISTEMA|APAGAR":
                self.reproducir_voz("Entendido, cerrando el programa. Hasta luego.")
                self.after(0, self.destroy)
                return "Cerrando el programa..."
        if resultado:
            if str(resultado).startswith("__CONTEXTO__"):
                contexto = str(resultado)[len("__CONTEXTO__"):]
                return self.brain.analizar(contexto) or "No pude procesar la solicitud."
            return str(resultado)
        if respuesta_texto:
            return respuesta_texto
        if comando_accion:
            return ""  # se ejecutó una acción sin texto de respuesta
        return "No entendí eso."

    # ---------------- Vista de ajustes ----------------
    def _crear_vista_ajustes(self, parent):
        cabecera = ctk.CTkFrame(parent, fg_color="transparent")
        cabecera.pack(fill="x", padx=30, pady=(22, 14))
        ctk.CTkButton(cabecera, text="\u2039  Volver al chat", width=150, height=32,
                      fg_color="transparent", hover_color=COLOR_HOVER,
                      text_color=COLOR_TEXTO_SUAVE, anchor="w", font=("Segoe UI", 13),
                      command=lambda: self._mostrar_vista("chat")).pack(side="left")
        ctk.CTkLabel(cabecera, text="Voz y Ajustes", font=("Segoe UI", 20, "bold"),
                     text_color=COLOR_TEXTO).pack(side="left", padx=18)

        # Indicador de estado
        self.lbl_estado = ctk.CTkLabel(parent, text="Estado: Inactivo",
                                       text_color="#e06c6c", font=("Segoe UI", 14))
        self.lbl_estado.pack(pady=(6, 6))
        self.btn_toggle = ctk.CTkButton(parent, text="Iniciar Escucha",
                                        command=self.toggle_escucha, fg_color=COLOR_ACENTO,
                                        hover_color=COLOR_ACENTO_HOVER, text_color="#ffffff",
                                        height=40, width=300, font=("Segoe UI", 14, "bold"))
        self.btn_toggle.pack(pady=(0, 18))

        # Marco de ajustes
        self.frame_config = ctk.CTkFrame(parent, fg_color=COLOR_TARJETA, corner_radius=12)
        self.frame_config.pack(fill="both", expand=True, padx=30, pady=(0, 26))

        self._campo_config(self.frame_config, "DeepSeek API Key:",
                           "entry_api", show="*", pady_top=(18, 4))
        self._campo_config(self.frame_config, "Nombre del Asistente (Wake Word):",
                           "entry_nombre")
        self._campo_config(self.frame_config, "Ruta Modelo RVC (.pth):", "entry_rvc")

        # Matriz en la nube: botones, cada uno con su propio diálogo de configuración
        ctk.CTkLabel(self.frame_config, text="Matriz en la nube (Excel + Azure DevOps):",
                     text_color=COLOR_TEXTO, font=("Segoe UI", 12)).pack(
                         anchor="w", padx=18, pady=(14, 4))
        frame_nube = ctk.CTkFrame(self.frame_config, fg_color="transparent")
        frame_nube.pack(fill="x", padx=18, pady=(0, 6))
        self.btn_cfg_ado = ctk.CTkButton(
            frame_nube, text="Configurar Azure DevOps",
            command=self._abrir_config_azure, fg_color=COLOR_ACENTO,
            hover_color=COLOR_ACENTO_HOVER, text_color="#ffffff", height=36,
            font=("Segoe UI", 12, "bold"))
        self.btn_cfg_ado.pack(side="left", expand=True, fill="x", padx=(0, 6))
        self.btn_cfg_excel = ctk.CTkButton(
            frame_nube, text="Excel (archivo o URL)",
            command=self._abrir_config_excel, fg_color=COLOR_ACENTO,
            hover_color=COLOR_ACENTO_HOVER, text_color="#ffffff", height=36,
            font=("Segoe UI", 12, "bold"))
        self.btn_cfg_excel.pack(side="left", expand=True, fill="x", padx=(6, 0))

        self.lbl_nube_excel = ctk.CTkLabel(self.frame_config,
                                           text="Excel: sin configurar",
                                           text_color=COLOR_TEXTO_SUAVE,
                                           font=("Segoe UI", 11), anchor="w",
                                           justify="left", wraplength=620)
        self.lbl_nube_excel.pack(anchor="w", padx=18, pady=(4, 2))
        self.lbl_nube_ado = ctk.CTkLabel(self.frame_config,
                                         text="Azure DevOps: sin configurar",
                                         text_color=COLOR_TEXTO_SUAVE,
                                         font=("Segoe UI", 11), anchor="w",
                                         justify="left", wraplength=620)
        self.lbl_nube_ado.pack(anchor="w", padx=18, pady=(0, 10))
        # Correo (Outlook nuevo / Microsoft 365): se conecta una sola vez con un
        # código de Microsoft; no hace falta contraseña ni Outlook clásico.
        ctk.CTkLabel(self.frame_config,
                     text="Correo (Outlook nuevo / Microsoft 365):",
                     text_color=COLOR_TEXTO, font=("Segoe UI", 12)).pack(
                         anchor="w", padx=18, pady=(14, 4))
        frame_correo = ctk.CTkFrame(self.frame_config, fg_color="transparent")
        frame_correo.pack(fill="x", padx=18, pady=(0, 6))
        self.btn_conectar_correo = ctk.CTkButton(
            frame_correo, text="Conectar correo", command=self._conectar_correo,
            fg_color=COLOR_ACENTO, hover_color=COLOR_ACENTO_HOVER, text_color="#ffffff",
            height=36, font=("Segoe UI", 12, "bold"))
        self.btn_conectar_correo.pack(side="left", expand=True, fill="x", padx=(0, 6))
        self.btn_desconectar_correo = ctk.CTkButton(
            frame_correo, text="Desconectar", command=self._desconectar_correo,
            fg_color=COLOR_PANEL, hover_color=COLOR_HOVER, text_color=COLOR_TEXTO,
            border_width=1, border_color=COLOR_BORDE, height=36,
            font=("Segoe UI", 12, "bold"))
        self.btn_desconectar_correo.pack(side="left", expand=True, fill="x", padx=(6, 0))

        self.lbl_correo = ctk.CTkLabel(self.frame_config, text="Correo: sin conectar",
                                       text_color=COLOR_TEXTO_SUAVE,
                                       font=("Segoe UI", 11), anchor="w",
                                       justify="left", wraplength=620)
        self.lbl_correo.pack(anchor="w", padx=18, pady=(4, 10))


        self.btn_guardar = ctk.CTkButton(self.frame_config, text="Guardar configuración",
                                         command=self.guardar_config, fg_color="#2f9e6b",
                                         hover_color="#26895d", text_color="#ffffff",
                                         height=38, font=("Segoe UI", 13, "bold"))
        self.btn_guardar.pack(pady=(18, 8))

        # Reglas permanentes (memoria del asistente: se añaden a cada consulta)
        ctk.CTkLabel(self.frame_config,
                     text="Reglas permanentes (Jarvis las recuerda siempre):",
                     text_color=COLOR_TEXTO, font=("Segoe UI", 12)).pack(
                         anchor="w", padx=18, pady=(10, 4))
        self.txt_reglas = ctk.CTkTextbox(self.frame_config, height=110,
                                         fg_color=COLOR_PANEL, border_width=1,
                                         border_color=COLOR_BORDE, text_color=COLOR_TEXTO,
                                         font=("Segoe UI", 12), wrap="word")
        self.txt_reglas.pack(padx=18, fill="x")
        self.btn_guardar_reglas = ctk.CTkButton(
            self.frame_config, text="Guardar reglas", command=self._guardar_reglas_ui,
            fg_color=COLOR_ACENTO, hover_color=COLOR_ACENTO_HOVER, text_color="#ffffff",
            height=34, font=("Segoe UI", 12, "bold"))
        self.btn_guardar_reglas.pack(pady=(10, 18))

    def _campo_config(self, parent, etiqueta, nombre_attr, show=None, pady_top=(14, 4)):
        ctk.CTkLabel(parent, text=etiqueta, text_color=COLOR_TEXTO,
                     font=("Segoe UI", 12)).pack(anchor="w", padx=18, pady=pady_top)
        entrada = ctk.CTkEntry(parent, show=show, fg_color=COLOR_PANEL,
                               border_color=COLOR_BORDE, text_color=COLOR_TEXTO, height=36)
        entrada.pack(padx=18, pady=(0, 10), fill="x")
        setattr(self, nombre_attr, entrada)

    # ---------------- Matriz en la nube: configuración con botones ----------------
    def _leer_config_json(self):
        """Devuelve el diccionario completo de config_usuario.json."""
        if os.path.exists(self.config_path):
            try:
                with open(self.config_path, "r", encoding="utf-8") as f:
                    data = json.load(f)
                if isinstance(data, dict):
                    return data
            except Exception as e:
                print(f"Error leyendo configuración: {e}")
        return {}

    def _escribir_config_json(self, data):
        """Escribe config_usuario.json (creando la carpeta si hace falta)."""
        carpeta = os.path.dirname(self.config_path)
        if carpeta:
            os.makedirs(carpeta, exist_ok=True)
        with open(self.config_path, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=4)

    def _actualizar_estado_nube(self):
        """Refresca las etiquetas que resumen la configuración de la matriz en la nube."""
        cfg = self._leer_config_json()
        matriz = cfg.get("matriz") if isinstance(cfg.get("matriz"), dict) else {}
        azure = cfg.get("azure_devops") if isinstance(cfg.get("azure_devops"), dict) else {}

        archivo = (matriz.get("excel_path") or "").strip()
        url = (matriz.get("excel_url") or "").strip()
        sincronizada = (matriz.get("excel_sync_dir") or "").strip()
        if archivo:
            texto_excel = f"Excel: {os.path.basename(archivo)}\n{archivo}"
        elif url:
            texto_excel = f"Excel (URL): {url}"
            if sincronizada:
                texto_excel += (f"\ncopia en la carpeta sincronizada (se sube a "
                                f"SharePoint): {sincronizada}")
            else:
                texto_excel += ("\ncopia local: los cambios NO se suben a SharePoint "
                                "(indica la carpeta sincronizada en el botón del Excel)")
        else:
            texto_excel = ("Excel: sin configurar\n"
                           f"(se usará {os.path.join(rutas.MATRICES, 'Matriz.xlsx')})")

        org = (azure.get("org") or "").strip()
        proyecto = (azure.get("project") or "").strip()
        if org or proyecto:
            token = ((azure.get("pat") or "").strip()
                     or os.environ.get("AZURE_DEVOPS_PAT", "").strip())
            texto_ado = (f"Azure DevOps: {org or '(sin organización)'} · "
                         f"{proyecto or '(sin proyecto)'} · "
                         f"{'token OK' if token else 'falta el token (PAT)'}")
        else:
            texto_ado = "Azure DevOps: sin configurar"

        if hasattr(self, "lbl_nube_excel"):
            self.lbl_nube_excel.configure(text=texto_excel)
        if hasattr(self, "lbl_nube_ado"):
            self.lbl_nube_ado.configure(text=texto_ado)
    def _actualizar_estado_correo(self):
        """Refresca la etiqueta que resume la conexión del correo."""
        cfg = self._leer_config_json()
        correo = cfg.get("correo") if isinstance(cfg.get("correo"), dict) else {}
        cuenta = (correo.get("cuenta") or "").strip()
        conectando = getattr(self, "_conectando_correo", False)
        if correo.get("token"):
            texto = f"Correo conectado: {cuenta}" if cuenta else "Correo conectado."
            if not cuenta:
                texto += " (pulsa 'Conectar correo' para comprobar la cuenta)"
            texto += ("\nJarvis puede leerlo, resumirlo y dejar los borradores "
                      "preparados; nunca envía nada por su cuenta.")
            color = "#2ecc71"
        elif conectando:
            return                      # hay una conexión en curso: no pisamos el código
        else:
            texto = ("Correo: sin conectar\nSe conecta una sola vez con un código de "
                     "Microsoft (el navegador se abre solo); no hace falta contraseña.")
            color = COLOR_TEXTO_SUAVE
        if hasattr(self, "lbl_correo"):
            self.lbl_correo.configure(text=texto, text_color=color)

    def _conectar_correo(self):
        """Conecta el correo con el código de dispositivo, sin bloquear la ventana."""
        if getattr(self, "_conectando_correo", False):
            return
        self._conectando_correo = True
        self.btn_conectar_correo.configure(text="Esperando el código...", state="disabled")
        self.lbl_correo.configure(text="Preparando la conexión con Microsoft...",
                                  text_color=COLOR_TEXTO)
        estado = {"abierto": False}

        def avisar(texto):
            """Enseña el código (acumulando líneas) y abre el navegador una vez."""
            if not estado["abierto"] and re.search(r"código\s+[A-Z0-9][A-Z0-9-]{4,}",
                                                   texto, re.IGNORECASE):
                estado["abierto"] = True
                try:
                    webbrowser.open("https://microsoft.com/devicelogin")
                except Exception as e:
                    print(f"No pude abrir el navegador: {e}")

            def pintar():
                # El código y el tiempo llegan en dos avisos seguidos: se acumulan
                # para que el código no desaparezca de la ventana.
                previo = (self.lbl_correo.cget("text") or "").strip()
                if previo.startswith("Preparando"):
                    previo = ""
                self.lbl_correo.configure(text=(f"{previo}\n{texto}" if previo else texto),
                                          text_color=COLOR_TEXTO)

            self.after(0, pintar)

        def tarea():
            from modules.correo_nube import CorreoNube
            try:
                ok, mensaje = CorreoNube(self.config_path).entrar(avisar=avisar)
            except Exception as e:
                ok, mensaje = False, f"No pude conectar tu correo: {e}"
            self.after(0, lambda: self._correo_resuelto(ok, mensaje))

        threading.Thread(target=tarea, daemon=True).start()

    def _correo_resuelto(self, ok, mensaje):
        """Termina la conexión del correo: pinta el resultado y rearma el botón."""
        self._conectando_correo = False
        self.btn_conectar_correo.configure(text="Conectar correo", state="normal")
        self.lbl_correo.configure(text=mensaje,
                                  text_color="#2ecc71" if ok else "#e06c6c")

    def _desconectar_correo(self):
        """Olvida la sesión de correo guardada, pidiendo confirmación."""
        cfg = self._leer_config_json()
        correo = cfg.get("correo") if isinstance(cfg.get("correo"), dict) else {}
        if not correo.get("token"):
            self.lbl_correo.configure(text="El correo todavía no está conectado.",
                                      text_color=COLOR_TEXTO_SUAVE)
            return

        dlg = ctk.CTkToplevel(self)
        dlg.title("Desconectar correo")
        dlg.geometry("470x230")
        dlg.configure(fg_color=COLOR_FONDO)
        dlg.transient(self)
        dlg.after(150, dlg.grab_set)
        ctk.CTkLabel(dlg, text="¿Olvidar la sesión de tu correo?",
                     font=("Segoe UI", 15, "bold"), text_color=COLOR_TEXTO).pack(
                         anchor="w", padx=24, pady=(22, 6))
        ctk.CTkLabel(dlg,
                     text=("Jarvis borrará el token guardado y dejará de leer tu correo.\n"
                           "Podrás volver a conectarlo cuando quieras."),
                     font=("Segoe UI", 12), text_color=COLOR_TEXTO_SUAVE,
                     justify="left").pack(anchor="w", padx=24)

        def confirmar():
            from modules.correo_nube import CorreoNube
            correo_nube = CorreoNube(self.config_path)
            correo_nube.salir()
            dlg.destroy()
            self.lbl_correo.configure(
                text="Correo desconectado. Jarvis ya no accede a tu buzón.",
                text_color=COLOR_TEXTO_SUAVE)

        botones = ctk.CTkFrame(dlg, fg_color="transparent")
        botones.pack(pady=(20, 12))
        ctk.CTkButton(botones, text="Desconectar", command=confirmar,
                      fg_color="#c0392b", hover_color="#a93226", text_color="#ffffff",
                      height=34, font=("Segoe UI", 12, "bold")).pack(side="left", padx=6)
        ctk.CTkButton(botones, text="Cancelar", command=dlg.destroy,
                      fg_color=COLOR_PANEL, hover_color=COLOR_HOVER,
                      text_color=COLOR_TEXTO, border_width=1, border_color=COLOR_BORDE,
                      height=34, font=("Segoe UI", 12)).pack(side="left", padx=6)


    def _abrir_config_azure(self):
        """Diálogo para configurar Azure DevOps (organización, proyecto, PAT...)."""
        cfg = self._leer_config_json()
        azure = cfg.get("azure_devops") if isinstance(cfg.get("azure_devops"), dict) else {}

        dlg = ctk.CTkToplevel(self)
        dlg.title("Configurar Azure DevOps")
        dlg.geometry("640x640")
        dlg.configure(fg_color=COLOR_FONDO)
        dlg.transient(self)
        dlg.after(150, dlg.grab_set)

        ctk.CTkLabel(dlg, text="Azure DevOps (work items)",
                     font=("Segoe UI", 16, "bold"),
                     text_color=COLOR_TEXTO).pack(anchor="w", padx=24, pady=(18, 2))
        ctk.CTkLabel(dlg,
                     text=("Genera un PAT en dev.azure.com → icono de usuario → "
                           "Tokens de acceso personal\n(permiso \"Work Items: Read & "
                           "Write\"). No necesitas una API key de administrador."),
                     font=("Segoe UI", 11), text_color=COLOR_TEXTO_SUAVE,
                     justify="left", wraplength=590).pack(anchor="w", padx=24,
                                                          pady=(0, 10))

        campos = [
            ("Organización (URL)", "org", "https://dev.azure.com/MiOrganizacion"),
            ("Proyecto", "project", "MiProyecto"),
            ("Token de acceso personal (PAT)", "pat", ""),
            ("Tipo de work item", "tipo_item", "Requirement"),
            ("Area Path (opcional)", "area_path", ""),
            ("Iteration Path (opcional)", "iteration_path", ""),
        ]
        entradas = {}
        for etiqueta, clave, pista in campos:
            ctk.CTkLabel(dlg, text=etiqueta, text_color=COLOR_TEXTO,
                         font=("Segoe UI", 12)).pack(anchor="w", padx=24, pady=(8, 2))
            entrada = ctk.CTkEntry(dlg, fg_color=COLOR_PANEL,
                                   border_color=COLOR_BORDE, text_color=COLOR_TEXTO,
                                   height=34, show="*" if clave == "pat" else None)
            entrada.pack(padx=24, fill="x")
            valor = str(azure.get(clave, "") or "")
            if valor:
                entrada.insert(0, valor)
            elif clave == "tipo_item":
                entrada.insert(0, "Requirement")
            elif pista:
                entrada.configure(placeholder_text=pista)
            entradas[clave] = entrada

        lbl_estado = ctk.CTkLabel(dlg, text="", text_color=COLOR_TEXTO_SUAVE,
                                  font=("Segoe UI", 11), justify="left",
                                  wraplength=590)
        lbl_estado.pack(anchor="w", padx=24, pady=(10, 0))

        def _probar():
            datos = {k: e.get().strip() for k, e in entradas.items()}
            lbl_estado.configure(text="Probando conexión...",
                                 text_color=COLOR_TEXTO_SUAVE)
            resultado = {}

            def _worker():
                try:
                    resultado["texto"] = self.system.cloud.probar_azure(
                        datos["org"], datos["project"], datos["pat"])
                except Exception as e:
                    resultado["texto"] = f"No pude conectar: {e}"

            def _revisar():
                """Se ejecuta en el hilo de la interfaz: la prueba va en un hilo aparte."""
                try:
                    if "texto" not in resultado:
                        dlg.after(250, _revisar)
                        return
                    texto = resultado["texto"]
                    color = "#2ecc71" if texto.startswith("Conexión") else "#e06c6c"
                    lbl_estado.configure(text=texto, text_color=color)
                except Exception:
                    pass

            threading.Thread(target=_worker, daemon=True).start()
            dlg.after(250, _revisar)

        def _guardar():
            datos = {k: e.get().strip() for k, e in entradas.items()}
            if not datos["org"] or not datos["project"]:
                lbl_estado.configure(text="Indica al menos la organización y el proyecto.",
                                     text_color="#e06c6c")
                return
            cfg_actual = self._leer_config_json()
            azure_cfg = cfg_actual.get("azure_devops")
            if not isinstance(azure_cfg, dict):
                azure_cfg = {}
            azure_cfg.update(datos)
            cfg_actual["azure_devops"] = azure_cfg
            try:
                self._escribir_config_json(cfg_actual)
            except Exception as e:
                lbl_estado.configure(text=f"No pude guardar: {e}", text_color="#e06c6c")
                return
            self._actualizar_estado_nube()
            dlg.destroy()

        frame_botones = ctk.CTkFrame(dlg, fg_color="transparent")
        frame_botones.pack(fill="x", padx=24, pady=16)
        ctk.CTkButton(frame_botones, text="Guardar", command=_guardar,
                      fg_color="#2f9e6b", hover_color="#26895d", text_color="#ffffff",
                      height=36, font=("Segoe UI", 12, "bold")).pack(side="right",
                                                                     padx=(8, 0))
        ctk.CTkButton(frame_botones, text="Probar conexión", command=_probar,
                      fg_color=COLOR_ACENTO, hover_color=COLOR_ACENTO_HOVER,
                      text_color="#ffffff", height=36,
                      font=("Segoe UI", 12, "bold")).pack(side="right", padx=(8, 0))
        ctk.CTkButton(frame_botones, text="Cancelar", command=dlg.destroy,
                      fg_color=COLOR_HOVER, hover_color=COLOR_BORDE,
                      text_color=COLOR_TEXTO, height=36).pack(side="right")

    def _cargar_reglas_ui(self):
        """Vuelca las reglas permanentes guardadas en el editor de la interfaz."""
        if not hasattr(self, "txt_reglas"):
            return
        try:
            reglas = self.system.productivity.leer_reglas()
        except Exception:
            reglas = ""
        self.txt_reglas.delete("1.0", "end")
        if reglas:
            self.txt_reglas.insert("1.0", reglas)

    def _guardar_reglas_ui(self):
        """Guarda el contenido del editor de reglas permanentes."""
        if not hasattr(self, "txt_reglas"):
            return
        contenido = self.txt_reglas.get("1.0", "end").strip()
        ok = self.system.productivity.guardar_reglas(contenido)
        self.btn_guardar_reglas.configure(
            text="Reglas guardadas \u2713" if ok else "Error al guardar")
        self.after(1600, lambda: self.btn_guardar_reglas.configure(
            text="Guardar reglas"))

    # ---------------- Lista de conversaciones (sidebar) ----------------
    def _cargar_conversaciones_sidebar(self):
        if not hasattr(self, "lista_conv"):
            return
        for w in self.lista_conv.winfo_children():
            w.destroy()

        archivos = sorted(glob.glob(os.path.join(rutas.CONVERSACIONES, "historial_chat_*.md")),
                          reverse=True)
        if not archivos:
            ctk.CTkLabel(self.lista_conv, text="Sin conversaciones todavía",
                         font=("Segoe UI", 11), text_color=COLOR_TEXTO_SUAVE,
                         wraplength=210, justify="left").pack(anchor="w", padx=10, pady=6)
            return

        for ruta in archivos:
            fecha = self._fecha_desde_archivo(ruta)
            titulo = self._primer_mensaje(ruta) or "Conversación"
            ctk.CTkButton(
                self.lista_conv, text=f"{fecha} · {titulo}", anchor="w",
                font=("Segoe UI", 11), fg_color="transparent", hover_color=COLOR_HOVER,
                text_color=COLOR_TEXTO, height=36, corner_radius=8,
                command=lambda r=ruta: self._abrir_conversacion(r)).pack(
                    fill="x", padx=2, pady=2)

    def _fecha_desde_archivo(self, ruta):
        m = re.search(r"(\d{4})-(\d{2})-(\d{2})", os.path.basename(ruta))
        if m:
            try:
                dt = datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
                meses = ["ene", "feb", "mar", "abr", "may", "jun",
                         "jul", "ago", "sep", "oct", "nov", "dic"]
                return f"{dt.day} {meses[dt.month - 1]} {dt.year}"
            except Exception:
                pass
        return os.path.basename(ruta)

    def _primer_mensaje(self, ruta):
        try:
            with open(ruta, "r", encoding="utf-8") as f:
                for linea in f:
                    m = re.match(r"\*\*\[\d{2}:\d{2}\]\s*(.+?):\*\*\s*(.*)", linea.strip())
                    if m and m.group(2).strip():
                        texto = m.group(2).strip()
                        return (texto[:34] + "…") if len(texto) > 34 else texto
        except Exception:
            pass
        return ""

    def _abrir_conversacion(self, ruta):
        for w in self.chat_scroll.winfo_children():
            w.destroy()
        self._ind_pensando = None
        self._pegado_abajo = True    # al abrir un hilo se muestra el final
        self._procesando = False
        self._burbujas = []
        self._seleccion = []
        self.conversacion_actual = ruta  # los mensajes nuevos siguen este hilo
        self._mostrar_vista("chat")

        pares = self._leer_conversacion(ruta)
        self.chat_activo = None
        if not pares:
            self._mostrar_estado_chat(False)
            return
        self._mostrar_estado_chat(True)
        for autor, texto in pares:
            self._insertar_burbuja(autor, texto)
        self.after(60, self._scroll_final)

    def _leer_conversacion(self, ruta):
        pares = []
        try:
            with open(ruta, "r", encoding="utf-8") as f:
                for linea in f:
                    m = re.match(r"\*\*\[\d{2}:\d{2}\]\s*(.+?):\*\*\s*(.*)", linea.strip())
                    if m:
                        pares.append((m.group(1).strip(), m.group(2).strip()))
        except Exception as e:
            print(f"Error leyendo conversación: {e}")
        return pares

    def _abrir_config_excel(self):
        """Diálogo para elegir el Excel de la matriz: archivo local o URL."""
        cfg = self._leer_config_json()
        matriz = cfg.get("matriz") if isinstance(cfg.get("matriz"), dict) else {}

        dlg = ctk.CTkToplevel(self)
        dlg.title("Configurar Excel de la matriz")
        dlg.geometry("700x580")
        dlg.configure(fg_color=COLOR_FONDO)
        dlg.transient(self)
        dlg.after(150, dlg.grab_set)

        ctk.CTkLabel(dlg, text="Excel de la matriz", font=("Segoe UI", 16, "bold"),
                     text_color=COLOR_TEXTO).pack(anchor="w", padx=24, pady=(18, 2))
        ctk.CTkLabel(dlg,
                     text=("Elige el archivo .xlsx o pega el enlace (URL) del Excel.\n"
                           "· Archivo dentro de una carpeta sincronizada de "
                           "OneDrive/SharePoint: se sube solo.\n"
                           "· URL + \"Carpeta sincronizada\": la copia se edita ahí y "
                           "OneDrive sube los cambios a SharePoint solo.\n"
                           "· URL sola: se edita una copia local y los cambios NO llegan a "
                           "SharePoint."),
                     font=("Segoe UI", 11), text_color=COLOR_TEXTO_SUAVE,
                     justify="left", wraplength=640).pack(anchor="w", padx=24,
                                                          pady=(0, 10))

        ctk.CTkLabel(dlg, text="Archivo (.xlsx) o URL:", text_color=COLOR_TEXTO,
                     font=("Segoe UI", 12)).pack(anchor="w", padx=24, pady=(6, 2))
        frame_ruta = ctk.CTkFrame(dlg, fg_color="transparent")
        frame_ruta.pack(fill="x", padx=24)
        entrada_ruta = ctk.CTkEntry(frame_ruta, fg_color=COLOR_PANEL,
                                    border_color=COLOR_BORDE, text_color=COLOR_TEXTO,
                                    height=34)
        entrada_ruta.pack(side="left", expand=True, fill="x", padx=(0, 8))
        valor = (matriz.get("excel_path") or matriz.get("excel_url") or "")
        if valor:
            entrada_ruta.insert(0, str(valor))

        def _examinar():
            actual = entrada_ruta.get().strip().strip('"')
            if not actual.lower().startswith(("http://", "https://")) \
                    and os.path.dirname(actual):
                inicial = os.path.dirname(actual)
            else:
                inicial = rutas.MATRICES
            elegido = filedialog.askopenfilename(
                title="Elegir Excel de la matriz", initialdir=inicial,
                filetypes=[("Hojas de cálculo", "*.xlsx *.xlsm"), ("Todos", "*.*")])
            if elegido:
                entrada_ruta.delete(0, "end")
                entrada_ruta.insert(0, elegido)

        ctk.CTkButton(frame_ruta, text="Examinar...", command=_examinar,
                      fg_color=COLOR_HOVER, hover_color=COLOR_BORDE,
                      text_color=COLOR_TEXTO, height=34).pack(side="left")

        def _descargar():
            """Baja de nuevo el Excel publicado en la URL (copia local actualizada)."""
            valor = entrada_ruta.get().strip().strip('"')
            if not valor.lower().startswith(("http://", "https://")):
                lbl_estado.configure(text="Escribe primero la URL del Excel.",
                                     text_color="#e06c6c")
                return
            # Se usa la carpeta sincronizada escrita aquí, aunque no se haya guardado aún.
            carpeta = os.path.expanduser(
                os.path.expandvars(entrada_sync.get().strip().strip('"')))
            if carpeta and not os.path.isdir(carpeta):
                lbl_estado.configure(
                    text=("La carpeta sincronizada no existe. Corrígela o déjala vacía para "
                          "bajar solo una copia local."),
                    text_color="#e06c6c")
                return
            lbl_estado.configure(text="Descargando...", text_color=COLOR_TEXTO_SUAVE)
            resultado = {}

            def _worker():
                try:
                    resultado["ruta"] = self.system.cloud.descargar_excel(
                        valor, forzar=True, carpeta_sincronizada=carpeta or "")
                except Exception as e:
                    self.system.cloud.error_ruta = f"No pude descargar el Excel: {e}"
                    resultado["ruta"] = None

            def _revisar():
                """Se ejecuta en el hilo de la interfaz: la descarga va en un hilo aparte."""
                try:
                    if "ruta" not in resultado:
                        dlg.after(250, _revisar)
                        return
                    bajada = resultado["ruta"]
                    if bajada:
                        prefijo = ("Copia actualizada en la carpeta sincronizada "
                                   "(OneDrive la subirá): ") if carpeta else \
                                  "Copia local actualizada: "
                        lbl_estado.configure(text=f"{prefijo}{bajada}",
                                             text_color="#2ecc71")
                    else:
                        lbl_estado.configure(
                            text=self.system.cloud.error_ruta or
                                 "No pude descargar el Excel.",
                            text_color="#e06c6c")
                except Exception:
                    pass

            threading.Thread(target=_worker, daemon=True).start()
            dlg.after(250, _revisar)

        ctk.CTkLabel(dlg, text="Hoja:", text_color=COLOR_TEXTO,
                     font=("Segoe UI", 12)).pack(anchor="w", padx=24, pady=(12, 2))
        entrada_hoja = ctk.CTkEntry(dlg, fg_color=COLOR_PANEL,
                                    border_color=COLOR_BORDE, text_color=COLOR_TEXTO,
                                    height=34)
        entrada_hoja.pack(padx=24, fill="x")
        entrada_hoja.insert(0, str(matriz.get("hoja") or "Matriz"))

        # Carpeta sincronizada (solo con URL): así los cambios se suben solos a SharePoint
        ctk.CTkLabel(dlg, text="Carpeta sincronizada de OneDrive/SharePoint (opcional):",
                     text_color=COLOR_TEXTO, font=("Segoe UI", 12)).pack(
                         anchor="w", padx=24, pady=(12, 2))
        frame_sync = ctk.CTkFrame(dlg, fg_color="transparent")
        frame_sync.pack(fill="x", padx=24)
        entrada_sync = ctk.CTkEntry(
            frame_sync, fg_color=COLOR_PANEL, border_color=COLOR_BORDE,
            text_color=COLOR_TEXTO, height=34,
            placeholder_text="C:\\Users\\tu_usuario\\OneDrive - Empresa\\Shared Documents")
        entrada_sync.pack(side="left", expand=True, fill="x", padx=(0, 8))
        valor_sync = str(matriz.get("excel_sync_dir") or "")
        if valor_sync:
            entrada_sync.insert(0, valor_sync)

        def _elegir_carpeta():
            actual = entrada_sync.get().strip().strip('"')
            inicial = actual if os.path.isdir(actual) else rutas.MATRICES
            carpeta = filedialog.askdirectory(
                title="Carpeta sincronizada de OneDrive/SharePoint", initialdir=inicial)
            if carpeta:
                entrada_sync.delete(0, "end")
                entrada_sync.insert(0, os.path.normpath(carpeta))

        ctk.CTkButton(frame_sync, text="Elegir carpeta...", command=_elegir_carpeta,
                      fg_color=COLOR_HOVER, hover_color=COLOR_BORDE,
                      text_color=COLOR_TEXTO, height=34).pack(side="left")
        ctk.CTkLabel(dlg,
                     text=("En SharePoint abre el Excel → \"Sincronizar\" y pon aquí esa "
                           "carpeta local.\nLa copia descargada se guarda en ella con el "
                           "mismo nombre, así OneDrive actualiza el archivo de SharePoint."),
                     font=("Segoe UI", 11), text_color=COLOR_TEXTO_SUAVE,
                     justify="left", wraplength=640).pack(anchor="w", padx=24, pady=(2, 0))

        lbl_estado = ctk.CTkLabel(dlg, text="", text_color=COLOR_TEXTO_SUAVE,
                                  font=("Segoe UI", 11), justify="left",
                                  wraplength=640)
        lbl_estado.pack(anchor="w", padx=24, pady=(10, 0))

        def _guardar():
            ruta_o_url = entrada_ruta.get().strip().strip('"')
            if not ruta_o_url:
                lbl_estado.configure(text="Indica un archivo o una URL.",
                                     text_color="#e06c6c")
                return
            cfg_actual = self._leer_config_json()
            matriz_cfg = cfg_actual.get("matriz")
            if not isinstance(matriz_cfg, dict):
                matriz_cfg = {}
            if ruta_o_url.lower().startswith(("http://", "https://")):
                matriz_cfg["excel_url"] = ruta_o_url
                matriz_cfg.pop("excel_path", None)
            else:
                if not os.path.isabs(ruta_o_url):
                    lbl_estado.configure(
                        text="Para archivos usa la ruta completa (C:\\...\\Matriz.xlsx).",
                        text_color="#e06c6c")
                    return
                if not os.path.exists(ruta_o_url):
                    lbl_estado.configure(text="El archivo no existe. Revisa la ruta.",
                                         text_color="#e06c6c")
                    return
                matriz_cfg["excel_path"] = ruta_o_url
                matriz_cfg.pop("excel_url", None)
            sincronizada = os.path.expanduser(
                os.path.expandvars(entrada_sync.get().strip().strip('"')))
            if sincronizada:
                if not os.path.isdir(sincronizada):
                    lbl_estado.configure(
                        text=("La carpeta sincronizada no existe. Elige una carpeta real "
                              "(usa \"Elegir carpeta...\")."),
                        text_color="#e06c6c")
                    return
                matriz_cfg["excel_sync_dir"] = sincronizada
            else:
                matriz_cfg.pop("excel_sync_dir", None)
            matriz_cfg["hoja"] = entrada_hoja.get().strip() or "Matriz"
            matriz_cfg.setdefault("columnas", list(COLUMNAS_DEFECTO))
            cfg_actual["matriz"] = matriz_cfg
            try:
                self._escribir_config_json(cfg_actual)
            except Exception as e:
                lbl_estado.configure(text=f"No pude guardar: {e}", text_color="#e06c6c")
                return
            self._actualizar_estado_nube()
            dlg.destroy()

        frame_botones = ctk.CTkFrame(dlg, fg_color="transparent")
        frame_botones.pack(fill="x", padx=24, pady=16)
        ctk.CTkButton(frame_botones, text="Guardar", command=_guardar,
                      fg_color="#2f9e6b", hover_color="#26895d", text_color="#ffffff",
                      height=36, font=("Segoe UI", 12, "bold")).pack(side="right",
                                                                     padx=(8, 0))
        ctk.CTkButton(frame_botones, text="Descargar URL ahora", command=_descargar,
                      fg_color=COLOR_ACENTO, hover_color=COLOR_ACENTO_HOVER,
                      text_color="#ffffff", height=36,
                      font=("Segoe UI", 12, "bold")).pack(side="right", padx=(8, 0))
        ctk.CTkButton(frame_botones, text="Cancelar", command=dlg.destroy,
                      fg_color=COLOR_HOVER, hover_color=COLOR_BORDE,
                      text_color=COLOR_TEXTO, height=36).pack(side="right")

    # ---------------- Configuración ----------------
    def _cargar_config_en_inputs(self):
        """Carga los valores guardados en config_usuario.json (o .env para la API key)."""
        api_key = ""
        nombre = "jarvis"
        rvc = ""
        usuario = ""
        if os.path.exists(self.config_path):
            try:
                with open(self.config_path, "r", encoding="utf-8") as f:
                    config = json.load(f)
                    api_key = config.get("api_key", "")
                    nombre = config.get("nombre_asistente", "jarvis")
                    rvc = config.get("modelo_rvc_path", "")
                    usuario = config.get("usuario", "")
            except Exception as e:
                print(f"Error al cargar configuración: {e}")

        # La API key puede venir de la variable de entorno (más segura)
        api_key_env = os.environ.get("DEEPSEEK_API_KEY", "").strip()
        if api_key_env:
            api_key = api_key_env

        usuario = usuario or os.environ.get("USERNAME", "") or "Usuario"
        self.lbl_usuario.configure(text=f"{usuario} · Free")
        self.avatar.configure(text=(usuario[:1].upper() if usuario else "J"))

        self.entry_api.delete(0, "end")
        self.entry_nombre.delete(0, "end")
        self.entry_rvc.delete(0, "end")
        self.entry_api.insert(0, api_key)
        self.entry_nombre.insert(0, nombre)
        self.entry_rvc.insert(0, rvc)

        # Matriz en la nube: se configura con sus propios botones/diálogos
        # Correo: idem (se conecta con 'Conectar correo')
        self._actualizar_estado_correo()
        self._actualizar_estado_nube()

    def guardar_config(self):
        """Guarda los datos de la interfaz en el JSON preservando claves existentes."""
        existente = {}
        if os.path.exists(self.config_path):
            try:
                with open(self.config_path, "r", encoding="utf-8") as f:
                    existente = json.load(f)
                if not isinstance(existente, dict):
                    existente = {}
            except Exception as e:
                print(f"Error leyendo configuración previa: {e}")

        # Si la API key viene de la variable de entorno, no la persistimos en el JSON
        api_key = self.entry_api.get().strip()
        api_key_env = os.environ.get("DEEPSEEK_API_KEY", "").strip()
        if api_key and api_key_env and api_key == api_key_env:
            api_key = ""

        existente["api_key"] = api_key
        existente["nombre_asistente"] = self.entry_nombre.get().strip().lower()
        existente["modelo_rvc_path"] = self.entry_rvc.get().strip()

        # La matriz en la nube (Excel + Azure DevOps) se configura con sus propios
        # botones y diálogos: aquí se conserva tal cual, sin modificarla.

        existente.setdefault("programas", {
            "discord": "discord://",
            "spotify": "spotify:",
            "vscode": "code"
        })

        os.makedirs(os.path.dirname(self.config_path), exist_ok=True)
        with open(self.config_path, "w", encoding="utf-8") as f:
            json.dump(existente, f, indent=4)

        # Re-inicializar módulos con la nueva configuración
        self.brain = APIBrain(self.config_path)
        self.voice.nombre_activacion = existente["nombre_asistente"]
        self.lbl_estado.configure(text="Configuración guardada", text_color="#2ecc71")

    # ---------------- Control por voz ----------------
    def toggle_escucha(self):
        if not self.escuchando:
            self.escuchando = True
            self.lbl_estado.configure(text="Estado: Escuchando...", text_color="#2ecc71")
            self.btn_toggle.configure(text="Detener Escucha", fg_color="#c0392b",
                                      hover_color="#a93226")
            threading.Thread(target=self._bucle_escucha, daemon=True).start()
        else:
            self.escuchando = False
            self.lbl_estado.configure(text="Estado: Inactivo", text_color="#e06c6c")
            self.btn_toggle.configure(text="Iniciar Escucha", fg_color=COLOR_ACENTO,
                                      hover_color=COLOR_ACENTO_HOVER)

    def _bucle_escucha(self):
        """Bucle principal de procesamiento de voz en hilo daemon."""
        while self.escuchando:
            comando_usuario = self.voice.escuchar_microfono()

            if comando_usuario and self.escuchando:
                # 1. Manejo de comandos de cierre
                if comando_usuario == "COMANDO_CERRAR":
                    self.reproducir_voz("Entendido, cerrando el programa. Hasta luego.")
                    self.after(0, self.destroy)
                    break
                if comando_usuario == "COMANDO_APAGAR":
                    self.reproducir_voz("Entendido, desactivando la escucha.")
                    self.after(0, self.toggle_escucha)
                    break

                # 1.5 Modo dictado continuo: escribir el texto en lugar de procesarlo
                if self.system.productivity.dictado_activo:
                    if self._es_fin_dictado(comando_usuario):
                        self.system.productivity.dictado_activo = False
                        self.reproducir_voz("Dictado desactivado.")
                    else:
                        self.system.productivity.dictar(comando_usuario)
                    continue

                print(f"Instrucción a procesar: {comando_usuario}")
                self._guardar_regla_si_aplica(comando_usuario)
                self._agregar_chat("Tú", comando_usuario)

                # 2-3-4. Procesar y responder por voz
                modelo_rvc = self.entry_rvc.get().strip()
                respuesta = self._procesar_mensaje(comando_usuario)
                if respuesta:
                    self._agregar_chat("Jarvis", respuesta)
                    self.reproducir_voz(respuesta, modelo_rvc or None)

    def _es_fin_dictado(self, texto):
        t = texto.lower()
        return any(p in t for p in ["detener dictado", "detén dictado", "fin del dictado",
                                    "salir del dictado", "terminar dictado"])

    def _metadatos_reunion(self):
        """Título e invitados de la reunión en curso, leídos del calendario.

        La grabación los guarda en el acta, y con ellos se puede enviar el acta a
        los asistentes sin tener que dictar los correos. Si no hay calendario
        disponible se devuelve un dict vacío y la grabación sigue igual.
        """
        datos = {}
        try:
            evento, _mensaje = self.system.calendario.evento_en_curso()
        except Exception as e:
            print(f"No pude leer el calendario para la reunión: {e}")
            return datos
        if evento:
            datos["titulo"] = evento.get("asunto", "")
            datos["asistentes"] = list(evento.get("correos") or [])
        return datos

    def _resumir_reunion(self, transcripcion):
        """Genera el acta de la reunión con la estructura fija de Jarvis."""
        prompt = (
            "Analiza esta transcripción de una reunión y escribe el ACTA en Markdown, "
            "en español, con EXACTAMENTE estas secciones y en este orden:\n"
            "## Objetivo\n"
            "## Puntos tratados\n"
            "## Acuerdos y decisiones\n"
            "## Tareas\n"
            "Aquí va una tabla Markdown con las columnas: Tarea | Responsable | Fecha. "
            "Si el responsable o la fecha no se mencionan, escribe 'por definir'.\n"
            "## Riesgos y bloqueos\n"
            "## Próxima reunión\n"
            "Reglas: no inventes nombres, fechas ni datos que no aparezcan en la "
            "transcripción; si una sección no tiene contenido, escribe 'Sin datos en la "
            "grabación'; no repitas la transcripción; no escribas nada antes del primer "
            "'##'.\n\nTranscripción:\n" + transcripcion
        )
        return self.brain.analizar(prompt)

    def _generar_preventa(self, tipo, contenido):
        prompts = {
            "propuesta": "Genera una propuesta técnica profesional para un proyecto de software con: alcance, objetivos, solución propuesta, arquitectura, plan de trabajo y entregables. Usa Markdown.",
            "requisitos": "Genera una matriz de cumplimiento de requisitos (tabla) con columnas: ID, Requisito, Prioridad, Tipo (funcional/no funcional) y Observaciones. Usa Markdown.",
            "historias": "Genera historias de usuario en formato 'Como [rol], quiero [funcionalidad], para [beneficio]', cada una con criterios de aceptación. Usa Markdown.",
            "glosario": "Genera un glosario de términos técnicos del contenido, con definiciones claras y concisas. Usa Markdown.",
            "resumen": "Genera un resumen ejecutivo con: objetivo, alcance, riesgos y recomendaciones. Usa Markdown.",
            "estimacion": "Genera una estimación de esfuerzo por fases (análisis, diseño, desarrollo, pruebas, despliegue) con horas y supuestos. Usa Markdown.",
        }
        prompt = prompts.get(tipo, "Procesa este contenido y genera un documento estructurado en Markdown.") + "\n\nContenido:\n" + contenido
        return self.brain.analizar(prompt)

    def reproducir_voz(self, texto, modelo_rvc=None):
        """Reproduce texto con edge-tts (voz española) y cae a pyttsx3 offline si falla."""
        if not texto:
            return
        try:
            self.voice.hablar(texto, rvc_model_path=modelo_rvc)
        except Exception as e:
            print(f"Error en edge-tts, usando pyttsx3 offline: {e}")
            if hasattr(self.voice_output, "hablar"):
                self.voice_output.hablar(texto)
